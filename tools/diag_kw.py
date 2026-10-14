"""적요 키워드 패턴 진단: verify.py --excel 이 만든 엑셀 계산 결과와 파이썬 기대값을 키워드별로 나란히 보여 준다.

사용: py tools/verify.py --excel  다음에  py tools/diag_kw.py
"""
import contextlib, io, os, runpy, sys
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "tools", "_out")
calc = [f for f in os.listdir(OUT) if f.startswith("calc_")][0]
wb = openpyxl.load_workbook(os.path.join(OUT, calc), data_only=True)
fx = openpyxl.load_workbook(os.path.join(OUT, calc[5:]), data_only=False)
kw, ca = wb["키워드사전"], wb["원인분석"]
print("원인분석 C9", ca["C9"].value, "C10~C13", [ca[c].value for c in ("C10", "C11", "C12", "C13")], "C16(증가분)", ca["C16"].value, "Z34", ca["Z34"].value, "Z35", ca["Z35"].value)
print("H5 수식:", fx["키워드사전"]["H5"].value)
sys.argv = ["build_datajs.py", os.devnull]
with contextlib.redirect_stdout(io.StringIO()):
    G = runpy.run_path(os.path.join(ROOT, "src", "build_datajs.py"))
E, P = G["E"], G["P"]
k = f'{ca["C9"].value}'
cur = [e for e in E if e["key"] == k and ca["C10"].value <= e["p"] <= ca["C11"].value]
print(f"당기 전표 {len(cur)}건 (파이썬)")
for r in range(5, 12):
    ks = [kw.cell(row=r, column=c).value for c in range(2, 7)]
    py = max([0] + [sum(e["v"] for e in cur if kk and kk in e["text"]) for kk in ks if kk])
    print(f"{r}행 {kw.cell(row=r, column=1).value}: 엑셀 H={kw.cell(row=r, column=8).value} I={kw.cell(row=r, column=9).value} · 파이썬 {py:.1f} · 키워드 {[x for x in ks if x]}")
je = wb["전표_원본"]
print("당기 표시(Q열) 전표 예시:")
n = 0
for row in je.iter_rows(min_row=5, max_row=je.max_row):
    if row[16].value == "당기":
        print("  ", row[8].value, "|", row[15].value); n += 1
        if n >= 6: break
