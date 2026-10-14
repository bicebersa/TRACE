"""TRACE 엑셀(TRACE_v0.5.xlsx)을 예시 데이터와 함께 생성한다.

사용: python src/build_trace.py [출력폴더]   (기본: 저장소 루트)
환경변수: JE_END(전표 수식 범위), SEL(검증용 원인분석 선택), EMBED_JSON(data.js 내장)
필요: pip install openpyxl
"""
import datetime as dt, json, os, random, sys, calendar
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.chart import BarChart, Reference
from openpyxl.utils import get_column_letter as CL

OUTDIR = sys.argv[1] if len(sys.argv) > 1 else os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
os.makedirs(OUTDIR, exist_ok=True)
XLSX = os.path.join(OUTDIR, "TRACE_v0.5.xlsx")
HERE = os.path.dirname(os.path.abspath(__file__))

FONT = "맑은 고딕"
NUM = '#,##0.0;"△"#,##0.0;"-"'
NUM0 = '#,##0;"△"#,##0;"-"'
PCT = '0%;"△"0%;"-"'
RATE = '+0%;-0%;0%'
NAVY = PatternFill("solid", fgColor="1F3864")
GRAYH = PatternFill("solid", fgColor="808080")
SUB = PatternFill("solid", fgColor="D9E1F2")
TOT = PatternFill("solid", fgColor="F2F2F2")
YEL = PatternFill("solid", fgColor="FFFF00")
BOXF = PatternFill("solid", fgColor="EAF3EE")
thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
BLUE, BLACK, GREEN, GRAY = "0000FF", "000000", "008000", "7F7F7F"

def F(color=BLACK, bold=False, size=10):
    return Font(name=FONT, color=color, bold=bold, size=size)
def title(ws, t, s):
    ws["A1"] = t if t.startswith("TRACE") else "TRACE · " + t; ws["A1"].font = F(bold=True, size=14)
    ws["A2"] = s; ws["A2"].font = F(GRAY, size=9)
def header(ws, row, cols, start=1, fill=NAVY, h=30):
    for i, v in enumerate(cols):
        c = ws.cell(row=row, column=start + i, value=v)
        c.font = Font(name=FONT, color="FFFFFF", bold=True, size=10)
        c.fill = fill; c.border = BOX
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = h
def put(ws, r, c, v, color=BLACK, fmt=None, bold=False, fill=None, align=None, wrap=False):
    cell = ws.cell(row=r, column=c, value=v)
    cell.font = F(color, bold); cell.border = BOX
    if fmt: cell.number_format = fmt
    if fill: cell.fill = fill
    if align or wrap: cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    return cell
def widths(ws, w):
    for i, x in enumerate(w, 1): ws.column_dimensions[CL(i)].width = x
def sec(ws, r, text):
    ws.cell(row=r, column=1, value=text).font = F(bold=True, size=11)

# =================================================================== 예시 데이터
rng = random.Random(int(os.environ.get("SEED", 11)))   # SEED를 바꾸면 숫자가 다른 예시 (패턴 시나리오는 같음)
Y, M = 2026, 9
PNOW = Y * 12 + M
MONTHS = [(2025, m) for m in range(1, 13)] + [(2026, m) for m in range(1, M + 1)]
ITEMS = {"B01": "여비교통비", "B02": "통신비", "B03": "소모품비", "B04": "도서인쇄비", "B05": "회의비", "B06": "업무추진비",
         "B07": "광고선전비", "B08": "교육훈련비", "B09": "지급수수료", "B10": "전산운영비", "B11": "임차료", "B12": "차량유지비", "B13": "판매수수료", "B14": "손해조사비"}
GLS = [("62101000", "여비교통비-국내", "B01", 0.85), ("62102000", "여비교통비-해외", "B01", 0.15),
       ("62201000", "통신비", "B02", 1.0), ("62301000", "소모품비", "B03", 1.0), ("62401000", "도서인쇄비", "B04", 1.0),
       ("62501000", "회의비", "B05", 1.0), ("62601000", "업무추진비", "B06", 1.0), ("62701000", "광고선전비", "B07", 1.0),
       ("62801000", "교육훈련비-사내", "B08", 0.4), ("62802000", "교육훈련비-외부위탁", "B08", 0.6),
       ("62901000", "지급수수료-컨설팅", "B09", 0.7), ("62902000", "지급수수료-기타", "B09", 0.3),
       ("63001000", "전산운영비-SW유지보수", "B10", 0.55), ("63002000", "전산운영비-클라우드", "B10", 0.45),
       ("63101000", "임차료", "B11", 1.0), ("63201000", "차량유지비", "B12", 1.0),
       ("63501000", "모집수수료", "B13", 0.8), ("63502000", "유지수수료", "B13", 0.2), ("63601000", "손해조사비", "B14", 1.0)]
GLNAME = {g[0]: g[1] for g in GLS}
DEPTS = {"D100": "경영기획팀", "D200": "재무회계팀", "D300": "인사팀", "D400": "상품개발팀",
         "D500": "마케팅팀", "D600": "영업지원팀", "D700": "IT운영팀", "D800": "고객서비스팀", "D900": "영업채널팀", "D950": "자산운용팀"}
CCS = [("1101", "경영기획", "D100", 1.0), ("1201", "재무회계", "D200", 1.0), ("1301", "인사", "D300", 1.0),
       ("1401", "상품개발", "D400", 1.0), ("1501", "마케팅전략", "D500", 0.6), ("1502", "디지털마케팅", "D500", 0.4),
       ("1601", "영업지원-본사", "D600", 0.55), ("1602", "영업지원-지역본부", "D600", 0.45),
       ("1701", "IT운영", "D700", 1.0), ("1801", "고객서비스", "D800", 1.0),
       ("1901", "GA채널", "D900", 0.6), ("1902", "전속채널", "D900", 0.4), ("1951", "자산운용", "D950", 1.0)]
BUDGET = {"D100": {"B01": 36000, "B02": 9600, "B03": 6000, "B05": 12000, "B06": 48000, "B09": 240000},
          "D200": {"B01": 18000, "B02": 7200, "B03": 7200, "B05": 6000, "B09": 420000},
          "D300": {"B01": 24000, "B02": 8400, "B03": 9600, "B05": 9600, "B06": 24000, "B08": 480000},
          "D400": {"B01": 30000, "B02": 7200, "B03": 4800, "B04": 12000, "B05": 7200, "B09": 360000},
          "D500": {"B01": 42000, "B02": 12000, "B03": 9600, "B04": 84000, "B05": 14400, "B06": 36000, "B07": 1800000},
          "D600": {"B01": 96000, "B02": 26400, "B03": 14400, "B05": 18000, "B06": 42000, "B11": 216000, "B12": 96000},
          "D700": {"B01": 14400, "B02": 180000, "B03": 12000, "B05": 4800, "B10": 2400000},
          "D800": {"B01": 12000, "B02": 240000, "B03": 18000, "B04": 96000, "B05": 6000, "B09": 600000, "B14": 420000},
          "D900": {"B13": 30000000},
          "D950": {"B01": 24000, "B02": 12000, "B09": 360000, "B10": 240000}}  # 천원
SEASON = {"B07": [0.6, 0.8, 1.4, 1.2, 1, 0.8, 0.7, 0.8, 1.4, 1.3, 1.0, 1.0]}
POOL = {"62101000": ["L여행사", "M항공", "(법인카드)"], "62102000": ["M항공", "N호텔"], "62201000": ["O텔레콤", "P통신"],
        "62301000": ["Q오피스", "R문구"], "62401000": ["G인쇄", "S출판"], "62501000": ["(법인카드)", "T케이터링"],
        "62601000": ["(법인카드)"], "62701000": ["E미디어", "F애드"], "62801000": ["U에듀", "(사내강사료)"],
        "62802000": ["I교육원", "V아카데미"], "62901000": ["K회계법인", "W컨설팅"], "62902000": ["J콜센터", "X서비스"],
        "63001000": ["Y소프트", "Z시스템"], "63002000": ["C클라우드"], "63101000": ["AA빌딩", "BB부동산"],
        "63201000": ["D정비", "H에너지"], "63501000": ["AA보험대리점", "BB에셋", "(전속설계사)"], "63502000": ["(전속설계사)", "CC금융서비스"], "63601000": ["DD손해사정", "EE손해사정"]}
SHORT = {"62101000": "국내출장", "62102000": "해외출장", "62201000": "통신요금", "62301000": "소모품 구입", "62401000": "인쇄물 제작",
         "62501000": "회의 다과", "62601000": "업무추진", "62701000": "광고 집행", "62801000": "사내교육", "62802000": "위탁교육",
         "62901000": "자문 수수료", "62902000": "업무 위탁 수수료", "63001000": "SW 유지보수", "63002000": "클라우드 사용료",
         "63101000": "사무실 임차료", "63201000": "차량 유지", "63501000": "모집수수료", "63502000": "유지수수료", "63601000": "손해사정 수수료"}

bud_gl = {}  # (dept, gl) -> 12개월 예산(원)
for d, items in BUDGET.items():
    for it, ann in items.items():
        w = SEASON.get(it, [1] * 12)
        for gl, gln, ii, gw in GLS:
            if ii != it: continue
            m = [round(ann * 1000 * gw * x / sum(w), -3) for x in w]
            bud_gl[(d, gl)] = m

vcode = {}
def vc(name):
    if name not in vcode: vcode[name] = f"V{10001 + len(vcode)}"
    return vcode[name]

E = []  # entries
def cc_of(d):
    lst = [(c[0], c[3]) for c in CCS if c[2] == d]
    x = rng.random(); acc = 0
    for cc, w in lst:
        acc += w
        if x <= acc: return cc
    return lst[-1][0]
def add(y, m, d, gl, amt, vname, text=None, day=None, typ="KR", cc=None):
    day = day or rng.randint(2, 27)
    E.append(dict(y=y, m=m, date=dt.date(y, m, day), typ=typ, gl=gl, cc=cc or cc_of(d), amt=round(amt, -3),
                  text=text or f"{m}월 {SHORT[gl]} - {vname}", vname=vname, d=d))
def spread(y, m, d, gl, total, pool=None):
    pool = pool or POOL[gl]
    names = []
    for v in pool:
        names += [v] * rng.randint(1, 2)
    w = [rng.uniform(0.6, 1.4) for _ in names]
    for v, x in zip(names, w):
        add(y, m, d, gl, total * x / sum(w), v)

CUSTOM = {("D500", "62701000"), ("D700", "63002000"), ("D600", "63201000"), ("D300", "62802000"), ("D800", "62902000"),
          ("D900", "63501000"), ("D900", "63502000")}
for (d, gl), bm in bud_gl.items():
    if (d, gl) in CUSTOM or d == "D950" or gl == "63601000": continue
    f = rng.uniform(0.80, 0.92); g = rng.uniform(0.78, 0.92)
    for y, m in MONTHS:
        spread(y, m, d, gl, bm[m - 1] * (g if y == 2025 else f) * rng.uniform(0.88, 1.12))

MIL = round(1_000_000 * float(os.environ.get("SCALE", 1)))   # SCALE: 예시 금액 배율 (다른 예시를 만들 때)
# S1 마케팅팀 광고선전비: 9월 단건 대형(TV캠페인) + 신규 거래처
s1_25 = [150, 140, 160, 150, 140, 150, 160, 150, 180, 190, 200, 201.4]
s1_26 = [190, 190, 190, 140, 130, 140, 190, 120]
for m, v in enumerate(s1_25, 1): spread(2025, m, "D500", "62701000", v * MIL)
for m, v in enumerate(s1_26, 1): spread(2026, m, "D500", "62701000", v * MIL)
spread(2026, 9, "D500", "62701000", 140 * MIL)
add(2026, 9, "D500", "62701000", 120 * MIL, "A광고대행사", "3분기 TV캠페인 집행분", day=18, cc="1501")
add(2026, 9, "D500", "62701000", 30 * MIL, "B디지털", "디지털 퍼포먼스 광고 9월", day=25, cc="1502")
# S2 IT운영팀 클라우드: 같은 건수, 금액 상승 (단가 인상)
for m, v in enumerate([64, 65, 66, 66, 67, 68, 68, 69, 70, 70, 71, 72], 1):
    add(2025, m, "D700", "63002000", v * MIL, "C클라우드", day=25)
for m, v in enumerate([78, 79, 78, 80, 79, 80, 86, 93, 101], 1):
    add(2026, m, "D700", "63002000", v * MIL, "C클라우드", day=25)
# S3 영업지원팀 차량유지비: 8월 누락 → 9월 2개월치 (시점 이동)
for y, m in MONTHS:
    base = 5.5 if y == 2025 else 6.0
    if (y, m) == (2026, 8):
        pass
    elif (y, m) == (2026, 9):
        add(y, m, "D600", "63201000", 12.4 * MIL, "D정비", "8~9월분 업무차량 정비비", day=20)
    else:
        add(y, m, "D600", "63201000", base * rng.uniform(0.95, 1.05) * MIL, "D정비")
    add(y, m, "D600", "63201000", (1.8 if y == 2025 else 2.0) * rng.uniform(0.95, 1.05) * MIL, "H에너지")
# S4 인사팀 외부위탁교육: 연간계약 일시 지급 (키워드)
for y, m in MONTHS:
    spread(y, m, "D300", "62802000", (14 if y == 2025 else 11) * rng.uniform(0.9, 1.1) * MIL)
add(2026, 9, "D300", "62802000", 60 * MIL, "I교육원", "2026 하반기 위탁교육 연간계약 일시 지급", day=15)
# S5 고객서비스팀 기타수수료: 건수 증가 (물량)
for y, m in MONTHS:
    n = 9 if y == 2025 else (14 if m == 9 else 10)
    for _ in range(n):
        add(y, m, "D800", "62902000", 1.8 * rng.uniform(0.97, 1.03) * MIL, "J콜센터", f"{m}월 콜 위탁 수수료(건별)")
    add(y, m, "D800", "62902000", 3.0 * MIL, "X서비스")
# S6 재무회계팀 컨설팅: 수기 대체(재분류) + 정정·역분개
add(2026, 9, "D200", "62901000", 25 * MIL, "(내부대체)", "경영기획팀 회계자문료 재분류 (8월분)", day=28, typ="SA", cc="1201")
add(2026, 9, "D100", "62901000", -25 * MIL, "(내부대체)", "재무회계팀으로 회계자문료 재분류 (8월분)", day=28, typ="SA", cc="1101")
add(2026, 9, "D200", "62901000", 8.0 * MIL, "K회계법인", "K회계법인 자문료 9월", day=10, cc="1201")
add(2026, 9, "D200", "62901000", -8.0 * MIL, "K회계법인", "K회계법인 자문료 9월 정정(역분개)", day=12, cc="1201")
add(2026, 9, "D200", "62901000", 8.5 * MIL, "K회계법인", "K회계법인 자문료 9월 (정정분)", day=12, cc="1201")
# S7 인사팀 광고선전비: 예산 없는 집행
for m, v in ((3, 4.8), (4, 3.2), (9, 5.5)):
    add(2026, m, "D300", "62701000", v * MIL, "LL채용", "경력직 채용공고", day=14, cc="1301")

rng2 = random.Random(31)
_save = rng; rng = rng2
for gl_ in ("63501000", "63502000"):
    bm = bud_gl[("D900", gl_)]
    for y, m in MONTHS:
        spread(y, m, "D900", gl_, bm[m - 1] * (0.86 if y == 2025 else 0.9) * rng.uniform(0.985, 1.015))
# S8 영업채널팀 모집수수료: 9월 신상품 GA 판매시책비 (신계약비 일회성 증가 → GMM CSM 경험조정)
add(2026, 9, "D900", "63501000", 650 * MIL, "AA보험대리점", "9월 신상품 GA 판매시책비 (추가 수수료)", day=26, cc="1901")
rng = random.Random(37)   # 자산운용팀: 별도 난수
for (d_, gl_), bm in bud_gl.items():
    if d_ != "D950": continue
    for y, m in MONTHS:
        spread(y, m, "D950", gl_, bm[m - 1] * (0.85 if y == 2025 else 0.88) * rng.uniform(0.95, 1.05))
rng = random.Random(43)   # 손해조사비(직접): 별도 난수 — 다른 계정 예시값이 바뀌지 않게
for (d_, gl_), bm in bud_gl.items():
    if gl_ != "63601000": continue
    for y, m in MONTHS:
        spread(y, m, d_, gl_, bm[m - 1] * (0.84 if y == 2025 else 0.87) * rng.uniform(0.93, 1.07))
rng = _save

E.sort(key=lambda e: (e["date"], e["gl"], e["cc"]))
seq = {}
for e in E:
    seq[e["y"]] = seq.get(e["y"], 0) + 1
    e["doc"] = f"{(5100000000 if e['typ'] == 'KR' else 100000000) + seq[e['y']]}"
    vc(e["vname"])

ADJ = {("D700", "63002000", 3): 50 * MIL, ("D300", "62802000", 3): -20 * MIL}
def actual(d, gl, y, months):
    return sum(e["amt"] for e in E if e["d"] == d and e["gl"] == gl and e["y"] == y and e["m"] in months)
BP = []  # 예산판 행
for (d, gl), bm in sorted(bud_gl.items()):
    for q in range(1, 5):
        ms = range(3 * q - 2, 3 * q + 1)
        alloc = sum(bm[m - 1] for m in ms); adj = ADJ.get((d, gl, q), 0)
        act = actual(d, gl, Y, [m for m in ms if m <= M])
        BP.append([Y, d, DEPTS[d], int(gl), GLNAME[gl], q, alloc, adj, act, alloc + adj - act])

for d, gl in sorted({(e["d"], e["gl"]) for e in E if e["y"] == Y} - set(bud_gl)):
    for q in range(1, 5):
        act = actual(d, gl, Y, [m for m in range(3 * q - 2, 3 * q + 1) if m <= M])
        BP.append([Y, d, DEPTS[d], int(gl), GLNAME[gl], q, 0, 0, act, -act])

TB = []
for y, m in MONTHS:
    for gl in [g[0] for g in GLS] + ["61101000", "61201000"]:
        if gl.startswith("61"):
            dr = round((800 if gl == "61101000" else 120) * rng.uniform(0.97, 1.03) * MIL, -3); cr = 0
        else:
            dr = sum(e["amt"] for e in E if e["gl"] == gl and (e["y"], e["m"]) == (y, m) and e["amt"] > 0)
            cr = -sum(e["amt"] for e in E if e["gl"] == gl and (e["y"], e["m"]) == (y, m) and e["amt"] < 0)
        TB.append([y, m, int(gl), GLNAME.get(gl, "급여" if gl == "61101000" else "복리후생비"), dr, cr])
bal = {}
for row in TB:
    k = (row[0], row[2]); bal[k] = bal.get(k, 0) + row[4] - row[5]; row.append(bal[k])

KEYS = sorted({(d, gl) for d, gl in bud_gl} | {(e["d"], e["gl"]) for e in E if e["y"] == Y})

# =================================================================== 레이아웃 상수
JE_END = int(os.environ.get("JE_END", 100000))
TB_END, BP_END, VA_END = 20000, 5000, 506
JR = lambda c: f"전표_원본!${c}$5:${c}${JE_END}"
TR = lambda c: f"TB_원본!${c}$5:${c}${TB_END}"
BR = lambda c: f"예산판_원본!${c}$5:${c}${BP_END}"
VR = lambda c: f"변동분석!${c}$7:${c}${VA_END}"
S = dict(co="설정!$C$5", yr="설정!$C$6", mon="설정!$C$7", ur="설정!$C$9", amt="설정!$C$10", warn="설정!$C$11",
         short="설정!$C$12", z="설정!$C$13", one="설정!$C$14", new="설정!$C$15", sa="설정!$C$16", kw="설정!$C$17",
         price="설정!$C$18", burn="설정!$C$19", P="설정!$C$20", Q="설정!$C$21", qs="설정!$C$22", fx="설정!$C$23")
JE_HDR = ["회계연도", "기간", "전표번호", "전기일", "전표유형", "G/L계정", "코스트센터", "금액", "적요", "거래처코드", "거래처명"]
TB_HDR = ["회계연도", "기간", "G/L계정", "계정명", "당기차변", "당기대변", "기말잔액"]
BP_HDR = ["회계연도", "부서코드", "부서명", "G/L계정", "계정명", "분기", "배정", "조정", "실적", "잔여"]
def plabel(ref):  # 기간번호 → "2026-09"
    return f'INT(({ref}-1)/12)&"-"&TEXT(MOD({ref}-1,12)+1,"00")'

wb = Workbook()

# =================================================================== 안내
ws = wb.active; ws.title = "안내"
title(ws, "TRACE 사업비 원인추적기 v0.5 (프로토타입)", "Excel 2016 호환 · 계산은 수식, 자동화는 VBA(modTrace.bas) · 예시 데이터는 가상 수치 (원본 단위: 원, 보고 단위: 백만원)")
r = 4
sec(ws, r, "이름의 뜻"); r += 1
header(ws, r, ["글자", "뜻", "도구에서 하는 일 (시트)"], h=20); r += 1
for a, b, c in [("T", "Trend 추이", "전월·전분기·전년동월 대비 변동 (변동분석 · 월별집계 · 계정변동)"),
                ("R", "Risk 예산 위험", "부서 × 계정 × 분기 한도 점검(변동분석 AJ~AO열)과 연간 전망(P~Y열)"),
                ("A", "Anomaly 이상치", "평소 변동폭을 벗어난 항목을 점수로 골라냄 (주의대상)"),
                ("C", "Cause 원인", "증가분을 거래처·전표로 쪼개 원인 후보와 근거 전표 제시 (원인분석 · 키워드사전)"),
                ("E", "Expense 사업비", "분석 대상: 사업비 계정 (계정매핑에서 사업비 여부 Y)")]:
    put(ws, r, 1, a, bold=True, align="center"); put(ws, r, 2, b, bold=True); put(ws, r, 3, c, GRAY); r += 1
r += 1
sec(ws, r, "흐름 단계와 시트"); r += 1
header(ws, r, ["단계", "하는 일", "시트"], h=20); r += 1
for a, b, c in [("① 자료 넣기", "SAP 예산판·TB·전표·BS·PL을 원본 시트 A열부터 붙여넣기", "예산판_원본 · TB_원본 · 전표_원본 · BS_원본 · PL_원본"),
                ("② 불러오기·대사", "매핑 적용, 전표 = TB = 예산판 실적 대사, 데이터 점검", "자가진단 · 계정변동"),
                ("③ 변동 분석", "전월·전분기·전년동월 대비 증감 (부서 × 계정)", "변동분석 F~O열 · 월별집계 · 계정변동"),
                ("④ 예산 점검", "분기 한도(배정+조정) 대비 분기 실적·분기말 예상 → 여유/주의/초과 예상/한도 초과, 연간 전망은 참고", "변동분석 AJ~AO열 · P~Y열"),
                ("⑤ 주의 대상", "증가액·증가율·이상치·예산 점수 합산 → TOP 10", "주의대상 · 변동분석 AA열 이후"),
                ("⑥ 원인 추정", "선택 항목을 거래처·전표로 분해, 8가지 패턴 판정, 근거 전표", "원인분석 · 키워드사전"),
                ("⑦ 보고·확인", "보고서·대시보드, 확인된 원인을 코멘트로 누적", "리포트 · 대시보드 · 차트데이터 · 코멘트 · TRACE_대시보드.html"),
                ("⑧ IFRS 17 연결", "계정대체(직접대체 → 재산관리비 → 부서별 비율) → 상품·포트폴리오 → 예정 대비 경험차이 → CSM(GMM)·당기손익, 계리 변동표와 대사", "IFRS17_사업비 · 계정대체기준 · 계정매핑 · 예정사업비 · 부채변동표")]:
    put(ws, r, 1, a, bold=True); put(ws, r, 2, b); put(ws, r, 3, c, GRAY); r += 1
r += 1
sec(ws, r, "매월 사용 순서"); r += 1
for ln in ["1. 설정 시트에서 기준연도·기준월 변경",
           "2. [예산판·TB·전표·BS·PL 불러오기] 버튼으로 SAP에서 내려받은 파일 선택 (매크로 없이: 원본 시트 A열부터 붙여넣고 보조수식 끌어 채우기)",
           "3. 자가진단이 '전체 정상'인지 확인",
           "4. 주의대상 TOP 10 확인 → 원인분석 시트 C5에 순위 입력 (비교 기준은 C6) → 원인 후보와 근거 전표 확인",
           "5. 담당 부서 확인 결과를 코멘트 시트에 입력 → 주의대상·리포트에 표시",
           "6. 새로 알게 된 적요 표현은 키워드사전에 추가 → 다음 달 원인 추정에 반영",
           "7. [예정사업비·부채변동표 불러오기] (분기·월 결산 후) → IFRS17_사업비 시트에서 CSM·손익 영향과 근거 전표 확인",
           "8. [대시보드 내보내기] 버튼 → 같은 폴더에 data.js 생성 → TRACE_대시보드.html을 열면 최신 결과 표시"]:
    ws.cell(row=r, column=1, value=ln).font = F(); r += 1
r += 1
sec(ws, r, "매크로 설치 (1회)"); r += 1
for ln in ["1. Alt+F11 → [파일] → [파일 가져오기] → modTrace.bas", "2. Alt+F8 → SetupButtons 실행 → 설정 시트 오른쪽에 버튼 10개 생성",
           "3. [다른 이름으로 저장] → 'Excel 매크로 사용 통합 문서 (*.xlsm)'", "4. TRACE_대시보드.html을 이 엑셀 파일과 같은 폴더에 둡니다"]:
    ws.cell(row=r, column=1, value=ln).font = F(); r += 1
r += 1
sec(ws, r, "입력 양식 바꾸는 법 (실제 SAP 양식이 정해졌을 때)"); r += 1
for ln in ["1. 열 이름·순서가 다름 → 헤더설정 시트의 '찾을 헤더 이름'을 실제 머리글로 고치고 [OO 불러오기] 버튼 사용 (열 순서는 상관없음)",
           "2. 같은 열을 부르는 이름이 여러 개 → '대체 이름'에 쉼표로 추가 (예: 금액, DMBTR, 현지통화금액)",
           "3. 전표가 차대구분(S/H) + 양수 금액 형식 → 입력옵션 시트에서 '전표 금액 방식'을 '차대구분+양수금액'으로",
           "4. BS·PL 항목코드가 회사 코드 → 입력옵션 시트의 항목 매핑에 실제 항목코드(또는 항목명) 입력",
           "5. 계정·부서 체계가 다름 → 계정매핑(사업비 여부, IFRS17 대체구분·직접대체 계정·공통비 성격)·부서매핑(코스트센터 → 부서)·계정대체기준(재산관리비 비율, 부서별 비율) 수정",
           "6. 선택 열이 없음 → 비워둬도 됩니다. 빠지는 기능은 헤더설정 '없을 때' 칸에 적혀 있습니다",
           "7. BS·PL을 감독원 양식(항목이 행, 금액이 열)으로 받음 → 입력옵션 'BS·PL 양식 형태'를 '감독원 (가로)'로, 당기·전기·누계 열 이름 지정 (기준월 한 달씩 합쳐짐)",
           "8. 감독원 양식의 항목코드·변동요인코드 → 입력옵션의 항목 매핑·변동요인 매핑에 실제 코드 입력 (예시 감독원 양식 코드가 기본값)",
           "9. 표준에 없는 새 열을 분석에 쓰고 싶음 → 외부망에서 버전업이 필요합니다 (src/build_trace.py 수정)",
           "매크로 없이 쓸 때는 원본 시트 4행 머리글 순서대로 열을 맞춰 붙여넣으면 됩니다."]:
    ws.cell(row=r, column=1, value=ln).font = F(); r += 1
r += 1
sec(ws, r, "원본 데이터 기준"); r += 1
for ln in ["예산판: 부서 × G/L × 분기 한 행 (배정·조정·실적·잔여). 화면이 분기를 옆으로 펼치는 형태면 알려주세요. 양식을 맞추겠습니다.",
           "TB: G/L × 월 한 행 (당기차변·당기대변·기말잔액). 당해와 전년을 함께 넣어야 전년동월 비교가 됩니다.",
           "전표: 사업비 계정 전표 라인 (최근 24개월 권장). 적요·거래처가 있어야 원인 추정이 됩니다.",
           "BS: 기간 × 항목 한 행 (기말잔액). PL: 기간 × 항목 한 행 (당월·누계, 비용은 양수). 항목코드는 BS A000/L000/E000, PL P100~P600 기준이며 실제 양식에 맞춰 바꿀 수 있습니다.",
           "거래처명이 '('로 시작하면 내부 거래(법인카드·내부대체 등)로 보고 신규 거래처 판정에서 뺍니다."]:
    ws.cell(row=r, column=1, value=ln).font = F(); r += 1
r += 1
sec(ws, r, "셀 색상"); r += 1
ws.cell(row=r, column=1, value="파란 글씨 = 입력값 · 검정 = 수식 · 초록 = 다른 시트 참조 수식 · 노란 배경 = 바꾸는 설정 · 회색 머리글 = 보조수식(수정 금지)").font = F(); r += 2
sec(ws, r, "한계와 전제"); r += 1
for ln in ["원인분석은 한 번에 한 항목씩 계산합니다. TOP 10 전체 원인을 리포트에 자동으로 채우는 기능은 VBA로 추가 예정입니다.",
           "변동분석의 부서·계정 목록은 고정입니다. 새 조합이 생기면 자가진단이 알려주며, 행을 직접 추가해야 합니다 (VBA 버전에서 자동).",
           "원인 후보는 규칙 기반 가설입니다. 반드시 근거 전표와 담당 부서 확인으로 확정하세요.",
           "IFRS 17 연결은 계정대체 규칙과 포트폴리오 비율로 전표를 나눈 추정입니다. 계리 변동표(경험조정-취득현금흐름)와 대사해 확인하세요. 측정모형별 처리는 회사 회계정책을 따릅니다.",
           "포트폴리오: GMM(보장성·저축성)과 VFA(변액)는 신계약비 차이가 CSM을 조정하고 유지·수금비 차이는 당기손익, PAA(단체)는 손익 영향만 봅니다. 변액은 특별계정 구분계리 대상이지만 손익에 영향을 주므로 분석에 포함합니다."]:
    ws.cell(row=r, column=1, value=ln).font = F(); r += 1
widths(ws, [18, 60, 62])

# =================================================================== 설정
st = wb.create_sheet("설정")
title(st, "설정", "노란 칸만 바꾸세요")
header(st, 4, ["항목", "값", "설명"], start=2, h=20)
SROWS = [(5, "회사명", os.environ.get("COMPANY", "예시보험㈜"), None, ""), (6, "기준연도", Y, "0", ""), (7, "기준월", M, "0", "1~12"),
         (8, "원본 금액단위", "원", None, "예산판·TB·전표 공통"), (9, "보고 단위", "백만원", None, "모든 분석 표시 단위"),
         (10, "증감 기준금액 (보고단위)", 10, NUM0, "이 금액 미만의 작은 증감은 주의 점수를 30%로 낮춤 · 증가액 점수 만점 = 5배"),
         (11, "예산 '주의' 기준 (예상 소진율)", 0.9, PCT, "예상 연간 실적 ÷ 가용예산이 이 값 초과면 주의"),
         (12, "예산 '부족 예상' 기준", 1.0, PCT, "이 값 초과면 부족 예상"),
         (13, "이상치 기준 (z)", 2.0, "0.0", "직전 12개월 평균 대비 표준편차 배수"),
         (14, "원인: 단건 비중", 0.5, PCT, "전표 1건이 증가분의 이 비율 이상이고, 그 거래처 평소 월평균의 3배 이상이면 '일회성 대형' (내부 거래처 제외)"),
         (15, "원인: 신규 거래처 비중", 0.3, PCT, "최근 12개월 거래 없던 거래처의 증가분 비중"),
         (16, "원인: SA 대체 비중", 0.3, PCT, "수기 대체전표(SA) 순증의 증가분 비중"),
         (17, "원인: 키워드 비중", 0.3, PCT, "키워드사전 해당 전표의 증가분 비중"),
         (18, "원인: 단가 상승률", 0.05, PCT, "같은 건수에서 이 비율 이상, 50% 이하로 오르면 '단가 인상'"),
         (19, "고갈 추세 기울기", os.environ.get("BURN_METHOD", "연내누적"), None, "연내누적(연초~기준월 누계에 맞춘 직선) / 최근월(당월 사용액) / 최근분기(최근 3개월 평균) / 전년추세(작년 같은 남은 기간의 월평균)")]
for rr, k, v, fm, d in SROWS:
    put(st, rr, 2, k, bold=True, fill=SUB)
    put(st, rr, 3, v, BLUE, fm, fill=YEL, align="center")
    put(st, rr, 4, d, GRAY)
for rr, k, f_, fm, d in [(20, "(계산) 기준 기간번호", f"=C6*12+C7", "0", "연도×12+월. 기간 비교에 사용"),
                         (21, "(계산) 기준 분기", f"=ROUNDUP(C7/3,0)", "0", ""),
                         (22, "(계산) 분기 시작 기간번호", f"=C6*12+(C21-1)*3+1", "0", ""),
                         (23, "(계산) 금액 환산계수", "=INDEX($G$5:$G$7,MATCH(C8,$F$5:$F$7,0))/INDEX($G$5:$G$7,MATCH(C9,$F$5:$F$7,0))", "0.000000", "원본 금액 × 계수 = 보고단위")]:
    put(st, rr, 2, k, bold=True, fill=SUB); put(st, rr, 3, f_, fmt=fm, align="center"); put(st, rr, 4, d, GRAY)
header(st, 4, ["단위", "원 환산"], start=6, h=20)
for i, (u, v) in enumerate([("원", 1), ("천원", 1000), ("백만원", 1000000)]):
    put(st, 5 + i, 6, u, BLUE, align="center"); put(st, 5 + i, 7, v, BLUE, "#,##0")
dv = DataValidation(type="list", formula1="=$F$5:$F$7"); st.add_data_validation(dv); dv.add("C8:C9")
dv = DataValidation(type="whole", operator="between", formula1="1", formula2="12"); st.add_data_validation(dv); dv.add("C7")
dv = DataValidation(type="list", formula1='"연내누적,최근월,최근분기,전년추세"'); st.add_data_validation(dv); dv.add("C19")
widths(st, [2, 30, 14, 62, 3, 9, 12])

# =================================================================== 매핑
gm = wb.create_sheet("계정매핑")
title(gm, "계정 매핑", "사업비 여부가 Y인 계정만 분석 · 코드 칸은 텍스트")
header(gm, 4, ["G/L 계정", "계정명", "사업비 항목", "사업비 여부", "IFRS17\n대체구분", "직접대체\nIFRS17 계정", "공통비\n성격"], h=32)
DIRECT = {"63501000": "신계약비", "63502000": "수금비", "62701000": "비귀속", "63601000": "손해조사비"}
NATURE_OTHER = {"62601000", "62801000", "62802000"}
rows = [(g[0], g[1], ITEMS[g[2]], "Y", "직접" if g[0] in DIRECT else "공통비", DIRECT.get(g[0], ""), "" if g[0] in DIRECT else ("기타" if g[0] in NATURE_OTHER else "유지비")) for g in GLS] \
       + [("61101000", "급여", "인건비", "제외", "", "", ""), ("61201000", "복리후생비", "인건비", "제외", "", "", "")]
for i, rw in enumerate(rows):
    for j, v in enumerate(rw):
        c = put(gm, 5 + i, j + 1, v or None, BLUE, align="center" if j in (0, 3, 4, 5, 6) else None)
        if j == 0: c.number_format = "@"
for rng_, lst in (("E5:E504", '"직접,공통비"'), ("F5:F504", '"신계약비,유지비,수금비,손해조사비,비귀속"'), ("G5:G504", '"유지비,기타"')):
    dv = DataValidation(type="list", formula1=lst, allow_blank=True); gm.add_data_validation(dv); dv.add(rng_)
gm.freeze_panes = "A5"; widths(gm, [12, 24, 14, 11, 11, 13, 10])

dm = wb.create_sheet("부서매핑")
title(dm, "부서 매핑", "코스트센터 → 부서 · 코드 칸은 텍스트")
header(dm, 4, ["코스트센터", "코스트센터명", "부서코드", "부서명"])
for i, (cc, ccn, d, _) in enumerate(CCS):
    for j, v in enumerate([cc, ccn, d, DEPTS[d]]):
        c = put(dm, 5 + i, j + 1, v, BLUE, align="center" if j in (0, 2) else None)
        if j in (0, 2): c.number_format = "@"
dm.freeze_panes = "A5"; widths(dm, [12, 20, 10, 14])

# =================================================================== 전표_원본
je = wb.create_sheet("전표_원본")
title(je, "전표 원본 (사업비 계정 전표 라인)", "A~K: 붙여넣기 · M~X: 보조수식 (회색, 수정 금지) · 최근 24개월 권장")
header(je, 4, JE_HDR)
header(je, 4, ["기간번호", "부서코드", "키", "금액\n(보고단위)", "선택구분", "선택거래처", "첫등장행", "당기\n근거점수", "비교\n근거점수", "목록포함", "사업비여부", "G/L텍스트",
               "비귀속", "신계약비", "유지비", "수금비", "신계약\n근거점수", "재산관리비", "대체구분", "직접대체\n계정", "재산관리비\n비율", "재산관리비\n약대"], start=13, fill=GRAYH)
for k, e in enumerate(E):
    r = 5 + k
    vals = [e["y"], e["m"], e["doc"], e["date"], e["typ"], int(e["gl"]), int(e["cc"]), e["amt"], e["text"], vcode[e["vname"]], e["vname"]]
    for j, v in enumerate(vals):
        c = je.cell(row=r, column=j + 1, value=v); c.font = F(BLUE)
        if j == 3: c.number_format = "yyyy-mm-dd"
        if j == 7: c.number_format = "#,##0"
JE_FILL = 4 + len(E)
for r in range(5, JE_FILL + 1):
    fm = [f'=IF($A{r}="","",VALUE($A{r}&"")*12+VALUE($B{r}&""))',
          f'=IF($A{r}="","",IFERROR(INDEX(부서매핑!$C$5:$C$504,MATCH($G{r}&"",부서매핑!$A$5:$A$504,0)),"미매핑"))',
          f'=IF($A{r}="","",N{r}&"|"&$F{r})',
          f'=IF($A{r}="",0,$H{r}*{S["fx"]})',
          f'=IF($A{r}="","",IF(O{r}<>원인분석!$C$9,"",IF(AND(M{r}>=원인분석!$C$10,M{r}<=원인분석!$C$11),"당기",IF(AND(M{r}>=원인분석!$C$12,M{r}<=원인분석!$C$13),"비교",""))))',
          f'=IF(Q{r}="","",$K{r}&"")',
          f'=IF(Q{r}="","",IF(MATCH(R{r},$R$5:$R${JE_END},0)=ROW()-4,ROW(),""))',
          f'=IF(Q{r}="당기",ABS(P{r})+ROW()/1000000000,"")',
          f'=IF(Q{r}="비교",ABS(P{r})+ROW()/1000000000,"")',
          f'=IF($A{r}="","",IF(COUNTIF({VR("E")},O{r})>0,"Y","N"))',
          f'=IF($A{r}="","",IFERROR(INDEX(계정매핑!$D$5:$D$504,MATCH($F{r}&"",계정매핑!$A$5:$A$504,0)),"미매핑"))',
          f'=IF($A{r}="","",$F{r}&"")']
    # Y 비귀속 · Z 신계약비 · AA 유지비 · AB 수금비 = 직접대체면 지정 계정에 전액, 공통비면 (1 − 재산관리비율) × 부서별 비율
    for cat_, col_ in (("비귀속", "F"), ("신계약비", "C"), ("유지비", "D"), ("수금비", "E")):
        hit = f'OR(AF{r}="유지비",AF{r}="손해조사비")' if cat_ == "유지비" else f'AF{r}="{cat_}"'
        fm.append(f'=IF($A{r}="",0,IF(AE{r}="직접",IF({hit},$P{r},0),$P{r}*(1-AG{r})*'
                  f'IFERROR(INDEX(계정대체기준!${col_}$20:${col_}$49,MATCH($N{r},계정대체기준!$A$20:$A$49,0)),{1 if cat_ == "비귀속" else 0})))')
    fm.append(f'=IF(AND($A{r}<>"",M{r}={S["P"]},Z{r}>0),Z{r}+ROW()/1000000000,"")')
    fm.append(f'=IF($A{r}="",0,$P{r}*AG{r})')
    fm.append(f'=IF($A{r}="","",IFERROR(INDEX(계정매핑!$E$5:$E$504,MATCH($F{r}&"",계정매핑!$A$5:$A$504,0))&"","공통비"))')
    fm.append(f'=IF($A{r}="","",IF(AE{r}="직접",IFERROR(INDEX(계정매핑!$F$5:$F$504,MATCH($F{r}&"",계정매핑!$A$5:$A$504,0))&"","비귀속"),""))')
    fm.append(f'=IF($A{r}="",0,IF(AND(AE{r}<>"직접",IFERROR(INDEX(계정매핑!$G$5:$G$504,MATCH($F{r}&"",계정매핑!$A$5:$A$504,0))&"","")="유지비",'
              f'COUNTIF(계정대체기준!$A$6:$A$15,$N{r})>0),IFERROR(INDEX(계정대체기준!$E$6:$E$15,MATCH(VALUE($A{r}&""),계정대체기준!$D$6:$D$15,0)),0),0))')
    fm.append(f'=IF($A{r}="","",IF(AD{r}=0,"",IF(OR(COUNTIF(계정대체기준!$A$77:$A$86,$N{r})>0,ISNUMBER(MATCH($X{r},계정대체기준!$D$77:$D$86,0)),'
              f'ISNUMBER(MATCH($F{r},계정대체기준!$D$77:$D$86,0))),"Y","N")))')
    for j, f_ in enumerate(fm):
        c = je.cell(row=r, column=13 + j, value=f_)
        c.font = F(GREEN if j in (1, 3, 4, 9, 10, 12, 13, 14, 15, 17, 18, 19, 20, 21) else BLACK)
        if j in (3, 12, 13, 14, 15, 17): c.number_format = NUM
        if j == 20: c.number_format = PCT
je.freeze_panes = "A5"; je.auto_filter.ref = "A4:AH4"
widths(je, [8, 5, 12, 11, 7, 10, 9, 13, 40, 9, 13, 2, 8, 8, 14, 10, 8, 12, 8, 9, 9, 8, 8, 10, 9, 9, 9, 9, 9, 9, 8, 9, 9, 9])

# =================================================================== TB_원본
tb = wb.create_sheet("TB_원본")
title(tb, "TB 원본 (월별 시산표)", "A~G: 붙여넣기 · I~K: 보조수식 · 당해와 전년을 함께 넣으세요")
header(tb, 4, TB_HDR)
header(tb, 4, ["기간번호", "G/L텍스트", "당월증감\n(보고단위)"], start=9, fill=GRAYH)
for k, row in enumerate(TB):
    for j, v in enumerate(row):
        c = tb.cell(row=5 + k, column=j + 1, value=v); c.font = F(BLUE)
        if j >= 4: c.number_format = "#,##0"
TB_FILL = 4 + len(TB)
for r in range(5, TB_FILL + 1):
    for j, f_ in enumerate([f'=IF($A{r}="","",VALUE($A{r}&"")*12+VALUE($B{r}&""))', f'=IF($A{r}="","",$C{r}&"")',
                            f'=IF($A{r}="",0,($E{r}-$F{r})*{S["fx"]})']):
        c = tb.cell(row=r, column=9 + j, value=f_); c.font = F(GREEN if j == 2 else BLACK)
        if j == 2: c.number_format = NUM
tb.freeze_panes = "A5"; widths(tb, [8, 5, 11, 22, 15, 13, 15, 2, 8, 10, 11])

# =================================================================== 예산판_원본
bp = wb.create_sheet("예산판_원본")
title(bp, "예산판 원본 (부서 × 계정 × 분기)", "A~J: 붙여넣기 · L~R: 보조수식 · 분기는 1~4, Q1, 1분기 모두 인식")
header(bp, 4, BP_HDR)
header(bp, 4, ["키", "연도", "분기", "가용\n(보고단위)", "실적\n(보고단위)", "목록포함", "G/L텍스트"], start=12, fill=GRAYH)
for k, row in enumerate(BP):
    for j, v in enumerate(row):
        c = bp.cell(row=5 + k, column=j + 1, value=v); c.font = F(BLUE)
        if j >= 6: c.number_format = "#,##0"
BP_FILL = 4 + len(BP)
for r in range(5, BP_FILL + 1):
    fm = [f'=IF($A{r}="","",$B{r}&"|"&$D{r})', f'=IF($A{r}="","",VALUE($A{r}&""))',
          f'=IF($A{r}="","",VALUE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(UPPER($F{r}&""),"Q",""),"분기","")," ","")))',
          f'=IF($A{r}="",0,($G{r}+$H{r})*{S["fx"]})', f'=IF($A{r}="",0,$I{r}*{S["fx"]})',
          f'=IF($A{r}="","",IF(COUNTIF({VR("E")},L{r})>0,"Y","N"))', f'=IF($A{r}="","",$D{r}&"")']
    for j, f_ in enumerate(fm):
        c = bp.cell(row=r, column=12 + j, value=f_); c.font = F(GREEN if j in (3, 4, 5) else BLACK)
        if j in (3, 4): c.number_format = NUM
bp.freeze_panes = "A5"; widths(bp, [8, 9, 12, 11, 20, 6, 14, 12, 14, 14, 2, 15, 6, 5, 11, 11, 8, 10])

# ---- 소진 추세 (선형). 점 (0,0), (1,누계1월) … (M,누계M월)에 최소제곱 직선 C(t)=α+βt → 연간가용에 닿는 t* = (가용−α)/β
def datestr(x):  # 날짜 일련번호 → "2026-11-17" (TEXT 날짜 서식은 지역 설정에 따라 달라 직접 조립)
    return f'YEAR({x})&"-"&TEXT(MONTH({x}),"00")&"-"&TEXT(DAY({x}),"00")'
def burn_formulas(r, blank, cum, avail, ytd, cB, cA, cT, cS, pyr):
    """cum: 누계 1~12월 범위(기준월 이후는 0), avail: 연간가용, ytd: 누계실적, pyr: 작년 같은 남은 기간 월평균.
    기울기는 설정 C19: 연내누적(점 (0,0)…(M,누계) 최소제곱) / 최근월 / 최근분기(최근 3개월, 연내) / 전년추세. 연내누적 외에는 (M, 누계)를 지나는 직선.
    반환 {열번호: (수식, 색, 서식)} (42~47열)"""
    m = S["mon"]; St = f"({m}*({m}+1)/2)"; Stt = f"({m}*({m}+1)*(2*{m}+1)/6)"
    SC = f"SUM({cum})"; StC = f"SUMPRODUCT(월별집계!$AB$5:$AM$5,{cum})"
    cM = f"INDEX({cum},{m})"
    reg = f"(({m}+1)*{StC}-{St}*{SC})/(({m}+1)*{Stt}-{St}^2)"
    last = f"({cM}-INDEX({cum},MAX(1,{m}-1))*({m}>1))"
    q3 = f"(({cM}-INDEX({cum},MAX(1,{m}-3))*({m}>3))/MIN(3,{m}))"
    meth = S["burn"]
    t = f"{cT}{r}"; st = f"{cS}{r}"; d = f'(DATE({S["yr"]},1,1)+ROUND({t}*365/12,0)-1)'
    return {
        42: (f'={blank},IF({meth}="최근월",{last},IF({meth}="최근분기",{q3},IF({meth}="전년추세",{pyr},{reg}))))', GREEN, NUM),
        43: (f'={blank},IF({meth}="연내누적",({SC}-{cB}{r}*{St})/({m}+1),{cM}-{cB}{r}*{m}))', GREEN, NUM),
        44: (f'={blank},{cA}{r}+{cB}{r}*12)', BLACK, NUM),
        45: (f'={blank},IF(OR({avail}<=0,{ytd}>={avail},{cB}{r}<=0),"",MAX({m},({avail}-{cA}{r})/{cB}{r})))', BLACK, "0.0"),
        46: (f'={blank},IF({t}="",IF({st}="이미 소진","소진됨",""),{datestr(d)}))', BLACK, None),
        47: (f'={blank},IF(AND({avail}<=0,{ytd}>0),"무예산",IF({avail}<=0,"-",IF({ytd}>={avail},"이미 소진",IF({cB}{r}<=0,"소진 안 됨",IF({t}<=12,"연내 고갈","연내 여유"))))))', GREEN, None),
    }
def qdepl(lim, act):  # 분기: 분기 시작부터 경과 개월의 평균 속도(직선)로 분기 한도에 닿는 날
    e = f'({S["P"]}-{S["qs"]}+1)'
    tq = f"MAX({e},{lim}/({act}/{e}))"
    d0 = f'DATE({S["yr"]},{S["qs"]}-{S["yr"]}*12,1)'
    return (f'IF({lim}<=0,"",IF({act}>={lim},"한도 소진",IF({act}<=0,"",IF({lim}/({act}/{e})>3,"분기 내 여유",'
            + datestr(f"({d0}+ROUND({tq}*365/12,0)-1)") + '))))')

# =================================================================== 변동분석
va = wb.create_sheet("변동분석")
title(va, "변동분석 (부서 × 계정)", "A·C열만 입력 · 전월/전분기/전년동월 비교 → 예산 점검 → 주의 점수 · 금액: 보고단위")
va["A3"] = f'="기준: "&{plabel(S["P"])}&" · 단위: "&{S["ur"]}'; va["A3"].font = F(GREEN, size=9)
cols = ["부서코드", "부서명", "G/L 계정", "계정명", "키",
        "당월", "전월", "전월\n대비", "전년\n동월", "전년동월\n대비", "당분기", "전분기", "전분기\n대비", "주요\n증가액", "이상치\n(z)",
        "연간가용\n(배정+조정)", "누계실적", "잔여", "남은기간\n예상 A\n(최근추세)", "남은기간\n예상 B\n(전년패턴)", "예상\n연간실적", "예상\n소진율", "예산판정", "부족\n예상액", "소진\n예상월",
        "빠른 신호 (당월)", "주의\n점수"]
header(va, 6, cols, h=44)
header(va, 6, ["정렬키", "단건비중", "SA순증", "음수건수", "증가액\n(30)", "증가율\n(20)", "이상치\n(25)", "예산\n(25)"], start=28, fill=GRAYH, h=44)
header(va, 6, ["분기 한도\n(배정+조정)", "분기 실적", "분기 잔여", "분기말\n예상", "분기\n소진율", "분기 판정"], start=36, h=44)
va.cell(row=5, column=36, value="④-2 분기 한도 점검 (예산 통제 단위: 부서 × 계정 × 분기)").font = F(bold=True)
header(va, 6, ["월 소진속도\n(기울기)", "절편", "추세\n연말 누계", "고갈\n예상(개월)", "고갈\n예상일", "추세 판정", "분기 고갈\n예상일"], start=42, h=44)
header(va, 6, ["고갈\n정렬키"], start=49, fill=GRAYH, h=44)
header(va, 6, ["전년동기\n누계", "전년동기\n누계 대비", "전년 남은기간\n월평균"], start=50, h=44)
va.cell(row=5, column=50, value="③-2 누계 비교").font = F(bold=True)
va.cell(row=5, column=42, value="④-3 소진 추세 (선형: 설정 C19의 기울기로 누계실적 직선을 그어 연간가용에 닿는 시점)").font = F(bold=True)
for c, t_ in [(6, "③ 변동"), (16, "④ 예산 점검"), (26, "⑤ 주의")]:
    va.cell(row=5, column=c, value=t_).font = F(bold=True)
VA_PRE = 7 + len(KEYS) + 40
MRow = lambda col, r: f"월별집계!${col}{r}"
for idx in range(VA_PRE - 7):
    r = 7 + idx
    d, gl = KEYS[idx] if idx < len(KEYS) else (None, None)
    put(va, r, 1, d, BLUE, align="center")
    c = put(va, r, 3, gl, BLUE, align="center"); c.number_format = "@"
    blank = f'IF($A{r}="",""'
    rate = lambda a, b: f'={blank},IF({b}<=0,IF({a}>0,1,0),{a}/{b}-1))'
    hdr = "월별집계!$C$5:$Z$5"; row_ = f"월별집계!$C{r}:$Z{r}"
    fm = {
        2: (f'={blank},IFERROR(INDEX(부서매핑!$D$5:$D$504,MATCH($A{r},부서매핑!$C$5:$C$504,0)),"부서매핑 없음"))', GREEN, None),
        4: (f'={blank},IFERROR(INDEX(계정매핑!$B$5:$B$504,MATCH($C{r}&"",계정매핑!$A$5:$A$504,0)),"계정매핑 없음"))', GREEN, None),
        5: (f'={blank},$A{r}&"|"&$C{r})', BLACK, None),
        6: (f'={blank},{MRow("Z", r)})', GREEN, NUM),
        7: (f'={blank},{MRow("Y", r)})', GREEN, NUM),
        8: (rate(f"F{r}", f"G{r}"), BLACK, RATE),
        9: (f'={blank},{MRow("N", r)})', GREEN, NUM),
        10: (rate(f"F{r}", f"I{r}"), BLACK, RATE),
        11: (f'={blank},SUMIFS({row_},{hdr},">="&{S["qs"]},{hdr},"<="&{S["P"]}))', GREEN, NUM),
        12: (f'={blank},SUMIFS({row_},{hdr},">="&({S["qs"]}-3),{hdr},"<="&({S["P"]}-3)))', GREEN, NUM),
        13: (rate(f"K{r}", f"L{r}"), BLACK, RATE),
        14: (f'={blank},MAX(F{r}-G{r},F{r}-I{r}))', BLACK, NUM),
        15: (f'={blank},IFERROR(({MRow("Z", r)}-AVERAGE(월별집계!N{r}:Y{r}))/STDEV(월별집계!N{r}:Y{r}),0))', GREEN, "0.0"),
        16: (f'={blank},SUMIFS({BR("O")},{BR("L")},$E{r},{BR("M")},{S["yr"]}))', GREEN, NUM),
        17: (f'={blank},SUMIFS({row_},{hdr},">="&({S["yr"]}*12+1),{hdr},"<="&{S["P"]}))', GREEN, NUM),
        18: (f'={blank},P{r}-Q{r})', BLACK, NUM),
        19: (f'={blank},AVERAGE(월별집계!X{r}:Z{r})*(12-{S["mon"]}))', GREEN, NUM),
        20: (f'={blank},IFERROR(Q{r}/(SUMIFS({row_},{hdr},">="&(({S["yr"]}-1)*12+1),{hdr},"<="&({S["P"]}-12))'
             f'/SUMIFS({row_},{hdr},">="&(({S["yr"]}-1)*12+1),{hdr},"<="&(({S["yr"]}-1)*12+12)))-Q{r},S{r}))', GREEN, NUM),
        21: (f'={blank},Q{r}+MAX(S{r},T{r}))', BLACK, NUM),
        22: (f'={blank},IF(P{r}=0,0,U{r}/P{r}))', BLACK, PCT),
        23: (f'={blank},IF(AND(P{r}=0,Q{r}>0),"무예산",IF(P{r}=0,"-",IF(V{r}>{S["short"]},"부족 예상",IF(V{r}>{S["warn"]},"주의","여유")))))', GREEN, None),
        24: (f'={blank},IF(P{r}=0,0,MAX(0,U{r}-P{r})))', BLACK, NUM),
        25: (f'={blank},IF(OR(W{r}<>"부족 예상",{S["mon"]}=12),"",IF(R{r}<=0,"이미 초과",MIN(12,{S["mon"]}+ROUNDUP(R{r}/(MAX(S{r},T{r})/(12-{S["mon"]})),0))&"월")))', GREEN, None),
        26: (f'={blank},TRIM(IF(AND(N(AC{r})>={S["one"]},N{r}>={S["amt"]}),"단건"&TEXT(MIN(AC{r},1),"0%")&" ","")&IF(AND(N{r}>0,AD{r}>0),IF(AD{r}/N{r}>={S["sa"]},"SA대체 ",""),"")'
             f'&IF(AE{r}>0,"역분개"&AE{r}&"건 ","")&IF(O{r}>={S["z"]},"이상치 ","")&IF(AND(I{r}=0,F{r}>0),"신규지출 ","")))', GREEN, None),
        27: (f'={blank},ROUND((AF{r}+AG{r}+AH{r}+AI{r})*IF(AND(F{r}<{S["amt"]},N{r}<{S["amt"]}),0.3,1),1))', GREEN, "0.0"),
        28: (f'={blank},AA{r}+ROW()/1000000)', BLACK, "0.000"),
        29: (f'=IF(OR($A{r}="",N{r}<=0),"",SUMPRODUCT(MAX(({JR("O")}=$E{r})*({JR("M")}={S["P"]})*{JR("P")}))/N{r})', GREEN, PCT),
        30: (f'={blank},SUMIFS({JR("P")},{JR("O")},$E{r},{JR("M")},{S["P"]},{JR("E")},"SA")-SUMIFS({JR("P")},{JR("O")},$E{r},{JR("M")},{S["P"]}-1,{JR("E")},"SA"))', GREEN, NUM),
        31: (f'={blank},COUNTIFS({JR("O")},$E{r},{JR("M")},{S["P"]},{JR("P")},"<0"))', GREEN, "0"),
        32: (f'={blank},MIN(MAX(0,N{r})/({S["amt"]}*5),1)*30)', GREEN, "0.0"),
        33: (f'={blank},MIN(MAX(0,H{r},J{r}),1)*20)', BLACK, "0.0"),
        34: (f'={blank},MIN(MAX(0,O{r})/3,1)*25)', BLACK, "0.0"),
        35: (f'={blank},IF(OR(W{r}="부족 예상",W{r}="무예산",AO{r}="한도 초과",AO{r}="초과 예상",AO{r}="무예산"),25,IF(OR(W{r}="주의",AO{r}="주의"),15,0)))', BLACK, "0.0"),
        36: (f'={blank},SUMIFS({BR("O")},{BR("L")},$E{r},{BR("M")},{S["yr"]},{BR("N")},{S["Q"]}))', GREEN, NUM),
        37: (f'={blank},K{r})', BLACK, NUM),
        38: (f'={blank},AJ{r}-AK{r})', BLACK, NUM),
        39: (f'={blank},AK{r}+AVERAGE(월별집계!X{r}:Z{r})*MAX(0,{S["qs"]}+2-{S["P"]}))', GREEN, NUM),
        40: (f'={blank},IF(AJ{r}=0,0,AM{r}/AJ{r}))', BLACK, PCT),
        41: (f'={blank},IF(AND(AJ{r}=0,AK{r}>0),"무예산",IF(AJ{r}=0,"-",IF(AK{r}>AJ{r},"한도 초과",IF(AN{r}>{S["short"]},"초과 예상",IF(AN{r}>{S["warn"]},"주의","여유"))))))', GREEN, None),
    }
    fm[52] = (f'={blank},SUMIFS({row_},{hdr},">"&({S["P"]}-12),{hdr},"<="&(({S["yr"]}-1)*12+12))/MAX(1,12-{S["mon"]}))', GREEN, NUM)
    fm.update(burn_formulas(r, blank, f"월별집계!$AB{r}:$AM{r}", f"P{r}", f"Q{r}", "AP", "AQ", "AS", "AU", f"AZ{r}"))
    fm[48] = (f'={blank},{qdepl(f"AJ{r}", f"AK{r}")})', GREEN, None)
    fm[50] = (f'={blank},SUMIFS({row_},{hdr},">="&(({S["yr"]}-1)*12+1),{hdr},"<="&({S["P"]}-12)))', GREEN, NUM)
    fm[51] = (rate(f"Q{r}", f"AX{r}"), BLACK, RATE)
    fm[49] = (f'=IF($A{r}="",999+ROW()/1000000,IF(AU{r}="이미 소진",ROW()/1000000,IF(AU{r}="연내 고갈",AS{r}+ROW()/1000000,99+ROW()/1000000)))', BLACK, "0.000")
    for c, (f_, col, fmt) in fm.items():
        put(va, r, c, f_, col, fmt, align="center" if c in (5, 23, 25, 41, 46, 47, 48) else None)
va.freeze_panes = "F7"; va.auto_filter.ref = f"A6:AZ{VA_PRE-1}"
widths(va, [8, 12, 10, 18, 8] + [9] * 10 + [11, 10, 10, 10, 10, 10, 8, 10, 9, 9, 20, 7] + [8] * 8 + [11, 10, 10, 10, 8, 10] + [10, 9, 10, 8, 11, 10, 12, 8] + [10, 9, 11])
for val, fill in (('"이미 소진"', "F8CBAD"), ('"연내 고갈"', "FCE4D6"), ('"무예산"', "D9C3E9")):
    va.conditional_formatting.add(f"AU7:AU{VA_END}", CellIsRule(operator="equal", formula=[val], fill=PatternFill("solid", fgColor=fill)))
for val, fill in (('"한도 초과"', "F8CBAD"), ('"초과 예상"', "FCE4D6"), ('"주의"', "FFE699"), ('"무예산"', "D9C3E9")):
    va.conditional_formatting.add(f"AO7:AO{VA_END}", CellIsRule(operator="equal", formula=[val], fill=PatternFill("solid", fgColor=fill)))
va.column_dimensions["E"].hidden = True
for rng_, rules in [(f"W7:W{VA_END}", [('"부족 예상"', "F8CBAD", "9C0006"), ('"주의"', "FFE699", None), ('"무예산"', "D9C3E9", None)])]:
    for val, fill, fc in rules:
        va.conditional_formatting.add(rng_, CellIsRule(operator="equal", formula=[val], fill=PatternFill("solid", fgColor=fill), font=Font(color=fc, bold=True) if fc else None))

# =================================================================== 월별집계
mz = wb.create_sheet("월별집계")
title(mz, "월별 집계 (부서 × 계정 × 최근 24개월)", "전표 기준 · 행 순서는 변동분석과 같음 · 금액: 보고단위")
put(mz, 4, 2, "기간", bold=True)
put(mz, 5, 1, "키", bold=True, fill=SUB); put(mz, 5, 2, "기간번호", bold=True, fill=SUB)
for k in range(24):
    c = 3 + k
    put(mz, 5, c, f"={S['P']}-{23-k}", GRAY, "0", align="center", fill=SUB)
    put(mz, 4, c, "=" + plabel(f"{CL(c)}$5"), bold=True, align="center", fill=SUB)
put(mz, 6, 1, "전사", bold=True, fill=TOT); put(mz, 6, 2, "사업비 합계", bold=True, fill=TOT)
for k in range(24):
    col = CL(3 + k)
    put(mz, 6, 3 + k, f"=SUM({col}7:{col}{VA_END})", fmt=NUM, bold=True, fill=TOT)
for idx in range(VA_PRE - 7):
    r = 7 + idx
    put(mz, r, 1, f'=IF(변동분석!$A{r}="","",변동분석!$E{r})', GREEN)
    put(mz, r, 2, f'=IF($A{r}="","",변동분석!$B{r}&" "&변동분석!$D{r})', GREEN)
    for k in range(24):
        col = CL(3 + k)
        put(mz, r, 3 + k, f'=IF($A{r}="","",SUMIFS({JR("P")},{JR("O")},$A{r},{JR("M")},{col}$5))', GREEN, NUM)
# 보조: 기준연도 누계 1~12월 (기준월 이후는 0) → 변동분석 AP~AU 소진 추세 회귀에 사용
put(mz, 4, 28, "당해 누계 (소진 추세용 · 기준월 이후 0)", bold=True)
for k in range(1, 13):
    c = 27 + k
    put(mz, 5, c, k, GRAY, "0", align="center", fill=SUB)
    put(mz, 6, c, f"=SUM({CL(c)}7:{CL(c)}{VA_END})", fmt=NUM, bold=True, fill=TOT)
for idx in range(VA_PRE - 7):
    r = 7 + idx
    for k in range(1, 13):
        put(mz, r, 27 + k, f'=IF(OR($A{r}="",{k}>{S["mon"]}),0,SUMIFS($C{r}:$Z{r},$C$5:$Z$5,">="&({S["yr"]}*12+1),$C$5:$Z$5,"<="&({S["yr"]}*12+{k})))', BLACK, NUM)
mz.freeze_panes = "C7"; widths(mz, [14, 26] + [8.5] * 24 + [2] + [8.5] * 12)

# =================================================================== 계정변동 (TB 기준)
av = wb.create_sheet("계정변동")
title(av, "계정변동 (TB 기준, 전사) 및 3자 대사", "TB로 본 계정별 변동 + 전표·예산판과 합계 대사 · 금액: 보고단위")
av["A3"] = f'="기준: "&{plabel(S["P"])}&" · 당분기: "&{S["yr"]}&"년 "&{S["Q"]}&"분기"'; av["A3"].font = F(GREEN, size=9)
header(av, 5, ["G/L 계정", "계정명", "TB 당월", "TB 전월", "전월\n대비", "TB 전년동월", "전년동월\n대비", "TB 당분기", "TB 전분기", "전분기\n대비",
               "전표 당월", "대사 차이\n(TB−전표)", "예산판\n당분기 실적", "전표\n당분기", "대사 차이\n(예산판−전표)", "판정"], h=40)
sgl = [g[0] for g in GLS]
for i, gl in enumerate(sgl):
    r = 6 + i
    c = put(av, r, 1, gl, BLUE, align="center"); c.number_format = "@"
    put(av, r, 2, f'=IFERROR(INDEX(계정매핑!$B$5:$B$504,MATCH($A{r},계정매핑!$A$5:$A$504,0)),"")', GREEN)
    tbs = lambda crit: f'SUMIFS({TR("K")},{TR("J")},$A{r},{crit})'
    put(av, r, 3, "=" + tbs(f'{TR("I")},{S["P"]}'), GREEN, NUM)
    put(av, r, 4, "=" + tbs(f'{TR("I")},{S["P"]}-1'), GREEN, NUM)
    put(av, r, 5, f"=IF(D{r}<=0,IF(C{r}>0,1,0),C{r}/D{r}-1)", fmt=RATE)
    put(av, r, 6, "=" + tbs(f'{TR("I")},{S["P"]}-12'), GREEN, NUM)
    put(av, r, 7, f"=IF(F{r}<=0,IF(C{r}>0,1,0),C{r}/F{r}-1)", fmt=RATE)
    put(av, r, 8, "=" + tbs(f'{TR("I")},">="&{S["qs"]},{TR("I")},"<="&{S["P"]}'), GREEN, NUM)
    put(av, r, 9, "=" + tbs(f'{TR("I")},">="&({S["qs"]}-3),{TR("I")},"<="&({S["P"]}-3)'), GREEN, NUM)
    put(av, r, 10, f"=IF(I{r}<=0,IF(H{r}>0,1,0),H{r}/I{r}-1)", fmt=RATE)
    put(av, r, 11, f'=SUMIFS({JR("P")},{JR("X")},$A{r},{JR("M")},{S["P"]})', GREEN, NUM)
    put(av, r, 12, f"=C{r}-K{r}", fmt=NUM)
    put(av, r, 13, f'=SUMIFS({BR("P")},{BR("R")},$A{r},{BR("M")},{S["yr"]},{BR("N")},{S["Q"]})', GREEN, NUM)
    put(av, r, 14, f'=SUMIFS({JR("P")},{JR("X")},$A{r},{JR("M")},">="&{S["qs"]},{JR("M")},"<="&{S["P"]})', GREEN, NUM)
    put(av, r, 15, f"=M{r}-N{r}", fmt=NUM)
    put(av, r, 16, f'=IF(AND(ABS(L{r})<0.05,ABS(O{r})<0.05),"OK","차이 확인")', align="center")
AT = 6 + len(sgl)
put(av, AT, 1, "합계", bold=True, fill=TOT); put(av, AT, 2, None, fill=TOT)
for c in (3, 4, 6, 8, 9, 11, 12, 13, 14, 15):
    put(av, AT, c, f"=SUM({CL(c)}6:{CL(c)}{AT-1})", fmt=NUM, bold=True, fill=TOT)
for c, (a, b) in {5: ("C", "D"), 7: ("C", "F"), 10: ("H", "I")}.items():
    put(av, AT, c, f"=IF({b}{AT}<=0,0,{a}{AT}/{b}{AT}-1)", fmt=RATE, bold=True, fill=TOT)
put(av, AT, 16, f'=IF(COUNTIF(P6:P{AT-1},"OK")={len(sgl)},"전체 일치","차이 있음")', bold=True, fill=TOT, align="center")
av.freeze_panes = "C6"; widths(av, [11, 20] + [10] * 13 + [10])
av.conditional_formatting.add(f"P6:P{AT}", CellIsRule(operator="equal", formula=['"차이 확인"'], fill=PatternFill("solid", fgColor="FFC7CE")))

# =================================================================== 코멘트
cm = wb.create_sheet("코멘트")
title(cm, "확인된 원인·조치", "담당 부서 확인 결과를 월별로 누적 · 기준연도·기준월이 같은 코멘트가 주의대상·리포트에 표시")
header(cm, 4, ["연도", "월", "부서코드", "G/L 계정", "원인 유형", "확인된 원인·조치", "작성자", "작성일"])
header(cm, 4, ["키"], start=9, fill=GRAYH)
CMTS = [("D500", "62701000", "일회성 대형", "3분기 신상품 TV캠페인 집행을 앞당김. 4분기 집행분에서 조정 예정", "마케팅팀 김OO"),
        ("D200", "62901000", "재분류", "8월 경영기획팀에 잘못 귀속된 회계자문료를 재분류. 실제 증가 아님", "재무회계팀 이OO")]
for k, (d, gl, t_, txt, who) in enumerate(CMTS):
    r = 5 + k
    for j, v in enumerate([Y, M, d, gl, t_, txt, who, dt.date(2026, 10, 8)]):
        c = put(cm, r, j + 1, v, BLUE, "yyyy-mm-dd" if j == 7 else None, align="center" if j < 5 else None)
        if j == 3: c.number_format = "@"
dv = DataValidation(type="list", formula1='"일회성 대형,신규 거래처,단가 인상,계상 시점 이동,정정·역분개,재분류,연간 일괄 계상,물량 증가,기타"', allow_blank=True)
cm.add_data_validation(dv); dv.add("E5:E2004")
for r in range(5, 2005):
    cm.cell(row=r, column=9, value=f'=IF($A{r}="","",$A{r}&"|"&$B{r}&"|"&$C{r}&"|"&$D{r})').font = F(GRAY)
cm.freeze_panes = "A5"; widths(cm, [7, 5, 9, 11, 14, 60, 16, 11, 22])

# =================================================================== 키워드사전
kw = wb.create_sheet("키워드사전")
title(kw, "적요 키워드 사전", "원인분석의 '키워드' 패턴이 사용 · 확인된 표현을 계속 추가하세요 (빈 칸은 무시)")
kw["A3"] = '="현재 원인분석 선택: "&원인분석!$E$9'; kw["A3"].font = F(GREEN, size=9)
header(kw, 4, ["원인 유형", "키워드1", "키워드2", "키워드3", "키워드4", "키워드5", "원인 설명 (출력 문장에 사용)", "당기 해당 금액", "증가분 대비"])
KWS = [("연간 일괄 계상", ["연간", "일시", "선급", "1년", "반기"], "연간·반기 계약을 한 번에 계상했을 가능성. 월할 안분 대상인지 확인"),
       ("캠페인·행사", ["캠페인", "행사", "이벤트", "프로모션", "박람회"], "일회성 캠페인·행사 비용일 가능성"),
       ("정정·취소", ["정정", "취소", "역분개", "오류", "환입"], "정정 전표 영향. 순액 기준으로 다시 확인"),
       ("이관·재분류", ["재분류", "이관", "대체", "귀속변경", ""], "다른 부서·계정에서 넘어온 금액일 가능성"),
       ("추정 계상", ["미지급", "추정", "선계상", "", ""], "월말 추정 계상분. 다음 달 역분개 여부 확인"),
       ("소급·정산", ["소급", "정산", "추가청구", "차액", ""], "과거분 정산·소급 지급일 가능성"),
       ("판매시책·추가수수료", ["시책", "추가 수수료", "인센티브", "특별수당", ""], "판매 시책비·추가 수수료. 신계약비 경험차이로 GMM CSM에 영향 가능")]
for i in range(15):
    r = 5 + i
    if i < len(KWS):
        nm, ks, desc = KWS[i]
        put(kw, r, 1, nm, BLUE, bold=True)
        for j, k_ in enumerate(ks): put(kw, r, 2 + j, k_ or None, BLUE)
        put(kw, r, 7, desc, BLUE)
    else:
        for c in range(1, 8): put(kw, r, c, None, BLUE)
    put(kw, r, 8, f'=IF($A{r}="","",SUMPRODUCT(MAX(($B{r}:$F{r}<>"")*SUMIFS({JR("P")},{JR("Q")},"당기",{JR("I")},"*"&$B{r}:$F{r}&"*"))))', GREEN, NUM)
    put(kw, r, 9, f'=IF(OR($A{r}="",원인분석!$C$16<=0),"",H{r}/원인분석!$C$16)', GREEN, PCT)
kw.freeze_panes = "A5"; widths(kw, [16, 9, 9, 9, 9, 9, 52, 12, 10])

# =================================================================== 원인분석
ca = wb.create_sheet("원인분석")
title(ca, "원인분석 (한 항목씩)", "C5~C8만 입력 · 선택 항목의 증가분을 거래처·전표로 쪼개고 8가지 패턴으로 원인 후보를 만듭니다 · 금액: 보고단위")
sel = [(5, "주의대상 순위", 1, "1~10"), (6, "비교 기준", "전월", "전월 / 전분기 / 전년동기누적 / 전년동월"),
       (7, "직접 지정: 부서코드", None, "비워두면 순위 사용"), (8, "직접 지정: G/L 계정", None, "부서코드와 함께 입력")]
if os.environ.get("SEL"):
    _d, _g, _b = os.environ["SEL"].split(",")
    sel = [(5, "주의대상 순위", 1, "1~10"), (6, "비교 기준", _b, ""), (7, "직접 지정: 부서코드", _d, ""), (8, "직접 지정: G/L 계정", _g, "")]
for rr, k, v, d in sel:
    put(ca, rr, 2, k, bold=True, fill=SUB); c = put(ca, rr, 3, v, BLUE, fill=YEL, align="center"); put(ca, rr, 4, d, GRAY)
    if rr == 8: c.number_format = "@"
dv = DataValidation(type="list", formula1='"전월,전분기,전년동기누적,전년동월"'); ca.add_data_validation(dv); dv.add("C6")
dv = DataValidation(type="whole", operator="between", formula1="1", formula2="10"); ca.add_data_validation(dv); dv.add("C5")
calc = [(9, "(계산) 선택 키", f'=IF(AND(C7<>"",C8<>""),C7&"|"&C8,IFERROR(INDEX(주의대상!$L$6:$L$15,C5),""))'),
        (10, "(계산) 당기 시작", f'=IF(C6="전분기",{S["qs"]},IF(C6="전년동기누적",{S["yr"]}*12+1,{S["P"]}))'), (11, "(계산) 당기 끝", f'={S["P"]}'),
        (12, "(계산) 비교 시작", f'=IF(C6="전월",{S["P"]}-1,IF(C6="전년동월",{S["P"]}-12,IF(C6="전년동기누적",({S["yr"]}-1)*12+1,{S["qs"]}-3)))'),
        (13, "(계산) 비교 끝", f'=IF(C6="전월",{S["P"]}-1,IF(C6="전년동월",{S["P"]}-12,IF(C6="전년동기누적",{S["P"]}-12,{S["P"]}-3)))')]
for rr, k, f_ in calc:
    put(ca, rr, 2, k, GRAY, fill=TOT); put(ca, rr, 3, f_, GREEN if rr != 11 else GREEN, "0" if rr > 9 else None, align="center", fill=TOT)
# 선택 항목 이름 · 기간 라벨
ca["E5"] = "선택 항목"; ca["E5"].font = F(GRAY, size=9)
ca["E6"] = ('=IF($C$9="","(선택 없음)",IFERROR(INDEX(부서매핑!$D$5:$D$504,MATCH(LEFT($C$9,FIND("|",$C$9)-1),부서매핑!$C$5:$C$504,0)),LEFT($C$9,FIND("|",$C$9)-1))'
            '&" · "&IFERROR(INDEX(계정매핑!$B$5:$B$504,MATCH(MID($C$9,FIND("|",$C$9)+1,20),계정매핑!$A$5:$A$504,0)),MID($C$9,FIND("|",$C$9)+1,20)))')
ca["E6"].font = F(GREEN, bold=True, size=13)
ca["E7"] = f'="당기 "&{plabel("$C$10")}&IF($C$10<>$C$11," ~ "&{plabel("$C$11")},"")&"  vs  비교 "&{plabel("$C$12")}&IF($C$12<>$C$13," ~ "&{plabel("$C$13")},"")'
ca["E7"].font = F(GREEN)
ca["E9"] = '=E6&" ("&C6&" 대비)"'; ca["E9"].font = F(GREEN, size=9)

header(ca, 15, ["당기 합계", "비교 합계", "증감", "증감률", "분기 판정", "연간 부족 예상액\n(추세 연말)", "고갈 예상일\n(추세)"], h=30)
put(ca, 16, 1, f'=SUMIFS({JR("P")},{JR("Q")},"당기")', GREEN, NUM, bold=True)
put(ca, 16, 2, f'=SUMIFS({JR("P")},{JR("Q")},"비교")', GREEN, NUM, bold=True)
put(ca, 16, 3, "=A16-B16", fmt='+#,##0.0;-#,##0.0;0', bold=True)
put(ca, 16, 4, "=IF(B16<=0,IF(A16>0,1,0),A16/B16-1)", fmt=RATE, bold=True)
put(ca, 16, 5, f'=IFERROR(INDEX({VR("AO")},MATCH($C$9,{VR("E")},0)),"-")', GREEN, bold=True, align="center")
put(ca, 16, 6, f'=IFERROR(IF(INDEX({VR("P")},MATCH($C$9,{VR("E")},0))<=0,0,MAX(0,INDEX({VR("AR")},MATCH($C$9,{VR("E")},0))-INDEX({VR("P")},MATCH($C$9,{VR("E")},0)))),0)', GREEN, NUM, bold=True)
put(ca, 16, 7, f'=IFERROR(INDEX({VR("AT")},MATCH($C$9,{VR("E")},0)),"")', GREEN, bold=True, align="center")
ca.merge_cells("A18:I18")
ca["A18"] = ('=IF($C$9="","선택된 항목이 없습니다.",$E$6&" "&$C$6&" 대비 "&TEXT($C$16,"+#,##0.0;-#,##0.0;0")&" ("&TEXT($D$16,"+0%;-0%;0%")&"). "'
             '&IF($C$16<=0,"감소 또는 변동 없음 - 원인 후보는 증가 항목에만 만듭니다. ",$D$21&". ")'
             '&"분기 한도: "&$E$16&IF($G$16="소진됨",", 연간 예산 이미 소진",IF(N($F$16)>0,", 연간 "&TEXT($F$16,"#,##0.0")&" 부족 예상 ("&$G$16&" 고갈)",""))&".")')
ca["A18"].font = F(bold=True, size=11); ca["A18"].alignment = Alignment(wrap_text=True, vertical="top"); ca["A18"].fill = BOXF
ca.row_dimensions[18].height = 48

# --- 패턴 판정 (보조 영역 N~T, 20~27행 / 보조값 Y~Z)
sec(ca, 19, "원인 후보 (근거 강도 순)")
header(ca, 20, ["순위", "패턴", "근거 강도", "설명 (근거 포함)"], start=1, h=20)
header(ca, 19, ["패턴", "해당", "강도", "강도점수", "비중", "문장", "정렬키"], start=14, fill=GRAYH, h=20)
header(ca, 19, ["보조값", "값"], start=25, fill=GRAYH, h=20)
aux = {  # Z열 보조값 (행: 수식)
    20: ("최대 단건 위치", f'=IFERROR(MATCH(LARGE({JR("T")},1),{JR("T")},0),"")'),
    21: ("최대 단건 금액", '=IF(Z20="","",INDEX(' + JR("P") + ',Z20))'),
    22: ("전표번호", '=IF(Z20="","",INDEX(' + JR("C") + ',Z20))'),
    23: ("전기일", '=IF(Z20="","",INDEX(' + JR("D") + ',Z20))'),
    24: ("거래처", '=IF(Z20="","",INDEX(' + JR("K") + ',Z20))'),
    25: ("적요", '=IF(Z20="","",INDEX(' + JR("I") + ',Z20))'),
    26: ("월별집계 행", f'=IFERROR(MATCH($C$9,월별집계!$A$7:$A${VA_END},0),"")'),
    27: ("당월 (기준월)", f'=IF(Z26="",0,INDEX(월별집계!$Z$7:$Z${VA_END},Z26))'),
    28: ("전월", f'=IF(Z26="",0,INDEX(월별집계!$Y$7:$Y${VA_END},Z26))'),
    29: ("직전3개월 평균 (전월 제외)", f'=IF(Z26="",0,(INDEX(월별집계!$V$7:$V${VA_END},Z26)+INDEX(월별집계!$W$7:$W${VA_END},Z26)+INDEX(월별집계!$X$7:$X${VA_END},Z26))/3)'),
    30: ("음수 전표 건수", f'=COUNTIFS({JR("Q")},"당기",{JR("P")},"<0")'),
    31: ("음수 전표 합계", f'=SUMIFS({JR("P")},{JR("Q")},"당기",{JR("P")},"<0")'),
    32: ("SA 순증", f'=SUMIFS({JR("P")},{JR("Q")},"당기",{JR("E")},"SA")-SUMIFS({JR("P")},{JR("Q")},"비교",{JR("E")},"SA")'),
    33: ("SA 첫 적요", f'=IFERROR(INDEX({JR("I")},MATCH(1,INDEX(({JR("Q")}="당기")*({JR("E")}="SA"),0),0)),"")'),
    34: ("키워드 최대 비중", '=IF(COUNT(키워드사전!$I$5:$I$19)=0,0,MAX(키워드사전!$I$5:$I$19))'),
    35: ("키워드 유형 행", '=IF(Z34<=0,"",MATCH(Z34,키워드사전!$I$5:$I$19,0))'),
    36: ("당기 건수 (+)", f'=COUNTIFS({JR("Q")},"당기",{JR("P")},">0")'),
    37: ("비교 건수 (+)", f'=COUNTIFS({JR("Q")},"비교",{JR("P")},">0")'),
    38: ("당기 금액 (+)", f'=SUMIFS({JR("P")},{JR("Q")},"당기",{JR("P")},">0")'),
    39: ("비교 금액 (+)", f'=SUMIFS({JR("P")},{JR("Q")},"비교",{JR("P")},">0")'),
    40: ("건당 당기", '=IF(Z36=0,0,Z38/Z36)'),
    41: ("건당 비교", '=IF(Z37=0,0,Z39/Z37)'),
    42: ("신규 거래처 증감", '=SUMIFS($R$32:$R$46,$V$32:$V$46,"Y")'),
    43: ("신규 거래처 수", '=COUNTIFS($V$32:$V$46,"Y",$R$32:$R$46,">0")'),
    44: ("첫 신규 거래처", '=IFERROR(INDEX($A$32:$A$46,MATCH("Y",$I$32:$I$46,0)),"")'),
    45: ("최대 단건 거래처 평소 월평균", f'=IF(Z20="",0,SUMIFS({JR("P")},{JR("O")},$C$9,{JR("K")},Z24,{JR("M")},">="&($C$10-12),{JR("M")},"<"&$C$10)/12)'),
}
for rr, (lab, f_) in aux.items():
    put(ca, rr, 25, lab, GRAY); put(ca, rr, 26, f_, GREEN if "!" in f_ else BLACK, NUM if rr in (21, 27, 28, 29, 31, 32, 38, 39, 40, 41, 42, 45) else ("m/d" if rr == 23 else None))
ca.cell(row=34, column=26).number_format = PCT
T = lambda v: f'TEXT({v},"#,##0.0")'
PAT = [
    ("일회성 대형 지출", "높음", 3,
     f'=IF(OR(Z20="",$C$16<=0),"N",IF(AND(Z21>0,Z21/$C$16>={S["one"]},LEFT(Z24,1)<>"(",Z21>=3*Z45),"Y","N"))',
     '=IF(O20="Y",MIN(1,Z21/$C$16),0)',
     f'="전표 "&Z22&" ("&TEXT(Z23,"m/d")&", "&Z24&", 적요 \'"&Z25&"\') "&{T("Z21")}&"이 증가분의 "&TEXT(R20,"0%")&" (거래처 평소 월평균 "&{T("Z45")}&"). 일회성 지출 가능성"'),
    ("신규 거래처", "중간", 2,
     f'=IF($C$16<=0,"N",IF(Z42/$C$16>={S["new"]},"Y","N"))',
     '=IF(O21="Y",MIN(1,Z42/$C$16),0)',
     '="최근 12개월 거래가 없던 거래처("&Z44&IF(Z43>1," 등 "&Z43&"곳","")&")가 증가분의 "&TEXT(R21,"0%")&". 신규 계약·거래 개시 가능성"'),
    ("단가 인상", "중간", 2,
     f'=IF(OR($C$16<=0,$A$32=""),"N",IF(AND($I$32="N",$F$32=$G$32,$G$32>0,$C$32>0),IF(AND($D$32/$C$32>={S["price"]},$D$32/$C$32<=0.5),"Y","N"),"N"))',
     '=IF(O22="Y",MIN(1,$D$32/$C$16),0)',
     '="같은 건수("&$F$32&"건)로 "&$A$32&" 금액이 "&TEXT(IF($C$32>0,$D$32/$C$32,0),"0%")&" 상승. 단가·계약 조건 변경 가능성 (최근 몇 달 추이 확인)"'),
    ("계상 시점 이동", "중간", 2,
     '=IF(AND(Z29>0,Z28<0.5*Z29,Z27>1.5*Z29),"Y","N")',
     '=IF(O23="Y",0.5,0)',
     f'="전월 "&{T("Z28")}&" (직전 3개월 평균 "&{T("Z29")}&") → 당월 "&{T("Z27")}&". 전월에 계상될 금액이 당월로 넘어왔을 가능성"'),
    ("정정 · 역분개", "높음", 3,
     '=IF(Z30>0,"Y","N")',
     '=IF(O24="Y",MIN(1,ABS(Z31)/MAX(ABS($C$16),0.000001)),0)',
     f'="음수(정정·역분개) 전표 "&Z30&"건, 합계 "&{T("Z31")}&". 정정 전후 순액 기준으로 다시 확인"'),
    ("재분류 (수기 대체)", "높음", 3,
     f'=IF($C$16<=0,"N",IF(Z32/$C$16>={S["sa"]},"Y","N"))',
     '=IF(O25="Y",MIN(1,Z32/$C$16),0)',
     f'="수기 대체전표(SA) 순증 "&{T("Z32")}&"이 증가분의 "&TEXT(R25,"0%")&". 다른 부서·계정에서 이관됐을 가능성 (적요: \'"&Z33&"\')"'),
    ("적요 키워드", "중간", 2,
     f'=IF(OR($C$16<=0,Z35=""),"N",IF(Z34>={S["kw"]},"Y","N"))',
     '=IF(O26="Y",MIN(1,Z34),0)',
     '=IF(Z35="","",INDEX(키워드사전!$A$5:$A$19,Z35)&": 적요에 \'"&INDEX(키워드사전!$B$5:$B$19,Z35)&IF(INDEX(키워드사전!$C$5:$C$19,Z35)<>"","·"&INDEX(키워드사전!$C$5:$C$19,Z35),"")'
     '&IF(INDEX(키워드사전!$D$5:$D$19,Z35)<>"","·"&INDEX(키워드사전!$D$5:$D$19,Z35),"")&"\' 등이 있는 전표 "&TEXT(INDEX(키워드사전!$H$5:$H$19,Z35),"#,##0.0")&" (증가분의 "&TEXT(Z34,"0%")&"). "&INDEX(키워드사전!$G$5:$G$19,Z35))'),
    ("물량 vs 단가", "참고", 1,
     '=IF(AND($C$16>0,Z36>0,Z37>0),"Y","N")',
     '=0',
     '="건수 "&Z37&"→"&Z36&" ("&TEXT(Z36/MAX(Z37,1)-1,"+0%;-0%;0%")&"), 건당 "&TEXT(Z41,"#,##0.0")&"→"&TEXT(Z40,"#,##0.0")&". "&IF(ABS((Z36-Z37)*Z41)>=ABS((Z40-Z41)*Z36),"건수(물량) 증가","건당 금액(단가) 상승")&" 영향이 더 큼"'),
]
for i, (nm, lvl, sc, f_yes, f_share, f_txt) in enumerate(PAT):
    r = 20 + i
    put(ca, r, 14, nm); put(ca, r, 15, f_yes, align="center"); put(ca, r, 16, lvl, align="center"); put(ca, r, 17, sc)
    put(ca, r, 18, f_share, fmt=PCT); put(ca, r, 19, f'=IF(O{r}="Y",{f_txt[1:]},"")', GREEN if "!" in f_txt else BLACK)
    put(ca, r, 20, f'=IF(O{r}="Y",Q{r}*10+R{r}+ROW()/1000000,-1)', fmt="0.000")
for k in range(1, 9):
    r = 20 + k
    put(ca, r, 1, k, align="center")
    put(ca, r, 11, f'=IF(LARGE($T$20:$T$27,{k})<0,"",MATCH(LARGE($T$20:$T$27,{k}),$T$20:$T$27,0))', GRAY)
    put(ca, r, 2, f'=IF($K{r}="",IF({k}=1,"뚜렷한 패턴 없음",""),INDEX($N$20:$N$27,$K{r}))', bold=k == 1)
    put(ca, r, 3, f'=IF($K{r}="","",INDEX($P$20:$P$27,$K{r}))', align="center")
    put(ca, r, 4, f'=IF($K{r}="",IF({k}=1,"근거 전표 목록을 직접 확인하세요",""),INDEX($S$20:$S$27,$K{r}))', wrap=True)
    ca.merge_cells(start_row=r, start_column=4, end_row=r, end_column=9)
    ca.row_dimensions[r].height = 30
for cf in (CellIsRule(operator="equal", formula=['"높음"'], fill=PatternFill("solid", fgColor="C6EFCE")),
           CellIsRule(operator="equal", formula=['"중간"'], fill=PatternFill("solid", fgColor="FFEB9C"))):
    ca.conditional_formatting.add("C21:C28", cf)

# --- 거래처별 분해 (표시 A~I 32~46행, 보조 N~W)
sec(ca, 30, "거래처별 증감 분해 (증가 큰 순)")
header(ca, 31, ["거래처", "당기", "비교", "증감", "증가분\n기여율", "당기\n건수", "비교\n건수", "당기\n최대 단건", "신규\n거래처"], h=32)
header(ca, 31, ["첫등장행", "거래처", "당기", "비교", "증감", "당기건수", "비교건수", "최대단건", "신규", "정렬키"], start=14, fill=GRAYH, h=32)
for k in range(1, 16):
    r = 31 + k
    put(ca, r, 14, f'=IFERROR(SMALL({JR("S")},{k}),"")', GREEN)
    put(ca, r, 15, f'=IF(N{r}="","",INDEX(전표_원본!$K$1:$K${JE_END},N{r}))', GREEN)
    put(ca, r, 16, f'=IF(O{r}="","",SUMIFS({JR("P")},{JR("Q")},"당기",{JR("R")},O{r}))', GREEN, NUM)
    put(ca, r, 17, f'=IF(O{r}="","",SUMIFS({JR("P")},{JR("Q")},"비교",{JR("R")},O{r}))', GREEN, NUM)
    put(ca, r, 18, f'=IF(O{r}="","",P{r}-Q{r})', fmt=NUM)
    put(ca, r, 19, f'=IF(O{r}="","",COUNTIFS({JR("Q")},"당기",{JR("R")},O{r},{JR("P")},">0"))', GREEN)
    put(ca, r, 20, f'=IF(O{r}="","",COUNTIFS({JR("Q")},"비교",{JR("R")},O{r},{JR("P")},">0"))', GREEN)
    put(ca, r, 21, f'=IF(O{r}="","",SUMPRODUCT(MAX(({JR("Q")}="당기")*({JR("R")}=O{r})*{JR("P")})))', GREEN, NUM)
    put(ca, r, 22, f'=IF(O{r}="","",IF(LEFT(O{r},1)="(","N",IF(COUNTIFS({JR("O")},$C$9,{JR("K")},O{r},{JR("M")},">="&($C$10-12),{JR("M")},"<"&$C$10)=0,"Y","N")))', GREEN, None)
    put(ca, r, 23, f'=IF(O{r}="","",R{r}+ROW()/1000000)', fmt="0.000")
    put(ca, r, 11, f'=IF(COUNT($W$32:$W$46)<{k},"",MATCH(LARGE($W$32:$W$46,{k}),$W$32:$W$46,0))', GRAY)
    for c, src, fmt in [(1, "O", None), (2, "P", NUM), (3, "Q", NUM), (4, "R", '+#,##0.0;-#,##0.0;0'), (6, "S", "0"), (7, "T", "0"), (8, "U", NUM), (9, "V", None)]:
        put(ca, r, c, f'=IF($K{r}="","",INDEX(${src}$32:${src}$46,$K{r}))', fmt=fmt, align="center" if c == 9 else None)
    put(ca, r, 5, f'=IF(OR($K{r}="",$C$16<=0),"",D{r}/$C$16)', fmt=PCT)
ca.conditional_formatting.add("I32:I46", CellIsRule(operator="equal", formula=['"Y"'], fill=PatternFill("solid", fgColor="FFEB9C"), font=Font(bold=True)))

# --- 근거 전표
sec(ca, 49, "근거 전표 · 당기 (금액 절댓값 큰 순)")
EVH = ["전기일", "전표번호", "유형", "거래처", "적요", "금액", "코스트센터"]
header(ca, 50, EVH, h=20)
def evidence(r0, n, col):
    for k in range(1, n + 1):
        r = r0 + k
        put(ca, r, 11, f'=IFERROR(MATCH(LARGE({JR(col)},{k}),{JR(col)},0),"")', GRAY)
        for c, src, fmt in [(1, "D", "yyyy-mm-dd"), (2, "C", None), (3, "E", None), (4, "K", None), (5, "I", None), (6, "P", NUM), (7, "G", None)]:
            put(ca, r, c, f'=IF($K{r}="","",INDEX({JR(src)},$K{r}))', GREEN, fmt, align="center" if c in (1, 3, 7) else None)
evidence(50, 15, "T")
sec(ca, 67, "근거 전표 · 비교 기간 (상위 5건)")
header(ca, 68, EVH, h=20)
evidence(68, 5, "U")
ca.conditional_formatting.add("F51:F65", CellIsRule(operator="lessThan", formula=["0"], font=Font(color="C00000", bold=True)))
widths(ca, [13, 16, 12, 22, 34, 11, 10, 11, 8, 2, 6, 2, 2, 16, 8, 8, 8, 10, 40, 9, 9, 8, 2, 2, 22, 22])
ca.column_dimensions["E"].width = 34
ca.freeze_panes = "A15"

# =================================================================== 주의대상
wt = wb.create_sheet("주의대상")
title(wt, "주의 대상 TOP 10", "점수 = 증가액(30) + 증가율(20) + 이상치(25) + 예산(25), 작은 금액은 30%로 낮춤 · 원인은 원인분석 시트 C5에 순위 입력")
wt["A3"] = f'="기준: "&{plabel(S["P"])}&" · 단위: "&{S["ur"]}&" · 분기 한도 초과 "&COUNTIF({VR("AO")},"한도 초과")&"건 · 초과 예상 "&COUNTIF({VR("AO")},"초과 예상")&"건 · 주의 "&COUNTIF({VR("AO")},"주의")&"건 · 무예산 "&COUNTIF({VR("AO")},"무예산")&"건"'
wt["A3"].font = F(GREEN, size=9)
header(wt, 5, ["순위", "부서", "계정", "당월", "전월\n대비", "전년동월\n대비", "분기 판정", "분기\n소진율", "연간 판정", "빠른 신호", "점수", "키", "확인된 원인 (코멘트)"], h=32)
header(wt, 5, ["행"], start=14, fill=GRAYH, h=32)
header(wt, 5, ["당분기", "전분기\n대비", "당해\n누계", "전년동기\n누계 대비"], start=15, h=32)
header(wt, 5, ["연간가용", "월 소진속도", "절편", "추세\n연말 누계", "고갈\n예상(개월)", "고갈\n예상일", "추세 판정", "분기 고갈\n예상일"] + [f"누계\n{m}월" for m in range(1, 13)] + ["전년 남은기간\n월평균"], start=19, fill=GRAYH, h=32)
header(wt, 5, ["전월\n증감액", "전분기\n증감액", "전년동기누계\n증감액"], start=40, h=32)
wt.cell(row=4, column=19, value="고갈 추세 (대시보드 주의 대상 상세 그래프용)").font = F(GRAY, size=9)
for k in range(1, 11):
    r = 5 + k
    put(wt, r, 1, k, align="center", bold=True)
    put(wt, r, 14, f'=IFERROR(MATCH(LARGE({VR("AB")},A{r}),{VR("AB")},0),"")', GRAY)
    for c, src, fmt in [(2, "B", None), (3, "D", None), (4, "F", NUM), (5, "H", RATE), (6, "J", RATE), (7, "AO", None), (8, "AN", PCT), (9, "W", None), (10, "Z", None), (11, "AA", "0.0"), (12, "E", None),
                        (15, "K", NUM), (16, "M", RATE), (17, "Q", NUM), (18, "AY", RATE),
                        (19, "P", NUM), (20, "AP", NUM), (21, "AQ", NUM), (22, "AR", NUM), (23, "AS", "0.0"), (24, "AT", None), (25, "AU", None), (26, "AV", None), (39, "AZ", NUM)]:
        put(wt, r, c, f'=IF($N{r}="","",INDEX({VR(src)},$N{r}))', GREEN, fmt, align="center" if c in (7, 9, 11) else None)
    for j in range(12):
        put(wt, r, 27 + j, f'=IF($N{r}="","",INDEX(월별집계!${CL(28 + j)}$7:${CL(28 + j)}${VA_END},$N{r}))', GREEN, NUM)
    for c, a_, b_ in ((40, "F", "G"), (41, "K", "L"), (42, "Q", "AX")):   # 증감액 (전월 · 전분기 · 전년동기누계)
        put(wt, r, c, f'=IF($N{r}="","",INDEX({VR(a_)},$N{r})-INDEX({VR(b_)},$N{r}))', GREEN, NUM)
    put(wt, r, 13, f'=IF($L{r}="","",IFERROR(INDEX(코멘트!$F$5:$F$2004,MATCH({S["yr"]}&"|"&{S["mon"]}&"|"&$L{r},코멘트!$I$5:$I$2004,0)),"(미확인)"))', GREEN, wrap=True)
    wt.row_dimensions[r].height = 30
for val, fill in (('"한도 초과"', "F8CBAD"), ('"초과 예상"', "FCE4D6"), ('"부족 예상"', "F8CBAD"), ('"주의"', "FFE699"), ('"무예산"', "D9C3E9")):
    wt.conditional_formatting.add("G6:G15", CellIsRule(operator="equal", formula=[val], fill=PatternFill("solid", fgColor=fill)))
widths(wt, [6, 12, 18, 9, 8, 9, 10, 9, 8, 22, 6, 4, 50, 5, 9, 8, 10, 9] + [9] * 21 + [10, 10, 12])
wt.column_dimensions["L"].hidden = True
wt.freeze_panes = "A6"

# =================================================================== 리포트
rp = wb.create_sheet("리포트")
rp["A1"] = f'={S["co"]}&" TRACE 사업비 변동 점검 ("&{plabel(S["P"])}&")"'; rp["A1"].font = F(bold=True, size=16)
rp["A2"] = f'="단위: "&{S["ur"]}&" · 전표·TB·예산판 대사: "&계정변동!$P${AT}&" · 자가진단: "&자가진단!$B$4'; rp["A2"].font = F(GRAY, size=9)
kp = [("사업비 당월", "=월별집계!Z6", NUM), ("전월 대비", "=IF(월별집계!Y6<=0,0,월별집계!Z6/월별집계!Y6-1)", RATE),
      ("전년동월 대비", "=IF(월별집계!N6<=0,0,월별집계!Z6/월별집계!N6-1)", RATE),
      ("분기 한도 초과(예상)", f'=COUNTIF({VR("AO")},"한도 초과")+COUNTIF({VR("AO")},"초과 예상")', '0"건"'), ("분기 한도 주의", f'=COUNTIF({VR("AO")},"주의")', '0"건"'),
      ("무예산 집행", f'=COUNTIF({VR("AO")},"무예산")', '0"건"')]
for k, (lab, f_, fm) in enumerate(kp):
    c = rp.cell(row=4, column=1 + k, value=lab); c.font = F(GRAY, size=9); c.fill = TOT; c.border = BOX; c.alignment = Alignment(horizontal="center")
    c = rp.cell(row=5, column=1 + k, value=f_); c.font = F(GREEN, bold=True, size=14); c.number_format = fm
    c.fill = TOT; c.border = BOX; c.alignment = Alignment(horizontal="center", vertical="center")
rp.row_dimensions[5].height = 30
sec(rp, 7, "1. 주의 대상 TOP 10")
header(rp, 8, ["순위", "부서", "계정", "당월", "전월대비", "예산", "빠른 신호", "확인된 원인"], h=20)
for k in range(10):
    r = 9 + k; wr = 6 + k
    for c, src, fm in [(1, "A", None), (2, "B", None), (3, "C", None), (4, "D", NUM), (5, "E", RATE), (6, "G", None), (7, "J", None), (8, "M", None)]:
        put(rp, r, c, f"=주의대상!{src}{wr}", GREEN, fm, align="center" if c in (1, 6) else None, wrap=c == 8)
    rp.row_dimensions[r].height = 28
for val, fill in (('"한도 초과"', "F8CBAD"), ('"초과 예상"', "FCE4D6"), ('"부족 예상"', "F8CBAD"), ('"주의"', "FFE699"), ('"무예산"', "D9C3E9")):
    rp.conditional_formatting.add("F9:F18", CellIsRule(operator="equal", formula=[val], fill=PatternFill("solid", fgColor=fill)))
sec(rp, 20, "2. 원인 분석 (원인분석 시트에서 선택한 항목)")
rp.merge_cells("A21:H21"); rp["A21"] = "=원인분석!A18"; rp["A21"].font = F(GREEN, bold=True)
rp["A21"].alignment = Alignment(wrap_text=True, vertical="top"); rp["A21"].fill = BOXF; rp.row_dimensions[21].height = 46
header(rp, 22, ["순위", "패턴", "근거 강도", "설명"], h=20); rp.merge_cells("D22:H22")
for k in range(4):
    r = 23 + k; cr = 21 + k
    put(rp, r, 1, f"=원인분석!A{cr}", GREEN, align="center"); put(rp, r, 2, f"=원인분석!B{cr}", GREEN)
    put(rp, r, 3, f"=원인분석!C{cr}", GREEN, align="center"); put(rp, r, 4, f"=원인분석!D{cr}", GREEN, wrap=True)
    for c in range(5, 9): put(rp, r, c, None)
    rp.merge_cells(start_row=r, start_column=4, end_row=r, end_column=8); rp.row_dimensions[r].height = 30
sec(rp, 28, "3. 월별 사업비 (전사, 당해 vs 전년)")
put(rp, 29, 11, "월", bold=True, fill=SUB); put(rp, 29, 12, "=" + S["yr"] + '&"년"', bold=True, fill=SUB); put(rp, 29, 13, "=" + S["yr"] + '-1&"년"', bold=True, fill=SUB)
for m in range(1, 13):
    r = 29 + m
    put(rp, r, 11, f"{m}월", align="center")
    put(rp, r, 12, f'=IF({m}>{S["mon"]},0,SUMIFS({JR("P")},{JR("M")},{S["yr"]}*12+{m}))', GREEN, NUM)
    put(rp, r, 13, f'=SUMIFS({JR("P")},{JR("M")},({S["yr"]}-1)*12+{m})', GREEN, NUM)
ch = BarChart(); ch.type = "col"; ch.grouping = "clustered"; ch.style = 10
ch.title = "월별 사업비"; ch.y_axis.numFmt = "#,##0"; ch.height = 8; ch.width = 24
ch.add_data(Reference(rp, min_col=12, max_col=13, min_row=29, max_row=41), titles_from_data=True)
ch.set_categories(Reference(rp, min_col=11, min_row=30, max_row=41))
rp.add_chart(ch, "A29")
widths(rp, [6, 12, 18, 10, 9, 10, 20, 46, 2, 2, 6, 10, 10])
rp.page_setup.orientation = "landscape"; rp.page_setup.fitToWidth = 1; rp.page_setup.fitToHeight = 0
rp.sheet_properties.pageSetUpPr.fitToPage = True; rp.print_area = "A1:H46"

# =================================================================== 자가진단
ck = wb.create_sheet("자가진단")
title(ck, "자가진단", "외부망으로 오류를 알릴 때는 [점검 항목]과 [상태]만 전달하세요")
header(ck, 6, ["No", "점검 항목", "결과", "상태", "조치 방법"], h=20)
def hdrchk(sheet, hdrs):
    mis = "&".join(f'IF({sheet}!{CL(i+1)}4<>"{h}","{CL(i+1)} ","")' for i, h in enumerate(hdrs))
    return f'=IF({mis}="","일치","불일치 열: "&{mis})'
cnt = lambda sheet, col, end: f"COUNTA({sheet}!{col}5:{col}{end})"
checks = [
    ("설정값", f'=IF(AND(ISNUMBER({S["mon"]}),{S["mon"]}>=1,{S["mon"]}<=12,ISNUMBER({S["fx"]})),"정상","기준월 또는 단위 오류")', 'IF(C{r}="정상","OK","확인필요")', "설정 시트 기준월(1~12)·단위 확인"),
    ("전표_원본 머리글", hdrchk("전표_원본", JE_HDR), 'IF(C{r}="일치","OK","확인필요")', "4행 머리글을 표준 이름으로"),
    ("TB_원본 머리글", hdrchk("TB_원본", TB_HDR), 'IF(C{r}="일치","OK","확인필요")', "4행 머리글을 표준 이름으로"),
    ("예산판_원본 머리글", hdrchk("예산판_원본", BP_HDR), 'IF(C{r}="일치","OK","확인필요")', "4행 머리글을 표준 이름으로"),
    ("전표 보조수식 범위", f'={cnt("전표_원본","A",JE_END)}&"행 / 보조 "&{cnt("전표_원본","O",JE_END)}&"행"', f'IF({cnt("전표_원본","O",JE_END)}>={cnt("전표_원본","A",JE_END)},"OK","확인필요")', "M~X열을 데이터 마지막 행까지 채우기"),
    ("TB 보조수식 범위", f'={cnt("TB_원본","A",TB_END)}&"행 / 보조 "&{cnt("TB_원본","I",TB_END)}&"행"', f'IF({cnt("TB_원본","I",TB_END)}>={cnt("TB_원본","A",TB_END)},"OK","확인필요")', "I~K열을 데이터 마지막 행까지 채우기"),
    ("예산판 보조수식 범위", f'={cnt("예산판_원본","A",BP_END)}&"행 / 보조 "&{cnt("예산판_원본","L",BP_END)}&"행"', f'IF({cnt("예산판_원본","L",BP_END)}>={cnt("예산판_원본","A",BP_END)},"OK","확인필요")', "L~R열을 데이터 마지막 행까지 채우기"),
    ("기준월 전표 존재", f'=COUNTIF({JR("M")},{S["P"]})&"건"', f'IF(COUNTIF({JR("M")},{S["P"]})>0,"OK","확인필요")', "기준월 전표가 들어 있는지 확인"),
    ("기준월 TB 존재", f'=COUNTIF({TR("I")},{S["P"]})&"행"', f'IF(COUNTIF({TR("I")},{S["P"]})>0,"OK","확인필요")', "기준월 TB가 들어 있는지 확인"),
    ("전년 동월 전표 존재", f'=COUNTIF({JR("M")},{S["P"]}-12)&"건"', f'IF(COUNTIF({JR("M")},{S["P"]}-12)>0,"OK","참고")', "전년동월 비교·이상치 판정에 필요 (최근 24개월 권장)"),
    ("미매핑 코스트센터", f'=COUNTIF({JR("N")},"미매핑")&"건"&IF(COUNTIF({JR("N")},"미매핑")>0," (예: "&INDEX({JR("G")},MATCH("미매핑",{JR("N")},0))&")","")', f'IF(COUNTIF({JR("N")},"미매핑")=0,"OK","확인필요")', "부서매핑에 코스트센터 추가"),
    ("미매핑 G/L 계정", f'=COUNTIF({JR("W")},"미매핑")&"건"&IF(COUNTIF({JR("W")},"미매핑")>0," (예: "&INDEX({JR("F")},MATCH("미매핑",{JR("W")},0))&")","")', f'IF(COUNTIF({JR("W")},"미매핑")=0,"OK","확인필요")', "계정매핑에 계정 추가"),
    ("TB ↔ 전표 당월 대사", f'="차이 "&TEXT(계정변동!L{AT},"#,##0.0")', f'IF(ABS(계정변동!L{AT})<0.05,"OK","확인필요")', "계정변동 시트에서 차이 나는 계정 확인 (전표 추출 범위 누락 등)"),
    ("예산판 ↔ 전표 당분기 대사", f'="차이 "&TEXT(계정변동!O{AT},"#,##0.0")', f'IF(ABS(계정변동!O{AT})<0.05,"OK","확인필요")', "예산판 추출 시점과 전표 기간 일치 여부 확인"),
    ("변동분석에 없는 조합 (전표)", f'=COUNTIFS({JR("V")},"N",{JR("A")},{S["yr"]})&"건"', f'IF(COUNTIFS({JR("V")},"N",{JR("A")},{S["yr"]})=0,"OK","확인필요")', "변동분석에 부서코드·G/L 행 추가"),
    ("변동분석에 없는 조합 (예산판)", f'=COUNTIFS({BR("Q")},"N",{BR("M")},{S["yr"]})&"행"', f'IF(COUNTIFS({BR("Q")},"N",{BR("M")},{S["yr"]})=0,"OK","확인필요")', "변동분석에 부서코드·G/L 행 추가"),
    ("무예산 집행", f'=COUNTIF({VR("AO")},"무예산")&"건"', f'IF(COUNTIF({VR("AO")},"무예산")=0,"OK","참고")', "업무상 예외 - 전용·추가 배정 여부 확인"),
    ("원인분석 선택 항목", '=원인분석!$E$9&" · 당기 전표 "&COUNTIF(' + JR("Q") + ',"당기")&"건"', f'IF(COUNTIF({JR("Q")},"당기")>0,"OK","참고")', "원인분석 C5~C8 선택값 확인"),
]
# ---- 입력양식: 실제 SAP 화면 양식이 정해지면 이 시트만 고쳐서 맞춘다 (VBA 불러오기가 사용) ----
FORM = [  # (원본 시트, 표준 열 목록[(열, 필수, 대체 이름, 없을 때)], 빈 행 판정 열 번호)
    ("예산판_원본", [("회계연도", "필수", "연도,년도,GJAHR", ""), ("부서코드", "필수", "부서,팀코드,조직코드", ""),
                    ("부서명", "선택", "팀명,조직명", "부서매핑의 부서명 사용"), ("G/L계정", "필수", "계정,계정코드,예산계정,G/L,RACCT,HKONT", ""),
                    ("계정명", "선택", "계정과목명,TXT50", "계정매핑의 계정명 사용"), ("분기", "필수", "분기구분,QTR", "1~4, Q1, 1분기 모두 가능"),
                    ("배정", "필수", "배정액,배정예산,예산", ""), ("조정", "선택", "조정액,전용", "0으로 처리"),
                    ("실적", "선택", "집행,집행액,실적액", "예산판↔전표 대사 불가"), ("잔여", "선택", "잔액,잔여예산", "참고용")], 2),
    ("TB_원본", [("회계연도", "필수", "연도,GJAHR,RYEAR", ""), ("기간", "필수", "월,회계기간,POPER", "1~12 또는 001~012"),
                ("G/L계정", "필수", "계정,계정코드,G/L,RACCT", ""), ("계정명", "선택", "계정과목명,TXT50", ""),
                ("당기차변", "필수", "차변,당월차변,차변금액", ""), ("당기대변", "필수", "대변,당월대변,대변금액", ""),
                ("기말잔액", "선택", "잔액,기말", "참고용")], 3),
    ("전표_원본", [("회계연도", "필수", "연도,GJAHR", ""), ("기간", "필수", "월,회계기간,POPER", ""),
                  ("전표번호", "선택", "BELNR,문서번호", "근거 전표에 번호 없이 표시"), ("전기일", "선택", "전기일자,일자,BUDAT", ""),
                  ("전표유형", "선택", "문서유형,BLART", "'재분류' 패턴 판정 불가"), ("G/L계정", "필수", "계정,계정코드,HKONT,RACCT", ""),
                  ("코스트센터", "필수", "비용센터,KOSTL,CC", ""), ("금액", "필수", "금액(원),현지통화금액,DMBTR,HSL", "차대구분 방식이면 양수 금액"),
                  ("적요", "선택", "텍스트,내역,SGTXT", "'적요 키워드' 패턴 판정 불가"), ("거래처코드", "선택", "거래처,공급업체,LIFNR", ""),
                  ("거래처명", "선택", "거래처명칭,공급업체명,NAME1", "거래처코드로 대신 표시 (둘 다 없으면 거래처 분석 불가)")], 6),
    ("BS_원본", [("회계연도", "필수", "연도,GJAHR", ""), ("기간", "필수", "월,회계기간", ""),
                ("항목코드", "필수", "코드,계정코드", "없으면 항목명으로 매핑"), ("항목명", "선택", "항목,계정명,과목", ""),
                ("구분", "선택", "", ""), ("레벨", "선택", "수준,LEVEL", ""), ("기말잔액", "필수", "당월말,잔액,금액", "")], 3),
    ("PL_원본", [("회계연도", "필수", "연도,GJAHR", ""), ("기간", "필수", "월,회계기간", ""),
                ("항목코드", "필수", "코드,계정코드", "없으면 항목명으로 매핑"), ("항목명", "선택", "항목,계정명,과목", ""),
                ("구분", "선택", "", ""), ("레벨", "선택", "수준,LEVEL", ""), ("당월금액", "필수", "당월,금액", ""),
                ("누계금액", "권장", "누계,누적", "당기순이익 누계 KPI가 0으로 표시")], 3),
    ("예정사업비", [("회계연도", "필수", "연도", ""), ("기간", "필수", "월,회계기간", ""), ("포트폴리오", "필수", "포트폴리오코드,PF", ""),
                   ("포트폴리오명", "선택", "", ""), ("신계약비", "필수", "예정신계약비", ""), ("유지비", "필수", "예정유지비", ""),
                   ("수금비", "필수", "예정수금비", "")], 3),
    ("부채변동표", [("회계연도", "선택", "연도", "설정의 기준연도로 채우고 그 달 자료만 교체"), ("기간", "선택", "월,회계기간", "설정의 기준월로 채움"), ("포트폴리오", "필수", "포트폴리오코드,PF", ""),
                   ("포트폴리오명", "선택", "", ""), ("측정모형", "필수", "모형,측정방법", "GMM·PAA·VFA"),
                   ("변동요인코드", "필수", "요인코드,코드", "없으면 변동요인명으로 매핑"), ("변동요인", "선택", "요인,항목,변동내역", ""),
                   ("BEL", "필수", "최선추정부채,최선추정", ""), ("RA", "필수", "위험조정", ""), ("CSM", "필수", "보험계약마진,계약서비스마진", "")], 3),
]
FM = {}  # 시트별 블록 시작 행 (헤더설정 시트)
r = 7
for nm_, cols_, key in FORM:
    FM[nm_] = (r + 2, len(cols_), key); r += 2 + len(cols_) + 2
FSS_CODE = {"A000": "10000", "A100": "11000", "A110": "11100", "A120": "11200", "A130": "11300", "A140": "11400", "A200": "12000", "A300": "13000",
            "A900": "19000", "L000": "20000", "L100": "21000", "L110": "21100", "L120": "21200", "L130": "21300", "L200": "22000", "L900": "29000",
            "E000": "30000", "E100": "31000", "E200": "32000", "E300": "33000", "P100": "40000", "P110": "41000", "P120": "42000", "P130": "43000",
            "P200": "50000", "P210": "51000", "P220": "52000", "P230": "53000", "P300": "60000", "P400": "70000", "P500": "80000", "P600": "90000"}
BS_STD = [("A000", "자산총계"), ("A100", "운용자산"), ("A110", "현금및예치금"), ("A120", "유가증권"), ("A130", "대출채권"), ("A140", "투자부동산"),
          ("A200", "보험계약자산"), ("A300", "재보험계약자산"), ("A900", "기타자산"), ("L000", "부채총계"), ("L100", "보험계약부채"),
          ("L110", "최선추정부채(BEL)"), ("L120", "위험조정(RA)"), ("L130", "보험계약마진(CSM)"), ("L200", "재보험계약부채"), ("L900", "기타부채"),
          ("E000", "자본총계"), ("E100", "자본금"), ("E200", "이익잉여금"), ("E300", "기타포괄손익누계액")]
PL_STD = [("P100", "보험손익"), ("P110", "보험수익"), ("P120", "보험서비스비용"), ("P130", "재보험손익"), ("P200", "투자손익"), ("P210", "투자수익"),
          ("P220", "보험금융비용"), ("P230", "재산관리비"), ("P300", "기타사업비용 (사업비)"), ("P400", "법인세차감전이익"), ("P500", "법인세비용"), ("P600", "당기순이익")]

# ---------- 헤더설정: 불러오기가 찾는 열 이름
hs = wb.create_sheet("헤더설정")
title(hs, "헤더설정 (불러오기가 찾는 열 이름)", "노란 칸만 고치세요 · SAP 화면에서 내려받은 파일의 머리글과 똑같이 적으면 그 열을 찾아 표준 위치에 넣습니다")
for i, ln in enumerate(["찾는 순서: ① '찾을 헤더 이름' → ② '대체 이름'(쉼표로 여러 개) 중 먼저 맞는 열. 띄어쓰기·대소문자·줄바꿈은 무시합니다.",
                        "필수 열을 찾지 못하면 불러오기를 멈추고 어떤 열이 없는지 알려줍니다. 선택 열은 없으면 빈칸으로 두고 '없을 때'에 적힌 기능만 빠집니다.",
                        "수식은 표준 열 기준이라 이 표를 고쳐도 다른 시트는 건드릴 필요가 없습니다."]):
    hs.cell(row=3 + i, column=2, value=ln).font = F(GRAY, size=9)
for nm_, cols_, key in FORM:
    r0, n, _ = FM[nm_]
    hs.cell(row=r0 - 2, column=2, value=f"{nm_}  (빈 행 판정 열: {cols_[key-1][0]})").font = F(bold=True, size=11)
    header(hs, r0 - 1, ["표준 열 (원본 시트)", "찾을 헤더 이름", "대체 이름 (쉼표로 구분)", "필수", "없을 때"], start=2, h=20)
    for i, (c_, req, alias, miss) in enumerate(cols_):
        put(hs, r0 + i, 2, f"{CL(i+1)}열 {c_}", bold=True, fill=SUB)
        put(hs, r0 + i, 3, c_, BLUE, fill=YEL)
        put(hs, r0 + i, 4, alias, BLUE)
        put(hs, r0 + i, 5, req, BLACK if req == "필수" else GRAY, align="center")
        put(hs, r0 + i, 6, miss, GRAY)
hs.freeze_panes = "A7"; widths(hs, [2, 24, 24, 38, 8, 52]); hs.sheet_properties.tabColor = "FFC000"

# ---------- 입력옵션: 행 제외, 전표 금액 방식, BS·PL 항목 매핑
os_ = wb.create_sheet("입력옵션")
title(os_, "입력옵션", "노란 칸만 고치세요 · 불러오기 규칙과 BS·PL 항목 매핑")
sec(os_, 4, "불러오기 규칙")
for rr, k, v, d in [(5, "제외할 행 키워드", "합계,소계,총계,Total", "빈 행 판정 열에 이 단어가 있으면 불러오지 않음 (쉼표로 구분)"),
                    (6, "머리글 찾는 범위 (행)", 20, "파일 위쪽 이 행 수 안에서 헤더 이름이 가장 많이 맞는 행을 머리글로 사용"),
                    (8, "전표 금액 방식", "부호 있는 금액", "'부호 있는 금액' 또는 '차대구분+양수금액' (SAP 차대구분 S/H)"),
                    (9, "차대구분 헤더 이름", "차대구분", "차대구분+양수금액 방식일 때 찾을 열 이름 (예: SHKZG)"),
                    (10, "차변 표시 값", "S", ""), (11, "대변 표시 값", "H", "이 값인 행은 금액에 −를 붙여 불러옴"),
                    (13, "BS·PL 양식 형태", "세로 (기간×항목)", "'감독원 (가로)'이면 항목이 행, 금액이 열인 양식을 기준월 한 달씩 불러와 기존 자료에 합침"),
                    (14, "가로: 당기 금액 헤더", "당기", "BS 기말잔액 · PL 당월 금액 열 이름 (쉼표로 대체 이름)"),
                    (15, "가로: 전기 금액 헤더", "전기", "BS 비교 열 (없으면 비워둠)"),
                    (16, "가로: PL 누계 헤더", "당기누계", "PL 누계 금액 열"),
                    (17, "가로: 전기 열의 기간", "전년말", "'전년말'(전년 12월) 또는 '전년동월'")]:
    put(os_, rr, 2, k, bold=True, fill=SUB); put(os_, rr, 3, v, BLUE, fill=YEL, align="center"); put(os_, rr, 4, d, GRAY)
dv = DataValidation(type="list", formula1='"부호 있는 금액,차대구분+양수금액"'); os_.add_data_validation(dv); dv.add("C8")
dv = DataValidation(type="list", formula1='"세로 (기간×항목),감독원 (가로)"'); os_.add_data_validation(dv); dv.add("C13")
dv = DataValidation(type="list", formula1='"전년말,전년동월"'); os_.add_data_validation(dv); dv.add("C17")
MAP_TOP = 21
MAP0 = MAP_TOP + 1; MAP1 = MAP0 + len(BS_STD) + len(PL_STD) - 1
sec(os_, MAP_TOP - 1, "BS·PL 항목 매핑 — 회사 재무제표의 항목코드(없으면 항목명)를 표준 코드에 연결. 차트·KPI·검산은 표준 코드만 사용")
header(os_, MAP_TOP, ["구분", "표준 코드", "표준 항목명", "실제 항목코드", "실제 항목명"], start=2, h=20)
for i, (cls, code, nm) in enumerate([("BS", c, n) for c, n in BS_STD] + [("PL", c, n) for c, n in PL_STD]):
    r = MAP0 + i
    put(os_, r, 2, cls, align="center"); put(os_, r, 3, code, align="center"); put(os_, r, 4, nm)
    c = put(os_, r, 5, FSS_CODE[code], BLUE, fill=YEL, align="center"); c.number_format = "@"
    put(os_, r, 6, nm, BLUE)
widths(os_, [2, 24, 18, 52, 14, 24]); os_.sheet_properties.tabColor = "FFC000"
MAPC = lambda col: f"입력옵션!${col}${MAP0}:${col}${MAP1}"
def code_map(r):  # BS/PL 보조수식: 실제 코드 → 표준 코드 (코드가 없으면 항목명으로, 둘 다 없으면 원래 코드)
    return (f'=IF($A{r}="","",IFERROR(INDEX({MAPC("C")},MATCH($C{r}&"",{MAPC("E")},0)),'
            f'IFERROR(INDEX({MAPC("C")},MATCH($D{r}&"",{MAPC("F")},0)),$C{r}&"")))')

# ---- v0.4: IFRS17(GMM) 사업비 배분 기준
DEPT_ALLOC = {"D100": (0, .10, 0, .90), "D200": (0, .20, .10, .70), "D300": (.05, .15, 0, .80), "D400": (.50, .10, 0, .40),
              "D500": (.15, 0, 0, .85), "D600": (.75, .10, 0, .15), "D700": (.10, .40, .10, .40), "D800": (0, .60, .30, .10),
              "D900": (.90, .05, .05, 0), "D950": (0, .30, 0, .70)}   # 공통비 부서별 대체비율 (신계약, 유지, 수금, 비귀속)
INV_DEPTS = ["D950"]; INV_RATE = {2025: 0.65, 2026: 0.70}           # 투자관련부서 유지비 공통비 → 재산관리비 (연 1회 변경)
CC2DEPT = {c[0]: c[2] for c in CCS}
PFS = [("PF01", "보장성", "GMM"), ("PF02", "저축성", "GMM"), ("PF03", "단체(단기)", "PAA"), ("PF04", "변액", "VFA")]
PF_DRV = {"PF01": (.55, .50, .50), "PF02": (.20, .30, .30), "PF03": (.10, .08, .08), "PF04": (.15, .12, .12)}  # (신계약, 유지, 수금)
# IFRS 17 반영 정책 (예시값 · 회사 회계정책으로 바꿈). B96(a): 보험료 연동 취득현금흐름의 예실차만 CSM, 나머지는 B97(c)로 당기손익
PF_LINK = {"PF01": 0.6, "PF02": 0.6, "PF03": 0.0, "PF04": 0.6}   # 신계약비 중 보험료에 연동되는 비율 (예: 보험료 비례 수수료)
PF_ONEROUS = {"PF01": "N", "PF02": "N", "PF03": "N", "PF04": "N"}   # 손실부담 집단이면 Y → 미래 서비스 관련 변동도 손실요소·당기손익
NB_POLICY = os.environ.get("NB_POLICY", "보험료연동분만 CSM")        # 보험료연동분만 CSM / 전액 CSM / 전액 당기손익
PAA_POLICY = "발생 즉시 비용"                                        # 발생 즉시 비용(59(a)) / 이연 후 상각
def alloc_of(e):   # → [비귀속, 신계약, 유지, 수금, 재산관리비] 비율
    if e["gl"] in DIRECT:
        return [1.0 if DIRECT[e["gl"]] == c else 0.0 for c in ("비귀속", "신계약비", "유지비", "수금비")] + [0.0]
    d = CC2DEPT[e["cc"]]
    inv = INV_RATE.get(e["y"], 0) if (e["gl"] not in NATURE_OTHER and d in INV_DEPTS) else 0.0
    sn, mt, cl, na = DEPT_ALLOC.get(d, (0, 0, 0, 1))
    return [(1 - inv) * na, (1 - inv) * sn, (1 - inv) * mt, (1 - inv) * cl, inv]
ATT = {}  # (y, m) -> [비귀속, 신계약, 유지, 수금, 재산관리비] 원
for e in E:
    a = ATT.setdefault((e["y"], e["m"]), [0, 0, 0, 0, 0])
    for i, rt in enumerate(alloc_of(e)): a[i] += e["amt"] * rt
irng = random.Random(41)
EXPROWS = []  # 예정사업비: 직전 12개월 실제 평균 기준 (계리 가정 대용)
for idx, (y, m) in enumerate(MONTHS):
    prev = [ATT[k] for k in MONTHS[max(0, idx - 12):idx]] or [ATT[k] for k in MONTHS[:12]]
    base = [sum(p[i] for p in prev) / len(prev) for i in range(4)]
    for pf, pfn, mdl in PFS:
        d_ = PF_DRV[pf]
        EXPROWS.append([y, m, pf, pfn] + [round(base[1 + j] * d_[j] * irng.uniform(0.98, 1.02), -3) for j in range(3)])

# ---- v0.3 확장: BS/PL 입력 시트, 차트데이터, 대시보드 (tool.py의 자가진단 섹션 앞에 삽입) ----
from openpyxl.chart import LineChart, DoughnutChart
BS_ITEMS = [("A000", "자산총계", "자산", 1), ("A100", "운용자산", "자산", 2), ("A110", "현금및예치금", "자산", 3), ("A120", "유가증권", "자산", 3),
            ("A130", "대출채권", "자산", 3), ("A140", "투자부동산", "자산", 3), ("A200", "보험계약자산", "자산", 2), ("A300", "재보험계약자산", "자산", 2),
            ("A900", "기타자산", "자산", 2), ("L000", "부채총계", "부채", 1), ("L100", "보험계약부채", "부채", 2), ("L110", "최선추정부채(BEL)", "부채", 3),
            ("L120", "위험조정(RA)", "부채", 3), ("L130", "보험계약마진(CSM)", "부채", 3), ("L200", "재보험계약부채", "부채", 2), ("L900", "기타부채", "부채", 2),
            ("E000", "자본총계", "자본", 1), ("E100", "자본금", "자본", 2), ("E200", "이익잉여금", "자본", 2), ("E300", "기타포괄손익누계액", "자본", 2)]
PL_ITEMS = [("P100", "보험손익", "손익", 1), ("P110", "보험수익", "수익", 2), ("P120", "보험서비스비용", "비용", 2), ("P130", "재보험손익", "손익", 2),
            ("P200", "투자손익", "손익", 1), ("P210", "투자수익", "수익", 2), ("P220", "보험금융비용", "비용", 2), ("P230", "재산관리비", "비용", 2),
            ("P300", "기타사업비용 (사업비)", "비용", 1), ("P400", "법인세차감전이익", "손익", 1), ("P500", "법인세비용", "비용", 1), ("P600", "당기순이익", "손익", 1)]
brng = random.Random(23)
BSROWS, PLROWS = [], []
cum = {}
for i, (y, m) in enumerate(MONTHS):
    n = lambda base, g, sd: base * (1 + g * i) * brng.uniform(1 - sd, 1 + sd)
    v = {"A110": n(85000, .002, .03), "A120": n(2100000, .004, .005), "A130": n(420000, .002, .01), "A140": 180000 - 650 * i,
         "A200": n(12000, 0, .05), "A300": n(41000, .001, .02), "A900": n(160000, 0, .02),
         "L110": n(2050000, .0035, .004), "L120": n(118000, .002, .01), "L130": 372000 + 1800 * i + brng.uniform(-900, 900),
         "L200": n(6800, 0, .05), "L900": n(95000, 0, .03), "E100": 100000, "E300": -40000 + brng.uniform(-5000, 5000)}
    v = {k: round(x, 1) for k, x in v.items()}
    v["A100"] = v["A110"] + v["A120"] + v["A130"] + v["A140"]; v["A000"] = v["A100"] + v["A200"] + v["A300"] + v["A900"]
    v["L100"] = v["L110"] + v["L120"] + v["L130"]; v["L000"] = v["L100"] + v["L200"] + v["L900"]
    v["E200"] = round(v["A000"] - v["L000"] - v["E100"] - v["E300"], 1); v["E000"] = v["E100"] + v["E200"] + v["E300"]
    for code, nm, cls, lv in BS_ITEMS:
        BSROWS.append([y, m, code, nm, cls, lv, round(v[code] * MIL)])
    p = {"P110": n(62000, .003, .03), "P120": n(50500, .003, .04), "P130": -1100 + brng.uniform(-200, 200),
         "P210": n(9500, 0, .08), "P220": n(7600, .003, .02)}
    p = {k: round(x * MIL, -3) for k, x in p.items()}
    p["P300"] = round(ATT[(y, m)][0]); p["P230"] = round(ATT[(y, m)][4])
    p["P100"] = p["P110"] - p["P120"] + p["P130"]; p["P200"] = p["P210"] - p["P220"] - p["P230"]
    p["P400"] = p["P100"] + p["P200"] - p["P300"]; p["P500"] = round(p["P400"] * 0.22, -3); p["P600"] = p["P400"] - p["P500"]
    for code, nm, cls, lv in PL_ITEMS:
        if m == 1: cum[code] = 0
        cum[code] = cum.get(code, 0) + p[code]
        PLROWS.append([y, m, code, nm, cls, lv, p[code], cum[code]])

FAC = [("F00", "기초잔액"), ("F10", "신계약 효과"), ("F21", "가정변경-사업비"), ("F22", "가정변경-기타"), ("F31", "경험조정-취득현금흐름"),
       ("F32", "경험조정-보험료 등"), ("F40", "보험금융비용(이자부리)"), ("F51", "CSM 상각"), ("F52", "RA 해제"), ("F60", "현금흐름 및 기타"), ("F99", "기말잔액")]
SH = {"BEL": {"PF01": .50, "PF02": .33, "PF03": .02, "PF04": .15}, "RA": {"PF01": .60, "PF02": .25, "PF03": .02, "PF04": .13},
      "CSM": {"PF01": .62 / .87 * .87, "PF02": .25, "PF03": 0, "PF04": .13}}
bsv_ = {(r_[0], r_[1], r_[2]): r_[6] / MIL for r_ in BSROWS}
frng = random.Random(53)
ROLL = []; csm_close = {pf: 370500 * SH["CSM"][pf] for pf, _, _ in PFS}
bel_prev = {pf: bsv_[(2025, 1, "L110")] / 1.0035 * SH["BEL"][pf] for pf, _, _ in PFS}
ra_prev = {pf: bsv_[(2025, 1, "L120")] / 1.002 * SH["RA"][pf] for pf, _, _ in PFS}
EXPV = {(r_[0], r_[1], r_[2]): r_[4] / MIL for r_ in EXPROWS}
CSM_TOT = {}
for (y, m) in MONTHS:
    nb_tot = 4500 * frng.uniform(0.85, 1.15)
    for pf, pfn, mdl in PFS:
        f = {k: [0.0, 0.0, 0.0] for k, _ in FAC}   # BEL, RA, CSM
        b0, r0, c0 = bel_prev[pf], ra_prev[pf], csm_close[pf]
        b1, r1 = bsv_[(y, m, "L110")] * SH["BEL"][pf], bsv_[(y, m, "L120")] * SH["RA"][pf]
        f["F00"] = [b0, r0, c0]
        if mdl != "PAA":
            nb = nb_tot * SH["CSM"][pf] / (1 - SH["CSM"]["PF03"])
            f["F10"] = [-1.25 * nb, 0.25 * nb, nb]
            a21 = -1200 if (pf == "PF01" and (y, m) == (2026, 9)) else (400 if (pf == "PF01" and (y, m) == (2025, 12)) else (-250 if (pf == "PF02" and (y, m) == (2025, 12)) else 0))
            f["F21"] = [-a21, 0, a21]
            act_nb = ATT[(y, m)][1] / MIL * PF_DRV[pf][0]
            f["F31"] = [0, 0, (EXPV[(y, m, pf)] - act_nb) * (1.0 if NB_POLICY == "전액 CSM" else 0.0 if NB_POLICY == "전액 당기손익" else PF_LINK[pf]) * (PF_ONEROUS[pf] != "Y")]
            f["F32"] = [0, 0, frng.uniform(-300, 300)]
            f["F40"] = [b0 * 0.0025, r0 * 0.002, c0 * 0.0025]
            f["F51"] = [0, 0, -c0 * 0.009]
        f["F52"] = [0, -r0 * 0.006, 0]
        c1 = sum(f[k][2] for k, _ in FAC if k not in ("F00", "F99"))
        c1 += c0
        f["F60"] = [b1 - b0 - sum(f[k][0] for k, _ in FAC if k not in ("F00", "F60", "F99")),
                    r1 - r0 - sum(f[k][1] for k, _ in FAC if k not in ("F00", "F60", "F99")), 0]
        f["F99"] = [b1, r1, c1]
        for k, nm in FAC:
            ROLL.append([y, m, pf, pfn, mdl, k, nm] + [round(v * MIL, -3) for v in f[k]])
        bel_prev[pf], ra_prev[pf], csm_close[pf] = b1, r1, c1
    CSM_TOT[(y, m)] = sum(csm_close.values())
# BS: CSM과 그 합계·이익잉여금을 부채변동표에 맞춰 다시 계산
_bs = {}
for r_ in BSROWS: _bs[(r_[0], r_[1], r_[2])] = r_
for (y, m) in MONTHS:
    g_ = lambda c: _bs[(y, m, c)][6]
    _bs[(y, m, "L130")][6] = sum(rr[9] for rr in ROLL if (rr[0], rr[1], rr[5]) == (y, m, "F99"))
    _bs[(y, m, "L110")][6] = sum(rr[7] for rr in ROLL if (rr[0], rr[1], rr[5]) == (y, m, "F99"))
    _bs[(y, m, "L120")][6] = sum(rr[8] for rr in ROLL if (rr[0], rr[1], rr[5]) == (y, m, "F99"))
    _bs[(y, m, "L100")][6] = g_("L110") + g_("L120") + g_("L130")
    _bs[(y, m, "L000")][6] = g_("L100") + g_("L200") + g_("L900")
    _bs[(y, m, "E200")][6] = g_("A000") - g_("L000") - g_("E100") - g_("E300")
    _bs[(y, m, "E000")][6] = g_("E100") + g_("E200") + g_("E300")

BSE, PLE = 20000, 20000
BSR = lambda c: f"BS_원본!${c}$5:${c}${BSE}"
PLR = lambda c: f"PL_원본!${c}$5:${c}${PLE}"
BS_HDR = ["회계연도", "기간", "항목코드", "항목명", "구분", "레벨", "기말잔액"]
PL_HDR = ["회계연도", "기간", "항목코드", "항목명", "구분", "레벨", "당월금액", "누계금액"]
for nm_, hdr, rows, hcols, sub in [("BS_원본", BS_HDR, BSROWS, ["기간번호", "항목코드\n텍스트", "기말잔액\n(보고단위)"], "재무상태표 · 월말 기준 · 한 행 = 기간 × 항목"),
                                   ("PL_원본", PL_HDR, PLROWS, ["기간번호", "항목코드\n텍스트", "당월\n(보고단위)", "누계\n(보고단위)"], "손익계산서 · 한 행 = 기간 × 항목 · 비용 항목은 양수로 입력")]:
    w_ = wb.create_sheet(nm_)
    title(w_, nm_.replace("_원본", "") + " 원본", f"A~{CL(len(hdr))}: 붙여넣기 · 보조수식은 회색 머리글 · {sub}")
    header(w_, 4, hdr)
    h0 = len(hdr) + 2
    header(w_, 4, hcols, start=h0, fill=GRAYH)
    for k, row in enumerate(rows):
        for j, v_ in enumerate(row):
            c = w_.cell(row=5 + k, column=j + 1, value=v_); c.font = F(BLUE)
            if j >= 6: c.number_format = "#,##0"
    for r in range(5, 5 + len(rows)):
        fm = [f'=IF($A{r}="","",VALUE($A{r}&"")*12+VALUE($B{r}&""))', code_map(r), f'=IF($A{r}="",0,$G{r}*{S["fx"]})']
        if nm_ == "PL_원본": fm.append(f'=IF($A{r}="",0,$H{r}*{S["fx"]})')
        for j, f_ in enumerate(fm):
            c = w_.cell(row=r, column=h0 + j, value=f_); c.font = F(GREEN if j >= 2 else BLACK)
            if j >= 2: c.number_format = NUM0
    w_.freeze_panes = "A5"; w_.auto_filter.ref = f"A4:{CL(h0 + len(hcols) - 1)}4"
    widths(w_, [8, 5, 9, 20, 6, 5, 16, 16, 2, 8, 9, 13, 13][:h0 + len(hcols)])
BS_FILL, PL_FILL = 4 + len(BSROWS), 4 + len(PLROWS)
bsv = lambda code, p: f'SUMIFS({BSR("K")},{BSR("J")},"{code}",{BSR("I")},{p})'
plv = lambda code, p, col="L": f'SUMIFS({PLR(col)},{PLR("K")},"{code}",{PLR("J")},{p})'  # PL 보조: J 기간번호, K 코드, L 당월, M 누계

# ---------------- 차트데이터
cd = wb.create_sheet("차트데이터")
title(cd, "차트데이터", "대시보드(엑셀·HTML) 원천 · 모두 수식 · 금액: 보고단위")
cd["A3"] = f'="기준: "&{plabel(S["P"])}&" · 단위: "&{S["ur"]}'; cd["A3"].font = F(GREEN, size=9)
CDA = 4   # 월별 사업비
header(cd, CDA, ["월", "당해", "전년", "예산\n(분기가용÷3)"], h=30)
for m in range(1, 13):
    r = CDA + m
    put(cd, r, 1, f"{m}월", align="center")
    put(cd, r, 2, f'=IF({m}>{S["mon"]},0,SUMIFS({JR("P")},{JR("M")},{S["yr"]}*12+{m}))', GREEN, NUM)
    put(cd, r, 3, f'=SUMIFS({JR("P")},{JR("M")},({S["yr"]}-1)*12+{m})', GREEN, NUM)
    put(cd, r, 4, f'=SUMIFS({BR("O")},{BR("M")},{S["yr"]},{BR("N")},ROUNDUP({m}/3,0))/3', GREEN, NUM)
CDB = 19  # 부서별
header(cd, CDB, ["부서코드", "부서명", "분기 한도", "분기 실적", "분기말 예상", "분기 소진율", "판정"], h=30)
DL = list(DEPTS)
for k, d in enumerate(DL):
    r = CDB + 1 + k
    put(cd, r, 1, d, BLUE, align="center")
    put(cd, r, 2, f'=IFERROR(INDEX(부서매핑!$D$5:$D$504,MATCH($A{r},부서매핑!$C$5:$C$504,0)),$A{r})', GREEN)
    for c, col in ((3, "AJ"), (4, "AK"), (5, "AM")):
        put(cd, r, c, f'=SUMIFS({VR(col)},{VR("A")},$A{r})', GREEN, NUM)
    put(cd, r, 6, f"=IF(C{r}=0,0,E{r}/C{r})", fmt=PCT)
    put(cd, r, 7, f'=IF(C{r}=0,"-",IF(D{r}>C{r},"한도 초과",IF(F{r}>{S["short"]},"초과 예상",IF(F{r}>{S["warn"]},"주의","여유"))))', GREEN, align="center")
CDC = 30  # 판정 분포
header(cd, CDC, ["분기 판정", "건수"], h=20)
STAT = ["여유", "주의", "초과 예상", "한도 초과", "무예산"]
for k, s_ in enumerate(STAT):
    put(cd, CDC + 1 + k, 1, s_, align="center"); put(cd, CDC + 1 + k, 2, f'=COUNTIF({VR("AO")},A{CDC+1+k})', GREEN, "0")
CDD = 37  # 계정별
header(cd, CDD, ["G/L 계정", "계정명", "당월", "전년동월", "증감", "전월", "당분기", "전분기", "당해 누계", "전년동기 누계"], h=30)
for i in range(len(sgl)):   # 모두 전표 기준(변동분석 합계) — 같은 출처끼리 비교
    r = CDD + 1 + i
    put(cd, r, 1, f"=계정변동!A{6+i}", GREEN, align="center"); put(cd, r, 2, f"=계정변동!B{6+i}", GREEN)
    for c, col in ((3, "F"), (4, "I"), (6, "G"), (7, "K"), (8, "L"), (9, "Q"), (10, "AX")):
        put(cd, r, c, f'=SUMIFS({VR(col)},{VR("C")},$A{r}&"")', GREEN, NUM)
    put(cd, r, 5, f"=C{r}-D{r}", fmt=NUM)
CDE = 56  # BS 추이 (최근 12개월)
header(cd, CDE, ["기간번호", "기간", "자산총계", "부채총계", "자본총계", "보험계약마진\n(CSM)", "운용자산"], h=30)
for k in range(12):
    r = CDE + 1 + k
    put(cd, r, 1, f"={S['P']}-{11-k}", GRAY, "0"); put(cd, r, 2, "=" + plabel(f"A{r}"), align="center")
    for c, code in ((3, "A000"), (4, "L000"), (5, "E000"), (6, "L130"), (7, "A100")):
        put(cd, r, c, "=" + bsv(code, f"$A{r}"), GREEN, NUM0)
CDF = 71  # PL 월별
header(cd, CDF, ["월", "당기순이익\n당해", "당기순이익\n전년", "보험손익\n당해", "투자손익\n당해", "사업비\n당해", "보험손익\n전년"], h=30)
for m in range(1, 13):
    r = CDF + m
    cur, prv = f"{S['yr']}*12+{m}", f"({S['yr']}-1)*12+{m}"
    put(cd, r, 1, f"{m}월", align="center")
    for c, code, p_ in ((2, "P600", cur), (3, "P600", prv), (4, "P100", cur), (5, "P200", cur), (6, "P300", cur), (7, "P100", prv)):
        f_ = plv(code, p_)
        put(cd, r, c, f"=IF({m}>{S['mon']},0,{f_})" if p_ == cur else "=" + f_, GREEN, NUM0)
CDG = 86  # KPI
header(cd, CDG, ["지표", "값", "비교값", "증감률", "비교 기준"], h=20)
KPI = [("총자산", bsv("A000", S["P"]), bsv("A000", S["P"] + "-12"), "전년동월말"),
       ("자본총계", bsv("E000", S["P"]), bsv("E000", S["P"] + "-12"), "전년동월말"),
       ("보험계약마진(CSM)", bsv("L130", S["P"]), bsv("L130", S["P"] + "-12"), "전년동월말"),
       ("당기순이익 누계", plv("P600", S["P"], "M"), plv("P600", S["P"] + "-12", "M"), "전년 누계"),
       ("사업비 누계", f'SUMIFS({JR("P")},{JR("M")},">="&({S["yr"]}*12+1),{JR("M")},"<="&{S["P"]})',
        f'SUMIFS({JR("P")},{JR("M")},">="&(({S["yr"]}-1)*12+1),{JR("M")},"<="&({S["P"]}-12))', "전년 누계"),
       ("사업비 당월", "월별집계!Z6", "월별집계!Y6", "전월"),
       ("사업비 당분기", f'SUM({VR("K")})', f'SUM({VR("L")})', "전분기"),
       ("분기 한도 초과(예상)", f'COUNTIF({VR("AO")},"한도 초과")+COUNTIF({VR("AO")},"초과 예상")', f'COUNTIF({VR("AO")},"주의")', "주의 건수")]
for k, (lab, a, b, basis) in enumerate(KPI):
    r = CDG + 1 + k
    last = k == len(KPI) - 1   # 마지막 '분기 한도 초과(예상)'는 건수
    put(cd, r, 1, lab, bold=True); put(cd, r, 2, "=" + a, GREEN, "0" if last else NUM0); put(cd, r, 3, "=" + b, GREEN, "0" if last else NUM0)
    put(cd, r, 4, "" if last else f"=IF(C{r}<=0,0,B{r}/C{r}-1)", fmt=RATE); put(cd, r, 5, basis, GRAY)
CDH = 96  # 소진 추세: 1행 전사 + 고갈이 빠른 부서×계정 15개 (변동분석 AW열 정렬키 순)
NBURN = 15
header(cd, CDH, ["고갈 순위", "부서", "계정", "연간가용", "누계실적", "월 소진속도", "절편", "추세\n연말 누계", "고갈\n예상(개월)", "고갈\n예상일", "추세 판정", "분기 고갈\n예상일"]
       + [f"누계\n{k}월" for k in range(1, 13)], h=30)
header(cd, CDH, ["변동분석\n행", "분기 한도", "분기 실적", "전년 남은기간\n월평균"], start=25, fill=SUB, h=30)
r = CDH + 1
put(cd, r, 1, "전사", bold=True, align="center"); put(cd, r, 2, "전사", bold=True); put(cd, r, 3, "사업비 합계", bold=True)
put(cd, r, 4, f'=SUM({VR("P")})', GREEN, NUM); put(cd, r, 5, f'=SUM({VR("Q")})', GREEN, NUM)
put(cd, r, 26, f'=SUM({VR("AJ")})', GREEN, NUM); put(cd, r, 27, f'=SUM({VR("AK")})', GREEN, NUM)
for k in range(12):
    put(cd, r, 13 + k, f"=월별집계!{CL(28 + k)}6", GREEN, NUM)
put(cd, r, 28, f'=SUMIFS(월별집계!$C$6:$Z$6,월별집계!$C$5:$Z$5,">"&({S["P"]}-12),월별집계!$C$5:$Z$5,"<="&(({S["yr"]}-1)*12+12))/MAX(1,12-{S["mon"]})', GREEN, NUM)
for c, (f_, col, fmt) in burn_formulas(r, f'IF($A{r}="",""', f"$M{r}:$X{r}", f"D{r}", f"E{r}", "F", "G", "I", "K", f"AB{r}").items():
    put(cd, r, c - 36, f_, col, fmt, align="center" if c in (46, 47) else None)
put(cd, r, 12, f"={qdepl(f'Z{r}', f'AA{r}')}", GREEN, align="center")
for k in range(1, NBURN + 1):
    r = CDH + 1 + k
    put(cd, r, 25, f'=IFERROR(IF(SMALL({VR("AW")},{k})>=99,"",MATCH(SMALL({VR("AW")},{k}),{VR("AW")},0)),"")', GRAY, "0")
    put(cd, r, 1, f'=IF($Y{r}="","",{k})', align="center")
    for c, src, fmt in ((2, "B", None), (3, "D", None), (4, "P", NUM), (5, "Q", NUM), (6, "AP", NUM), (7, "AQ", NUM), (8, "AR", NUM), (9, "AS", "0.0"),
                        (10, "AT", None), (11, "AU", None), (12, "AV", None), (26, "AJ", NUM), (27, "AK", NUM), (28, "AZ", NUM)):
        put(cd, r, c, f'=IF($Y{r}="","",INDEX({VR(src)},$Y{r}))', GREEN, fmt, align="center" if c in (10, 11, 12) else None)
    for j in range(12):
        put(cd, r, 13 + j, f'=IF($Y{r}="","",INDEX(월별집계!${CL(28 + j)}$7:${CL(28 + j)}${VA_END},$Y{r}))', GREEN, NUM)
widths(cd, [12, 14, 14, 14, 14, 14, 14, 11, 9, 11, 10, 12] + [8] * 12 + [8, 10, 10])
CD_LAYOUT = dict(CDA=CDA, CDB=CDB, CDC=CDC, CDD=CDD, CDE=CDE, CDF=CDF, CDG=CDG, CDH=CDH, NBURN=NBURN, NDEPT=len(DL), NGL=len(sgl), NKPI=len(KPI))

# ---------------- 대시보드 (엑셀)
db = wb.create_sheet("대시보드")
db["A1"] = f'={S["co"]}&" TRACE 대시보드 ("&{plabel(S["P"])}&")"'; db["A1"].font = F(bold=True, size=16)
db["A2"] = f'="단위: "&{S["ur"]}&" · 자가진단: "&자가진단!$B$4&" · 3자 대사: "&계정변동!$P${AT}'; db["A2"].font = F(GRAY, size=9)
tiles = [(0, "총자산"), (1, "자본총계"), (2, "CSM"), (3, "당기순이익 누계"), (4, "사업비 누계"), (6, "분기 한도 초과(예상)")]
for k, (ki, lab) in enumerate(tiles):
    c0 = 1 + k * 2
    r = CDG + 1 + ki
    db.merge_cells(start_row=4, start_column=c0, end_row=4, end_column=c0 + 1)
    db.merge_cells(start_row=5, start_column=c0, end_row=5, end_column=c0 + 1)
    db.merge_cells(start_row=6, start_column=c0, end_row=6, end_column=c0 + 1)
    c = db.cell(row=4, column=c0, value=lab); c.font = F(GRAY, size=9); c.alignment = Alignment(horizontal="center"); c.fill = TOT
    c = db.cell(row=5, column=c0, value=f"=차트데이터!B{r}"); c.font = F(GREEN, bold=True, size=14); c.number_format = "#,##0" if ki < 6 else '0"건"'
    c.alignment = Alignment(horizontal="center"); c.fill = TOT
    sub = (f'="주의 "&차트데이터!C{r}&"건"' if ki == 6 else f'=차트데이터!E{r}&" 대비 "&TEXT(차트데이터!D{r},"+0.0%;-0.0%;0.0%")')
    c = db.cell(row=6, column=c0, value=sub); c.font = F(GREEN, size=9); c.alignment = Alignment(horizontal="center"); c.fill = TOT
for c in range(1, 13): db.column_dimensions[CL(c)].width = 10.5

def place(ch, anchor, w=15.5, h=7.2):
    ch.width = w; ch.height = h; ch.style = 10; db.add_chart(ch, anchor)
c1 = BarChart(); c1.type = "col"; c1.title = "월별 사업비 (당해 vs 전년)"; c1.y_axis.numFmt = "#,##0"
c1.add_data(Reference(cd, min_col=2, max_col=3, min_row=CDA, max_row=CDA + 12), titles_from_data=True)
c1.set_categories(Reference(cd, min_col=1, min_row=CDA + 1, max_row=CDA + 12)); place(c1, "A8")
c2 = BarChart(); c2.type = "bar"; c2.title = "부서별 분기 한도 vs 분기말 예상"; c2.x_axis.numFmt = "#,##0"
c2.add_data(Reference(cd, min_col=3, max_col=3, min_row=CDB, max_row=CDB + len(DL)), titles_from_data=True)
c2.add_data(Reference(cd, min_col=5, max_col=5, min_row=CDB, max_row=CDB + len(DL)), titles_from_data=True)
c2.set_categories(Reference(cd, min_col=2, min_row=CDB + 1, max_row=CDB + len(DL))); place(c2, "G8")
c3 = BarChart(); c3.type = "bar"; c3.title = "계정별 당월 vs 전년동월"
c3.add_data(Reference(cd, min_col=3, max_col=4, min_row=CDD, max_row=CDD + len(sgl)), titles_from_data=True)
c3.set_categories(Reference(cd, min_col=2, min_row=CDD + 1, max_row=CDD + len(sgl))); place(c3, "A23", h=9)
c4 = DoughnutChart(); c4.title = "분기 한도 판정 분포 (부서×계정)"
c4.add_data(Reference(cd, min_col=2, min_row=CDC, max_row=CDC + 5), titles_from_data=True)
c4.set_categories(Reference(cd, min_col=1, min_row=CDC + 1, max_row=CDC + 5)); place(c4, "G23", h=9)
c5 = LineChart(); c5.title = "재무상태 추이 (최근 12개월)"; c5.y_axis.numFmt = "#,##0"
c5.add_data(Reference(cd, min_col=3, max_col=5, min_row=CDE, max_row=CDE + 12), titles_from_data=True)
c5.set_categories(Reference(cd, min_col=2, min_row=CDE + 1, max_row=CDE + 12)); place(c5, "A42")
c6 = BarChart(); c6.type = "col"; c6.title = "월별 당기순이익 (당해 vs 전년)"; c6.y_axis.numFmt = "#,##0"
c6.add_data(Reference(cd, min_col=2, max_col=3, min_row=CDF, max_row=CDF + 12), titles_from_data=True)
c6.set_categories(Reference(cd, min_col=1, min_row=CDF + 1, max_row=CDF + 12)); place(c6, "G42")
sec(db, 58, "주의 대상 TOP 10")
header(db, 59, ["순위", "부서", "계정", "당월", "전월대비", "예산", "신호", "", "", "", "", ""], h=20)
db.merge_cells("G59:L59")
for k in range(10):
    r = 60 + k; wr = 6 + k
    for c, src, fm in ((1, "A", None), (2, "B", None), (3, "C", None), (4, "D", NUM), (5, "E", RATE), (6, "G", None), (7, "J", None)):
        put(db, r, c, f"=주의대상!{src}{wr}", GREEN, fm, align="center" if c in (1, 6) else None)
    for c in range(8, 13): put(db, r, c, None)
    db.merge_cells(start_row=r, start_column=7, end_row=r, end_column=12)
for val, fill in (('"한도 초과"', "F8CBAD"), ('"초과 예상"', "FCE4D6"), ('"부족 예상"', "F8CBAD"), ('"주의"', "FFE699"), ('"무예산"', "D9C3E9")):
    db.conditional_formatting.add("F60:F69", CellIsRule(operator="equal", formula=[val], fill=PatternFill("solid", fgColor=fill)))
sec(db, 71, "선택 항목 원인 (원인분석 시트)")
db.merge_cells("A72:L73"); db["A72"] = "=원인분석!A18"; db["A72"].font = F(GREEN, bold=True)
db["A72"].alignment = Alignment(wrap_text=True, vertical="top"); db["A72"].fill = BOXF
db.sheet_view.showGridLines = False

# ================= v0.4: IFRS17(GMM) 사업비 연결 시트 =================
PCT1 = '0.0%;"△"0.0%;"-"'
# ---------- 입력옵션: 변동요인 매핑 (감독원 양식 요인코드 → TRACE 요인코드)
FSS_FAC = {"F00": "B00", "F10": "B10", "F21": "B21", "F22": "B22", "F31": "B31", "F32": "B32", "F40": "B40", "F51": "B51", "F52": "B52", "F60": "B60", "F99": "B99"}
FSS_FACN = {"F00": "기초", "F10": "신계약", "F21": "추정변경(사업비)", "F22": "추정변경(기타)", "F31": "경험조정(보험취득현금흐름)", "F32": "경험조정(기타)",
            "F40": "보험금융비용", "F51": "보험계약마진 상각", "F52": "위험조정 변동(당기)", "F60": "현금흐름 및 기타", "F99": "기말"}
FMAP_TOP = MAP1 + 4
FMAP0 = FMAP_TOP + 1; FMAP1 = FMAP0 + len(FAC) - 1
sec(os_, FMAP_TOP - 1, "부채변동표 변동요인 매핑 — 감독원 양식의 요인코드(없으면 요인명)를 TRACE 요인코드에 연결")
header(os_, FMAP_TOP, ["구분", "TRACE 요인코드", "TRACE 요인명", "실제 요인코드", "실제 요인명"], start=2, h=20)
for i, (code, nm) in enumerate(FAC):
    r = FMAP0 + i
    put(os_, r, 2, "변동표", align="center"); put(os_, r, 3, code, align="center"); put(os_, r, 4, nm)
    c = put(os_, r, 5, FSS_FAC[code], BLUE, fill=YEL, align="center"); c.number_format = "@"
    put(os_, r, 6, FSS_FACN[code], BLUE)
FMC = lambda col: f"입력옵션!${col}${FMAP0}:${col}${FMAP1}"

# ---------- 계정대체기준 (입력)
ab = wb.create_sheet("계정대체기준")
title(ab, "계정대체기준 (IFRS 4·9 → IFRS 17)", "① 직접대체 계정은 계정매핑 시트에서 지정 · ② 투자관련부서의 유지비 성격 공통비 → 재산관리비 · ③ 남은 공통비 → 부서별 비율")
sec(ab, 4, "1. 재산관리비 대체 (투자관련부서가 쓴 유지비 성격 공통비)")
header(ab, 5, ["투자관련부서", "부서명"], h=20); header(ab, 5, ["적용연도", "재산관리비 비율"], start=4, h=20)
for i in range(10):
    r = 6 + i
    d_ = INV_DEPTS[i] if i < len(INV_DEPTS) else None
    put(ab, r, 1, d_, BLUE, align="center"); put(ab, r, 2, DEPTS.get(d_) if d_ else None, BLUE)
    yr_ = sorted(INV_RATE)[i] if i < len(INV_RATE) else None
    put(ab, r, 4, yr_, BLUE, "0", align="center"); put(ab, r, 5, INV_RATE.get(yr_) if yr_ else None, BLUE, PCT, fill=YEL if yr_ else None)
ab.cell(row=16, column=1, value="비율은 1년에 한 번 바뀝니다. 새 연도는 아래 행에 추가하세요. 전표의 회계연도에 맞는 비율을 씁니다.").font = F(GRAY, size=9)
sec(ab, 18, "2. 부서별 공통비 대체비율 (재산관리비로 넘긴 뒤 남은 공통비)")
header(ab, 19, ["부서코드", "부서명", "신계약비", "유지비", "수금비", "비귀속", "합계"], h=20)
for i in range(30):
    r = 20 + i
    if i < len(DEPT_ALLOC):
        d_, rt = list(DEPT_ALLOC.items())[i]
        put(ab, r, 1, d_, BLUE, align="center"); put(ab, r, 2, DEPTS[d_], BLUE)
        for j, v in enumerate(rt): put(ab, r, 3 + j, v, BLUE, PCT, fill=YEL)
    else:
        for c in range(1, 7): put(ab, r, c, None, BLUE, PCT if c > 2 else None)
    put(ab, r, 7, f'=IF($A{r}="",0,SUM(C{r}:F{r}))', fmt=PCT)
ab.cell(row=50, column=1, value="표에 없는 부서의 공통비는 전액 비귀속으로 처리하고 자가진단이 알려줍니다.").font = F(GRAY, size=9)
sec(ab, 52, "3. 상품·계약집단 귀속 드라이버 (상품·증권번호별 귀속 결과를 받기 전까지 포트폴리오 비율로 대신)")
header(ab, 53, ["포트폴리오", "포트폴리오명", "측정모형", "신계약비\n(신계약 건수)", "유지비\n(보유계약)", "수금비\n(수입보험료)", "비고", "보험료 연동\n비율(신계약비)", "손실부담\n(Y/N)"], h=32)
for i in range(10):
    r = 54 + i
    if i < len(PFS):
        pf, pfn, mdl = PFS[i]
        put(ab, r, 1, pf, BLUE, align="center"); put(ab, r, 2, pfn, BLUE); put(ab, r, 3, mdl, BLUE, align="center")
        for j, v in enumerate(PF_DRV[pf]): put(ab, r, 4 + j, v, BLUE, PCT, fill=YEL)
        put(ab, r, 7, "변액 · 특별계정 (CSM 조정 · 손익 영향 분석 포함)" if mdl == "VFA" else ("CSM 없음 · 손익 반영" if mdl == "PAA" else "분석 대상"), GRAY)
        put(ab, r, 8, PF_LINK[pf], BLUE, PCT, fill=YEL); put(ab, r, 9, PF_ONEROUS[pf], BLUE, fill=YEL, align="center")
    else:
        for c in range(1, 10): put(ab, r, c, None, BLUE, PCT if (4 <= c <= 6 or c == 8) else None)
put(ab, 64, 3, "합계", bold=True, fill=TOT)
for c in (4, 5, 6): put(ab, 64, c, f"=SUM({CL(c)}54:{CL(c)}63)", fmt=PCT, bold=True, fill=TOT)
dv = DataValidation(type="list", formula1='"GMM,PAA,VFA"'); ab.add_data_validation(dv); dv.add("C54:C63")
dv = DataValidation(type="list", formula1='"Y,N"'); ab.add_data_validation(dv); dv.add("I54:I63")
sec(ab, 66, "4. IFRS 17 반영 정책 (회사 회계정책에 맞춰 고름 · K-IFRS 1117)")
for rr, k, v, opts, d in [(67, "신계약비 예실차 반영처", NB_POLICY, "보험료연동분만 CSM,전액 CSM,전액 당기손익",
                           "B96(a): 당기 보험료 중 미래 서비스 관련분과 연동된 취득현금흐름(예: 보험료 비례 수수료)의 예실차만 CSM. 나머지 경험조정은 B97(c)로 당기손익. 비율은 3번 표 H열"),
                          (68, "PAA 보험취득현금흐름", PAA_POLICY, "발생 즉시 비용,이연 후 상각",
                           "59(a): 보장기간 1년 이하면 발생 즉시 비용 선택 가능. 이연이면 잔여보장부채에서 상각(경험차이도 당기에 바로 손익이 되지 않음)")]:
    put(ab, rr, 1, k, bold=True, fill=SUB); ab.merge_cells(start_row=rr, start_column=1, end_row=rr, end_column=2)
    put(ab, rr, 3, v, BLUE, fill=YEL, align="center"); ab.merge_cells(start_row=rr, start_column=3, end_row=rr, end_column=4)
    put(ab, rr, 5, d, GRAY); ab.merge_cells(start_row=rr, start_column=5, end_row=rr, end_column=9)
    dv = DataValidation(type="list", formula1=f'"{opts}"'); ab.add_data_validation(dv); dv.add(f"C{rr}")
for k_, ln in enumerate(["· 당기 신계약의 취득현금흐름은 경험조정이 아니라 최초 측정(38항, 신계약 효과 F10)에 들어갑니다. 변동표 F31과 대사할 때 회사 변동표의 정의를 확인하세요.",
                         "· 손실부담 집단(Y)은 미래 서비스 관련 변동도 CSM이 아닌 손실요소·당기손익으로 갑니다 (44(c), 48·50항).",
                         "· 집단 인식 전에 지급한 취득현금흐름은 '보험취득현금흐름 자산'(28B~28F)으로 따로 관리하며 TRACE 범위 밖입니다.",
                         "· VFA의 투자 관련 서비스 원가, GMM의 투자수익 서비스 원가는 이행현금흐름에 포함될 수 있습니다 (B65(ka)). 1번 표 재산관리비 대체 대상인지 확인하세요."]):
    ab.cell(row=70 + k_, column=1, value=ln).font = F(GRAY, size=9)
widths(ab, [13, 18, 10, 12, 14, 12, 34, 12, 9]); ab.freeze_panes = "A4"

# ---------- 예정사업비 (입력)
EX_HDR = ["회계연도", "기간", "포트폴리오", "포트폴리오명", "신계약비", "유지비", "수금비"]
EXE = 20000; EXR = lambda c: f"예정사업비!${c}$5:${c}${EXE}"
ex_ = wb.create_sheet("예정사업비")
title(ex_, "예정사업비 (계리 · BEL 예상현금흐름의 사업비)", "A~G: 붙여넣기 · I~L: 보조수식 · 포트폴리오 × 월 · 금액은 원본 단위")
header(ex_, 4, EX_HDR); header(ex_, 4, ["기간번호", "신계약비\n(보고단위)", "유지비\n(보고단위)", "수금비\n(보고단위)"], start=9, fill=GRAYH)
for k, row in enumerate(EXPROWS):
    r = 5 + k
    for j, v in enumerate(row):
        c = ex_.cell(row=r, column=j + 1, value=v); c.font = F(BLUE)
        if j >= 4: c.number_format = "#,##0"
    ex_.cell(row=r, column=9, value=f'=IF($A{r}="","",VALUE($A{r}&"")*12+VALUE($B{r}&""))').font = F()
    for j, col in enumerate("EFG"):
        c = ex_.cell(row=r, column=10 + j, value=f'=IF($A{r}="",0,${col}{r}*{S["fx"]})'); c.font = F(GREEN); c.number_format = NUM
ex_.freeze_panes = "A5"; widths(ex_, [8, 5, 10, 12, 14, 14, 14, 2, 8, 11, 11, 11])

# ---------- 부채변동표 (입력)
RF_HDR = ["회계연도", "기간", "포트폴리오", "포트폴리오명", "측정모형", "변동요인코드", "변동요인", "BEL", "RA", "CSM"]
RFE = 20000; RFR = lambda c: f"부채변동표!${c}$5:${c}${RFE}"
rf = wb.create_sheet("부채변동표")
title(rf, "보험계약부채 변동표 (BEL · RA · CSM)", "A~J: 붙여넣기 (감독원 양식·계리 결과) · L~P: 보조수식 · 요인코드는 입력옵션의 변동요인 매핑으로 TRACE 코드에 연결")
header(rf, 4, RF_HDR); header(rf, 4, ["기간번호", "BEL\n(보고단위)", "RA\n(보고단위)", "CSM\n(보고단위)", "TRACE\n요인코드"], start=12, fill=GRAYH)
for k, row in enumerate(ROLL):
    r = 5 + k
    for j, v in enumerate(row):
        c = rf.cell(row=r, column=j + 1, value=v); c.font = F(BLUE)
        if j >= 7: c.number_format = "#,##0"
    rf.cell(row=r, column=12, value=f'=IF($A{r}="","",VALUE($A{r}&"")*12+VALUE($B{r}&""))').font = F()
    for j, col in enumerate("HIJ"):
        c = rf.cell(row=r, column=13 + j, value=f'=IF($A{r}="",0,${col}{r}*{S["fx"]})'); c.font = F(GREEN); c.number_format = NUM
    rf.cell(row=r, column=16, value=(f'=IF($A{r}="","",IFERROR(INDEX({FMC("C")},MATCH($F{r}&"",{FMC("E")},0)),'
                                     f'IFERROR(INDEX({FMC("C")},MATCH($G{r}&"",{FMC("F")},0)),$F{r}&"")))')).font = F(GREEN)
rf.freeze_panes = "A5"; rf.auto_filter.ref = "A4:P4"; widths(rf, [8, 5, 10, 11, 8, 11, 22, 15, 13, 13, 2, 8, 12, 11, 11, 9])

# ---------- IFRS17_사업비 (분석)
ia = wb.create_sheet("IFRS17_사업비")
title(ia, "IFRS 17 사업비 영향 (GMM · VFA · PAA)", "사업비 전표 → 계정대체 → 포트폴리오 → 예정 대비 경험차이 → CSM(GMM·VFA)·당기손익 · 금액: 보고단위")
ia["A3"] = f'="기준: "&{plabel(S["P"])}&" · 단위: "&{S["ur"]}'; ia["A3"].font = F(GREEN, size=9)
SUMR = 5
ia.merge_cells(f"A{SUMR}:H{SUMR+1}")
cats = [("비귀속", "Y"), ("신계약비", "Z"), ("유지비", "AA"), ("수금비", "AB"), ("재산관리비", "AD")]
A0 = 9   # 1. 배분 요약
sec(ia, A0 - 1, "1. 계정대체 결과 요약 (전표 기준)")
header(ia, A0, ["구분", "당월", "전월", "누계", "당월 비중"], h=20)
for i, (nm, col) in enumerate(cats):
    r = A0 + 1 + i
    put(ia, r, 1, nm, bold=True)
    put(ia, r, 2, f'=SUMIFS({JR(col)},{JR("M")},{S["P"]})', GREEN, NUM)
    put(ia, r, 3, f'=SUMIFS({JR(col)},{JR("M")},{S["P"]}-1)', GREEN, NUM)
    put(ia, r, 4, f'=SUMIFS({JR(col)},{JR("M")},">="&({S["yr"]}*12+1),{JR("M")},"<="&{S["P"]})', GREEN, NUM)
    put(ia, r, 5, f'=IF($B${A0+6}=0,0,B{r}/$B${A0+6})', fmt=PCT1)
put(ia, A0 + 6, 1, "합계 (= 전표 사업비)", bold=True, fill=TOT)
for c in (2, 3, 4): put(ia, A0 + 6, c, f"=SUM({CL(c)}{A0+1}:{CL(c)}{A0+5})", fmt=NUM, bold=True, fill=TOT)
put(ia, A0 + 6, 5, f"=SUM(E{A0+1}:E{A0+5})", fmt=PCT1, bold=True, fill=TOT)
ia.cell(row=A0 + 7, column=1, value="재산관리비는 IFRS 17 이행현금흐름 밖 (PL 투자손익에서 차감) · 비귀속은 기타사업비용").font = F(GRAY, size=9)
B0 = A0 + 10  # 2. 포트폴리오별 경험차이
sec(ia, B0 - 1, "2. 포트폴리오별 경험차이 (당월, 예정 − 실제 · 음수 = 불리)")
header(ia, B0, ["포트폴리오", "포트폴리오명", "측정모형", "구분", "실제", "예정", "경험차이", "반영처", "CSM 반영", "당기손익 반영"], h=20)
NBP, PAAP = "계정대체기준!$C$67", "계정대체기준!$C$68"
PFR = lambda col: f"계정대체기준!${col}$54:${col}$63"
rows_b = []
for i, (pf, pfn, mdl) in enumerate(PFS):
    for j, (cat, drvcol, excol) in enumerate((("신계약비", "D", "J"), ("유지비", "E", "K"), ("수금비", "F", "L"))):
        r = B0 + 1 + i * 3 + j; rows_b.append(r)
        put(ia, r, 1, pf, BLUE, align="center")
        put(ia, r, 2, f'=IFERROR(INDEX({PFR("B")},MATCH($A{r},{PFR("A")},0)),"")', GREEN)
        put(ia, r, 3, f'=IFERROR(INDEX({PFR("C")},MATCH($A{r},{PFR("A")},0)),"")', GREEN, align="center")
        put(ia, r, 4, cat, align="center")
        put(ia, r, 5, f'=$B${A0+2+j}*IFERROR(INDEX({PFR(drvcol)},MATCH($A{r},{PFR("A")},0)),0)', GREEN, NUM)
        put(ia, r, 6, f'=SUMIFS({EXR(excol)},{EXR("I")},{S["P"]},{EXR("C")},$A{r})', GREEN, NUM)
        put(ia, r, 7, f"=F{r}-E{r}", fmt='+#,##0.0;-#,##0.0;"-"', bold=True)
        lk = f'IF({NBP}="전액 CSM",1,IF({NBP}="전액 당기손익",0,IFERROR(INDEX({PFR("H")},MATCH($A{r},{PFR("A")},0)),0)))'
        on = f'(IFERROR(INDEX({PFR("I")},MATCH($A{r},{PFR("A")},0)),"N")="Y")'
        csm_ok = f'AND(D{r}="신계약비",OR(C{r}="GMM",C{r}="VFA"),NOT({on}))'
        put(ia, r, 9, f'=IF({csm_ok},G{r}*{lk},0)', fmt='+#,##0.0;-#,##0.0;"-"')
        put(ia, r, 10, f'=IF(AND(D{r}="신계약비",C{r}="PAA",{PAAP}="이연 후 상각"),0,G{r}-I{r})', fmt='+#,##0.0;-#,##0.0;"-"')
        put(ia, r, 8, f'=IF(D{r}<>"신계약비","당기손익(보험서비스결과)",IF(C{r}="PAA",IF({PAAP}="이연 후 상각","잔여보장부채 이연(상각)","당기손익"),'
                      f'IF({on},"당기손익(손실요소)",IF({lk}>=1,"CSM 조정 ("&C{r}&")",IF({lk}<=0,"당기손익","CSM "&TEXT({lk},"0%")&" ("&C{r}&") · 나머지 손익")))))', align="center")
B1 = rows_b[-1]
BT = B1 + 1
for k, (lab, f_) in enumerate((("CSM 조정 (GMM 취득현금흐름 예실차)", f'=SUMIFS($I${B0+1}:$I${B1},$C${B0+1}:$C${B1},"GMM")'),
                                ("CSM 조정 (VFA 취득현금흐름 예실차)", f'=SUMIFS($I${B0+1}:$I${B1},$C${B0+1}:$C${B1},"VFA")'),
                                ("당기손익 (유지·수금, 연동 외 신계약비, PAA)", f"=SUM($J${B0+1}:$J${B1})"))):
    r = BT + k
    put(ia, r, 1, lab, bold=True, fill=TOT); ia.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    put(ia, r, 7, f_, fmt='+#,##0.0;-#,##0.0;"-"', bold=True, fill=TOT)
    for c in (8, 9, 10): put(ia, r, c, None, fill=TOT)
ia.conditional_formatting.add(f"H{B0+1}:H{B1}", FormulaRule(formula=[f'LEFT(H{B0+1},3)="CSM"'], fill=PatternFill("solid", fgColor="E3EBF7")))
ia.conditional_formatting.add(f"G{B0+1}:G{B1}", CellIsRule(operator="lessThan", formula=["0"], font=Font(color="C00000", bold=True)))
C0_ = BT + 6  # 3. 변동표 대사
sec(ia, C0_ - 1, "3. 변동표 대사 (계리 결과 '경험조정-취득현금흐름' CSM vs TRACE CSM 반영분)")
header(ia, C0_, ["포트폴리오", "측정모형", "계리 F31 (CSM)", "TRACE 추정\n(CSM 반영분)", "차이", "판정"], h=30)
for i, (pf, pfn, mdl) in enumerate(PFS):
    r = C0_ + 1 + i
    put(ia, r, 1, pf, BLUE, align="center")
    put(ia, r, 2, f'=IFERROR(INDEX({PFR("C")},MATCH($A{r},{PFR("A")},0)),"")', GREEN, align="center")
    put(ia, r, 3, f'=SUMIFS({RFR("O")},{RFR("L")},{S["P"]},{RFR("C")},$A{r},{RFR("P")},"F31")', GREEN, NUM)
    put(ia, r, 4, f'=IF(B{r}="PAA",0,SUMIFS($I${B0+1}:$I${B1},$A${B0+1}:$A${B1},$A{r},$D${B0+1}:$D${B1},"신계약비"))', fmt=NUM)
    put(ia, r, 5, f"=C{r}-D{r}", fmt=NUM)
    put(ia, r, 6, f'=IF(B{r}="PAA","해당없음 (CSM 없음)",IF(ABS(E{r})<0.5,"일치","차이 확인"))', align="center")
C1_ = C0_ + len(PFS)
ia.conditional_formatting.add(f"F{C0_+1}:F{C1_}", CellIsRule(operator="equal", formula=['"차이 확인"'], fill=PatternFill("solid", fgColor="FFC7CE")))
def bridge(d0, mdl, label):
    sec(ia, d0 - 1, label)
    header(ia, d0, ["요인코드", "변동요인", "BEL", "RA", "CSM", "사업비 관련"], h=20)
    for i, (code, nm) in enumerate(FAC):
        r = d0 + 1 + i
        bold_ = code in ("F00", "F99")
        put(ia, r, 1, code, align="center", bold=bold_, fill=TOT if bold_ else None)
        put(ia, r, 2, nm, bold=bold_, fill=TOT if bold_ else None)
        for c, col in ((3, "M"), (4, "N"), (5, "O")):
            put(ia, r, c, f'=SUMIFS({RFR(col)},{RFR("L")},{S["P"]},{RFR("P")},$A{r},{RFR("E")},"{mdl}")', GREEN, NUM0, bold=bold_, fill=TOT if bold_ else None)
        put(ia, r, 6, "사업비" if code in ("F21", "F31") else None, align="center", bold=True)
    d1 = d0 + len(FAC)
    r21, r31 = d0 + 1 + [f[0] for f in FAC].index("F21"), d0 + 1 + [f[0] for f in FAC].index("F31")
    for k, (lab, f_) in enumerate((("CSM 당월 변동 (기말 − 기초)", f"=E{d1}-E{d0+1}"), ("그중 사업비 관련 (F21 + F31)", f"=E{r21}+E{r31}"))):
        r = d1 + 1 + k
        put(ia, r, 2, lab, bold=True, fill=SUB); put(ia, r, 5, f_, fmt='+#,##0;-#,##0;"-"', bold=True, fill=SUB)
    ia.conditional_formatting.add(f"F{d0+1}:F{d1}", CellIsRule(operator="equal", formula=['"사업비"'], fill=PatternFill("solid", fgColor="E3EBF7")))
    return d1, r21, r31
D0 = C1_ + 3  # 4. CSM 브리지 (GMM)
D1, R21, R31 = bridge(D0, "GMM", "4. GMM 변동 브리지 (보장성 + 저축성, 당월)")
DV0 = D1 + 6  # 4-2. VFA 브리지
DV1, RV21, RV31 = bridge(DV0, "VFA", "4-2. VFA 변동 브리지 (변액, 당월)")
D1_, D0_ = D1, D0
E0 = DV1 + 5  # 5. 신계약비 증가 기여 (부서 × 계정)
sec(ia, E0 - 1, "5. 신계약비 증가 기여 (부서 × 계정, 당월 − 전월, 귀속 신계약비 기준)")
header(ia, E0, ["순위", "부서", "계정", "당월", "전월", "증감"], h=20)
header(ia, 6, ["키", "당월 신계약", "전월 신계약", "증감", "정렬키"], start=12, fill=GRAYH, h=20)
for idx in range(VA_PRE - 7):
    r = 7 + idx
    put(ia, r, 12, f'=IF(변동분석!$A{r}="","",변동분석!$E{r})', GREEN)
    put(ia, r, 13, f'=IF(L{r}="",0,SUMIFS({JR("Z")},{JR("O")},L{r},{JR("M")},{S["P"]}))', GREEN, NUM)
    put(ia, r, 14, f'=IF(L{r}="",0,SUMIFS({JR("Z")},{JR("O")},L{r},{JR("M")},{S["P"]}-1))', GREEN, NUM)
    put(ia, r, 15, f"=M{r}-N{r}", fmt=NUM)
    put(ia, r, 16, f'=IF(L{r}="","",O{r}+ROW()/1000000)', fmt="0.000")
HL_ = 7 + VA_PRE - 8
for k in range(1, 6):
    r = E0 + k
    put(ia, r, 1, k, align="center")
    put(ia, r, 11, f'=IFERROR(MATCH(LARGE($P$7:$P${HL_},{k}),$P$7:$P${HL_},0),"")', GRAY)
    put(ia, r, 2, f'=IF($K{r}="","",INDEX(변동분석!$B$7:$B${HL_},$K{r}))', GREEN)
    put(ia, r, 3, f'=IF($K{r}="","",INDEX(변동분석!$D$7:$D${HL_},$K{r}))', GREEN)
    for c, col in ((4, "M"), (5, "N"), (6, "O")):
        put(ia, r, c, f'=IF($K{r}="","",INDEX(${col}$7:${col}${HL_},$K{r}))', fmt='+#,##0.0;-#,##0.0;"-"' if c == 6 else NUM)
F0 = E0 + 8  # 6. 근거 전표
sec(ia, F0 - 1, "6. 근거 전표 (당월, 귀속 신계약비 큰 순)")
header(ia, F0, ["전기일", "전표번호", "코스트센터", "G/L 계정", "거래처", "적요", "금액", "신계약 귀속"], h=20)
for k in range(1, 11):
    r = F0 + k
    put(ia, r, 11, f'=IFERROR(MATCH(LARGE({JR("AC")},{k}),{JR("AC")},0),"")', GRAY)
    for c, src, fmt in ((1, "D", "yyyy-mm-dd"), (2, "C", None), (3, "G", None), (4, "F", None), (5, "K", None), (6, "I", None), (7, "P", NUM), (8, "Z", NUM)):
        put(ia, r, c, f'=IF($K{r}="","",INDEX({JR(src)},$K{r}))', GREEN, fmt, align="center" if c in (1, 3, 4) else None)
nb_tot = f"SUMIFS($G${B0+1}:$G${B1},$D${B0+1}:$D${B1},\"신계약비\")"
ia[f"A{SUMR}"] = (f'="GMM CSM 당월 변동 "&TEXT(E{D1+1},"+#,##0;-#,##0;0")&" 중 사업비 관련 "&TEXT(E{D1+2},"+#,##0.0;-#,##0.0;0")'
                  f'&" (취득현금흐름 경험조정 "&TEXT(E{R31},"+#,##0.0;-#,##0.0;0")&", 사업비 가정변경 "&TEXT(E{R21},"+#,##0.0;-#,##0.0;0")&"). "'
                  f'&"신계약비 실제가 예정보다 "&TEXT(ABS({nb_tot}),"#,##0.0")&IF({nb_tot}<0," 많았고"," 적었고")'
                  f'&IF(B{E0+1}="","",", 증가 기여 1위는 "&B{E0+1}&" "&C{E0+1}&" ("&TEXT(F{E0+1},"+#,##0.0;-#,##0.0;0")&")")'
                  f'&IF(B{F0+1}="","",", 대표 전표는 "&E{F0+1}&" \'"&F{F0+1}&"\' "&TEXT(H{F0+1},"#,##0.0"))&". "'
                  f'&"당기손익 반영 경험차이 "&TEXT(G{BT+2},"+#,##0.0;-#,##0.0;0")&". "'
                  f'&"VFA(변액) CSM 당월 변동 "&TEXT(E{DV1+1},"+#,##0;-#,##0;0")&" 중 사업비 관련 "&TEXT(E{DV1+2},"+#,##0.0;-#,##0.0;0")&"."')
ia[f"A{SUMR}"].font = F(GREEN, bold=True); ia[f"A{SUMR}"].alignment = Alignment(wrap_text=True, vertical="top"); ia[f"A{SUMR}"].fill = BOXF
ia.row_dimensions[SUMR].height = 30; ia.row_dimensions[SUMR + 1].height = 30
widths(ia, [12, 26, 12, 14, 14, 30, 14, 22, 2, 2, 6, 14, 11, 11, 11, 9])
ia.freeze_panes = "A8"
IF_LAYOUT = dict(SUMR=SUMR, A0=A0, B0=B0, B1=B1, BT=BT, C0=C0_, C1=C1_, D0=D0, D1=D1, DV0=DV0, DV1=DV1, E0=E0, F0=F0)

# 대시보드: GMM CSM 변동 요인 차트 + 요약
sec(db, 75, "IFRS 17 사업비 영향 (GMM · VFA)")
db.merge_cells("A76:L77"); db["A76"] = f"=IFRS17_사업비!A{SUMR}"; db["A76"].font = F(GREEN, bold=True)
db["A76"].alignment = Alignment(wrap_text=True, vertical="top"); db["A76"].fill = BOXF
c7 = BarChart(); c7.type = "bar"; c7.title = "GMM CSM 변동 요인 (당월)"; c7.x_axis.numFmt = "#,##0"
c7.add_data(Reference(ia, min_col=5, max_col=5, min_row=D0, max_row=D1 - 1), titles_from_data=True)
c7.set_categories(Reference(ia, min_col=2, min_row=D0 + 1, max_row=D1 - 1)); place(c7, "A79", w=31, h=8)

IFRS17_CHECKS = [
    ("대체비율 합계 100%", f'="부서 "&SUMPRODUCT((계정대체기준!$A$20:$A$49<>"")*(ABS(계정대체기준!$G$20:$G$49-1)>0.0001))&"행 · 포트폴리오 "&(ABS(계정대체기준!$D$64-1)+ABS(계정대체기준!$E$64-1)+ABS(계정대체기준!$F$64-1)>0.0001)*1',
     'IF(SUMPRODUCT((계정대체기준!$A$20:$A$49<>"")*(ABS(계정대체기준!$G$20:$G$49-1)>0.0001))+(ABS(계정대체기준!$D$64-1)+ABS(계정대체기준!$E$64-1)+ABS(계정대체기준!$F$64-1)>0.0001)=0,"OK","확인필요")',
     "계정대체기준 시트에서 합계가 100%가 아닌 행 수정"),
    ("대체비율표에 없는 부서 (공통비)", f'=SUMPRODUCT(({JR("AE")}="공통비")*(COUNTIF(계정대체기준!$A$20:$A$49,{JR("N")})=0))&"건 (전액 비귀속 처리)"',
     f'IF(SUMPRODUCT(({JR("AE")}="공통비")*(COUNTIF(계정대체기준!$A$20:$A$49,{JR("N")})=0))=0,"OK","확인필요")', "계정대체기준 부서별 비율표에 부서 추가"),
    ("계정대체 전후 합계", f'="차이 "&TEXT(SUMIFS({JR("P")},{JR("M")},{S["P"]})-SUMIFS({JR("Y")},{JR("M")},{S["P"]})-SUMIFS({JR("Z")},{JR("M")},{S["P"]})-SUMIFS({JR("AA")},{JR("M")},{S["P"]})-SUMIFS({JR("AB")},{JR("M")},{S["P"]})-SUMIFS({JR("AD")},{JR("M")},{S["P"]}),"#,##0.0")',
     f'IF(ABS(SUMIFS({JR("P")},{JR("M")},{S["P"]})-SUMIFS({JR("Y")},{JR("M")},{S["P"]})-SUMIFS({JR("Z")},{JR("M")},{S["P"]})-SUMIFS({JR("AA")},{JR("M")},{S["P"]})-SUMIFS({JR("AB")},{JR("M")},{S["P"]})-SUMIFS({JR("AD")},{JR("M")},{S["P"]}))<0.05,"OK","확인필요")',
     "직접대체 계정 지정, 부서별 비율 합계 확인"),
    ("PL 재산관리비 ↔ 재산관리비 대체액", f'="차이 "&TEXT({plv("P230", S["P"])}-SUMIFS({JR("AD")},{JR("M")},{S["P"]}),"#,##0.0")',
     f'IF(ABS({plv("P230", S["P"])}-SUMIFS({JR("AD")},{JR("M")},{S["P"]}))<0.5,"OK","확인필요")', "투자관련부서 목록·재산관리비 비율, PL 재산관리비 범위 확인"),
    ("기준월 예정사업비 존재", f'=COUNTIF({EXR("I")},{S["P"]})&"행"', f'IF(COUNTIF({EXR("I")},{S["P"]})>0,"OK","확인필요")', "계리 예정사업비 기준월 자료 확인"),
    ("부채변동표 롤포워드", f'="차이 "&TEXT(SUMIFS({RFR("M")},{RFR("L")},{S["P"]},{RFR("P")},"<>F99")-SUMIFS({RFR("M")},{RFR("L")},{S["P"]},{RFR("P")},"F99"),"#,##0.0")&" / "&TEXT(SUMIFS({RFR("O")},{RFR("L")},{S["P"]},{RFR("P")},"<>F99")-SUMIFS({RFR("O")},{RFR("L")},{S["P"]},{RFR("P")},"F99"),"#,##0.0")',
     f'IF(ABS(SUMIFS({RFR("M")},{RFR("L")},{S["P"]},{RFR("P")},"<>F99")-SUMIFS({RFR("M")},{RFR("L")},{S["P"]},{RFR("P")},"F99"))+ABS(SUMIFS({RFR("N")},{RFR("L")},{S["P"]},{RFR("P")},"<>F99")-SUMIFS({RFR("N")},{RFR("L")},{S["P"]},{RFR("P")},"F99"))+ABS(SUMIFS({RFR("O")},{RFR("L")},{S["P"]},{RFR("P")},"<>F99")-SUMIFS({RFR("O")},{RFR("L")},{S["P"]},{RFR("P")},"F99"))<1,"OK","확인필요")',
     "기초 + 요인 합계 = 기말인지, 변동요인 매핑(입력옵션) 누락 확인"),
    ("변동표 기말 ↔ BS (BEL·RA·CSM)", f'="차이 "&TEXT(SUMIFS({RFR("M")},{RFR("L")},{S["P"]},{RFR("P")},"F99")-{bsv("L110", S["P"])},"#,##0")&" / "&TEXT(SUMIFS({RFR("N")},{RFR("L")},{S["P"]},{RFR("P")},"F99")-{bsv("L120", S["P"])},"#,##0")&" / "&TEXT(SUMIFS({RFR("O")},{RFR("L")},{S["P"]},{RFR("P")},"F99")-{bsv("L130", S["P"])},"#,##0")',
     f'IF(ABS(SUMIFS({RFR("M")},{RFR("L")},{S["P"]},{RFR("P")},"F99")-{bsv("L110", S["P"])})+ABS(SUMIFS({RFR("N")},{RFR("L")},{S["P"]},{RFR("P")},"F99")-{bsv("L120", S["P"])})+ABS(SUMIFS({RFR("O")},{RFR("L")},{S["P"]},{RFR("P")},"F99")-{bsv("L130", S["P"])})<1,"OK","확인필요")',
     "BS 항목 매핑(L110·L120·L130)과 변동표 포트폴리오 범위 확인"),
    ("변동표 F31 ↔ TRACE 신계약비 경험차이", f'=COUNTIF(IFRS17_사업비!$F${C0_+1}:$F${C1_},"차이 확인")&"개 포트폴리오 차이"',
     f'IF(COUNTIF(IFRS17_사업비!$F${C0_+1}:$F${C1_},"차이 확인")=0,"OK","참고")', "배분 드라이버·예정사업비 기준이 계리와 같은지 확인 (차이는 정보성)"),
]

# ---------------- 자가진단 추가 항목
EXTRA_CHECKS = [
    ("BS_원본 머리글", hdrchk("BS_원본", BS_HDR), 'IF(C{r}="일치","OK","확인필요")', "4행 머리글을 표준 이름으로"),
    ("PL_원본 머리글", hdrchk("PL_원본", PL_HDR), 'IF(C{r}="일치","OK","확인필요")', "4행 머리글을 표준 이름으로"),
    ("BS·PL 보조수식 범위", f'="BS "&{cnt("BS_원본","A",BSE)}&"/"&{cnt("BS_원본","I",BSE)}&" · PL "&{cnt("PL_원본","A",PLE)}&"/"&{cnt("PL_원본","J",PLE)}',
     f'IF(AND({cnt("BS_원본","I",BSE)}>={cnt("BS_원본","A",BSE)},{cnt("PL_원본","J",PLE)}>={cnt("PL_원본","A",PLE)}),"OK","확인필요")', "보조수식을 데이터 마지막 행까지 채우기"),
    ("기준월 BS·PL 존재", f'="BS "&COUNTIF({BSR("I")},{S["P"]})&"행 · PL "&COUNTIF({PLR("J")},{S["P"]})&"행"',
     f'IF(AND(COUNTIF({BSR("I")},{S["P"]})>0,COUNTIF({PLR("J")},{S["P"]})>0),"OK","확인필요")', "기준월 BS·PL이 들어 있는지 확인"),
    ("BS 대차 (자산 = 부채 + 자본)", f'="차이 "&TEXT({bsv("A000", S["P"])}-{bsv("L000", S["P"])}-{bsv("E000", S["P"])},"#,##0")',
     f'IF(ABS({bsv("A000", S["P"])}-{bsv("L000", S["P"])}-{bsv("E000", S["P"])})<1,"OK","확인필요")', "BS 합계 행(A000·L000·E000) 확인"),
    ("PL 당기순이익 검산", f'="차이 "&TEXT({plv("P100", S["P"])}+{plv("P200", S["P"])}-{plv("P300", S["P"])}-{plv("P500", S["P"])}-{plv("P600", S["P"])},"#,##0.0")',
     f'IF(ABS({plv("P100", S["P"])}+{plv("P200", S["P"])}-{plv("P300", S["P"])}-{plv("P500", S["P"])}-{plv("P600", S["P"])})<0.05,"OK","확인필요")', "보험손익+투자손익−사업비−법인세 = 당기순이익"),
    ("PL 기타사업비용 ↔ 비귀속 사업비", f'="차이 "&TEXT({plv("P300", S["P"])}-SUMIFS({JR("Y")},{JR("M")},{S["P"]}),"#,##0.0")',
     f'IF(ABS({plv("P300", S["P"])}-SUMIFS({JR("Y")},{JR("M")},{S["P"]}))<0.5,"OK","확인필요")', "직접대체·부서별 비율의 비귀속과 PL 기타사업비용 범위 확인"),
] + IFRS17_CHECKS + [
]

checks += EXTRA_CHECKS
C0 = 7
for k, (nm, res, stt, how) in enumerate(checks):
    r = C0 + k
    put(ck, r, 1, k + 1, align="center"); put(ck, r, 2, nm, bold=True); put(ck, r, 3, res, GREEN)
    put(ck, r, 4, "=" + stt.format(r=r), align="center", bold=True); put(ck, r, 5, how, GRAY)
C1 = C0 + len(checks) - 1
ck["A4"] = "종합"; ck["A4"].font = F(bold=True, size=12)
ck["B4"] = f'=IF(COUNTIF(D{C0}:D{C1},"확인필요")=0,"전체 정상"&IF(COUNTIF(D{C0}:D{C1},"참고")>0," (참고 "&COUNTIF(D{C0}:D{C1},"참고")&"건)",""),"확인필요 "&COUNTIF(D{C0}:D{C1},"확인필요")&"건")'
ck["B4"].font = F(bold=True, size=12)
for val, fill in (('"OK"', "C6EFCE"), ('"확인필요"', "FFC7CE"), ('"참고"', "FFEB9C")):
    ck.conditional_formatting.add(f"D{C0}:D{C1}", CellIsRule(operator="equal", formula=[val], fill=PatternFill("solid", fgColor=fill)))
widths(ck, [5, 28, 40, 10, 60])

# 대시보드 결과 내장 (EMBED_JSON=data.js 경로): 엑셀만 올려도 HTML 대시보드가 읽도록 숨김 시트에 30,000자씩 저장
if os.environ.get("EMBED_JSON") and os.path.exists(os.environ["EMBED_JSON"]):
    _js = open(os.environ["EMBED_JSON"], encoding="utf-8").read()
    _ws = wb.create_sheet("대시보드데이터")
    _ch = [_js[i:i + 30000] for i in range(0, len(_js), 30000)]
    _ws["A1"] = "TRACE_DATA"; _ws["B1"] = len(_ch); _ws["C1"] = "예시 데이터"
    for _i, _c in enumerate(_ch): _ws.cell(row=2 + _i, column=1, value=_c).number_format = "@"
    _ws.sheet_state = "hidden"
order = ["안내", "설정", "헤더설정", "입력옵션", "자가진단", "대시보드", "IFRS17_사업비", "리포트", "주의대상", "원인분석", "변동분석", "계정변동", "월별집계", "차트데이터", "코멘트", "키워드사전",
         "예산판_원본", "TB_원본", "전표_원본", "BS_원본", "PL_원본", "부채변동표", "예정사업비", "계정대체기준", "계정매핑", "부서매핑"]
if "대시보드데이터" in wb.sheetnames: order.append("대시보드데이터")
wb._sheets = [wb[n] for n in order]
for w in wb.worksheets: w.sheet_view.showGridLines = False
for n, colr in [("설정", "FFC000"), ("자가진단", "70AD47"), ("대시보드", "1F3864"), ("리포트", "1F3864"), ("BS_원본", "5B9BD5"), ("PL_원본", "5B9BD5"), ("부채변동표", "5B9BD5"), ("예정사업비", "5B9BD5"), ("IFRS17_사업비", "7030A0"), ("주의대상", "C00000"), ("원인분석", "C00000"),
                ("예산판_원본", "5B9BD5"), ("TB_원본", "5B9BD5"), ("전표_원본", "5B9BD5")]:
    wb[n].sheet_properties.tabColor = colr
wb.active = 3
wb.calculation.fullCalcOnLoad = True
wb.save(XLSX)

json.dump(dict(E=[{**e, "date": e["date"].isoformat()} for e in E], BP=BP, KEYS=KEYS, bud_gl={f"{k[0]}|{k[1]}": v for k, v in bud_gl.items()},
               VA_PRE=VA_PRE, AT=AT, C0=C0, C1=C1, JE_END=JE_END, CD=CD_LAYOUT, BS=BSROWS, PL=PLROWS, FM=FM, ROLL=ROLL, EXP=EXPROWS, IF=IF_LAYOUT), open(os.path.join(HERE, "tool_data.json"), "w"), ensure_ascii=False)
print("saved", XLSX, "entries", len(E), "TB", len(TB), "BP", len(BP), "keys", len(KEYS))
