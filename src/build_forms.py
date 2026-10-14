"""TRACE 입력 양식(빈 양식 + 예시 3행) 생성.

TRACE_v0.5.xlsx의 헤더설정 시트(열 이름·필수·대체 이름)와 각 원본 시트의 예시 행을 읽어
samples/TRACE_입력양식_v0.5.xlsx 를 만든다. SAP·감독원 자료를 이 머리글로 맞춰 내려받거나 붙여넣으면 된다.

사용: python src/build_forms.py   (build_trace.py 다음에 실행)
"""
import json, os
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
SRC_BOOK = os.path.join(ROOT, "TRACE_v0.5.xlsx")
OUT = os.path.join(ROOT, "samples", "TRACE_입력양식_v0.5.xlsx")

DESC = {  # 시트별 설명 (어디서 받는 자료인지)
    "예산판_원본": "SAP 예산판 화면 · 부서 × 계정 × 분기 (배정·조정·실적·잔여)",
    "TB_원본": "SAP 월별 시산표 · 당해와 전년을 함께",
    "전표_원본": "SAP 사업비 계정 전표 라인 · 최근 24개월 권장",
    "BS_원본": "재무상태표 · 월말 기준 · 한 행 = 기간 × 항목 (감독원 가로 양식도 불러오기 가능)",
    "PL_원본": "손익계산서 · 한 행 = 기간 × 항목 · 비용 항목은 양수",
    "예정사업비": "계리 예정사업비 · 포트폴리오 × 월",
    "부채변동표": "보험계약부채 변동표 (BEL · RA · CSM) · 감독원 양식·계리 결과",
}
N_EX = 3

GRAY_FILL = PatternFill("solid", fgColor="D9D9D9")
REQ_FILL = PatternFill("solid", fgColor="FFF2CC")
thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)


def main():
    fm = json.load(open(os.path.join(HERE, "tool_data.json"), encoding="utf-8"))["FM"]
    src = openpyxl.load_workbook(SRC_BOOK)
    hs = src["헤더설정"]
    wb = openpyxl.Workbook()
    guide = wb.active; guide.title = "안내"
    guide["B2"] = "TRACE 입력 양식 v0.5"; guide["B2"].font = Font(bold=True, size=14)
    for i, ln in enumerate([
        "각 시트의 4행 머리글이 TRACE 불러오기가 찾는 열 이름입니다. 노란 머리글 = 필수, 회색 = 선택.",
        "5행부터 있는 예시 3행(가상 데이터)은 지우고 실제 자료를 붙여넣거나, SAP에서 이 머리글로 내려받으세요.",
        "열 순서는 달라도 됩니다. 머리글 이름이 다르면 TRACE 본체의 '헤더설정' 시트에서 '찾을 헤더 이름'을 고치면 됩니다.",
        "TRACE 본체 설정 시트의 [OO 불러오기] 버튼으로 이 파일을 고르면 원본 시트에 들어갑니다.",
    ]):
        guide.cell(row=4 + i, column=2, value=ln)
    guide.cell(row=9, column=2, value="시트").font = Font(bold=True)
    guide.cell(row=9, column=3, value="받는 자료").font = Font(bold=True)
    for c in (guide["B9"], guide["C9"]): c.fill = GRAY_FILL
    guide.column_dimensions["A"].width = 2; guide.column_dimensions["B"].width = 16; guide.column_dimensions["C"].width = 90

    for k, (sheet, (r0, n, _key)) in enumerate(fm.items()):
        guide.cell(row=10 + k, column=2, value=sheet).hyperlink = f"#'{sheet}'!A1"
        guide.cell(row=10 + k, column=3, value=DESC.get(sheet, ""))
        cols = [(hs.cell(row=r0 + i, column=3).value, hs.cell(row=r0 + i, column=4).value or "",
                 hs.cell(row=r0 + i, column=5).value, hs.cell(row=r0 + i, column=6).value or "") for i in range(n)]
        ws = wb.create_sheet(sheet)
        ws["A1"] = f"TRACE · {sheet} 입력 양식"; ws["A1"].font = Font(bold=True, size=12)
        ws["A2"] = DESC.get(sheet, "") + " · 예시 행은 가상 데이터"; ws["A2"].font = Font(color="808080", size=9)
        so = src[sheet]
        for j, (name, alias, req, miss) in enumerate(cols, start=1):
            h = ws.cell(row=4, column=j, value=name)
            h.font = Font(bold=True); h.fill = REQ_FILL if req == "필수" else GRAY_FILL
            h.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); h.border = BOX
            note = f"{req}" + (f" · 대체 이름: {alias}" if alias else "") + (f" · 없을 때: {miss}" if miss else "")
            ws.cell(row=3, column=j, value=note).font = Font(color="808080", size=8)
            ws.cell(row=3, column=j).alignment = Alignment(wrap_text=True, vertical="bottom")
            for i in range(N_EX):
                s = so.cell(row=5 + i, column=j)
                c = ws.cell(row=5 + i, column=j, value=s.value)
                c.font = Font(color="0000FF"); c.border = BOX; c.number_format = s.number_format
            ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = max(12, min(30, len(str(name)) * 2 + 6))
        ws.row_dimensions[3].height = 48
        ws.freeze_panes = "A5"
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    wb.save(OUT)
    print("saved", OUT)


if __name__ == "__main__":
    main()
