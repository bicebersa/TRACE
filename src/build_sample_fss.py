"""예시 감독원 양식(가로형) 파일: 재무상태표 · 손익계산서 · 보험계약부채 변동표 (기준월 한 달)."""
import json, os, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
HERE = os.path.dirname(os.path.abspath(__file__))
D = json.load(open(os.path.join(HERE, "tool_data.json")))
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "samples", "샘플_감독원양식_2026-09.xlsx")
Y, M = 2026, 9
FSS_BS = [("10000", "자산총계", "A000"), ("11000", "  운용자산", "A100"), ("11100", "    현금및예치금", "A110"), ("11200", "    유가증권", "A120"),
          ("11300", "    대출채권", "A130"), ("11400", "    투자부동산", "A140"), ("12000", "  보험계약자산", "A200"), ("13000", "  재보험계약자산", "A300"),
          ("19000", "  기타자산", "A900"), ("20000", "부채총계", "L000"), ("21000", "  보험계약부채", "L100"), ("21100", "    최선추정부채(BEL)", "L110"),
          ("21200", "    위험조정(RA)", "L120"), ("21300", "    보험계약마진(CSM)", "L130"), ("22000", "  재보험계약부채", "L200"), ("29000", "  기타부채", "L900"),
          ("30000", "자본총계", "E000"), ("31000", "  자본금", "E100"), ("32000", "  이익잉여금", "E200"), ("33000", "  기타포괄손익누계액", "E300")]
FSS_PL = [("40000", "보험손익", "P100"), ("41000", "  보험수익", "P110"), ("42000", "  보험서비스비용", "P120"), ("43000", "  재보험손익", "P130"),
          ("50000", "투자손익", "P200"), ("51000", "  투자수익", "P210"), ("52000", "  보험금융비용", "P220"), ("60000", "기타사업비용", "P300"),
          ("70000", "법인세비용차감전순이익", "P400"), ("80000", "법인세비용", "P500"), ("90000", "당기순이익", "P600")]
FAC = {"F00": ("B00", "기초"), "F10": ("B10", "신계약"), "F21": ("B21", "추정변경(사업비)"), "F22": ("B22", "추정변경(기타)"), "F31": ("B31", "경험조정(보험취득현금흐름)"),
       "F32": ("B32", "경험조정(기타)"), "F40": ("B40", "보험금융비용"), "F51": ("B51", "보험계약마진 상각"), "F52": ("B52", "위험조정 변동(당기)"),
       "F60": ("B60", "현금흐름 및 기타"), "F99": ("B99", "기말")}
bs = {(r[0], r[1], r[2]): r[6] for r in D["BS"]}
pl = {(r[0], r[1], r[2]): (r[6], r[7]) for r in D["PL"]}
thin = Side(style="thin", color="999999"); B = Border(left=thin, right=thin, top=thin, bottom=thin)
H = PatternFill("solid", fgColor="DDE3EA"); FN = "맑은 고딕"
wb = Workbook()
def sheet(ws, title, cols, rows, widths):
    ws["A1"] = "보험회사 업무보고서 (예시 양식)"; ws["A1"].font = Font(name=FN, size=9, color="666666")
    ws["A2"] = title; ws["A2"].font = Font(name=FN, size=14, bold=True)
    ws["A3"] = f"회사명: 예시보험㈜    기준: {Y}년 {M}월    단위: 원"; ws["A3"].font = Font(name=FN, size=9)
    for j, c in enumerate(cols):
        x = ws.cell(row=5, column=j + 1, value=c); x.font = Font(name=FN, bold=True); x.fill = H; x.border = B; x.alignment = Alignment(horizontal="center")
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            x = ws.cell(row=6 + i, column=j + 1, value=v); x.font = Font(name=FN, bold=isinstance(row[1], str) and not row[1].startswith(" ")); x.border = B
            if isinstance(v, (int, float)) and j >= 2: x.number_format = "#,##0"
    for j, w in enumerate(widths): ws.column_dimensions[chr(65 + j)].width = w
ws = wb.active; ws.title = "재무상태표"
sheet(ws, "재무상태표", ["항목코드", "항목명", "당기", "전기"], [[c, n, bs[(Y, M, s)], bs[(Y - 1, 12, s)]] for c, n, s in FSS_BS], [11, 26, 20, 20])
ws = wb.create_sheet("손익계산서")
sheet(ws, "손익계산서", ["항목코드", "항목명", "당기", "당기누계", "전년동기누계"], [[c, n, pl[(Y, M, s)][0], pl[(Y, M, s)][1], pl[(Y - 1, M, s)][1]] for c, n, s in FSS_PL], [11, 26, 18, 20, 20])
ws = wb.create_sheet("보험계약부채변동")
rows = [[r[2], r[3], r[4], FAC[r[5]][0], FAC[r[5]][1], r[7], r[8], r[9]] for r in D["ROLL"] if (r[0], r[1]) == (Y, M)]
sheet(ws, "보험계약부채 변동 (포트폴리오별)", ["포트폴리오", "포트폴리오명", "측정모형", "요인코드", "변동요인", "최선추정부채", "위험조정", "보험계약마진"], rows, [11, 12, 9, 9, 26, 18, 16, 16])
wb.save(OUT); print("wrote", OUT, len(rows))
