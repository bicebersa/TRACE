"""초기 산출물: SAP 보험사 시산표·자산운용 예시 양식 생성 (legacy/에 저장)."""
import datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter as CL

import os
OUT = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "SAP_보험사_시산표_자산운용_양식.xlsx")
FONT = "맑은 고딕"
NUM = '#,##0;(#,##0);"-"'
NUM2 = '#,##0.00;(#,##0.00);"-"'
PCT = '0.0%;(0.0%);"-"'
PCT2 = '0.00%;(0.00%);"-"'
DATE = "yyyy-mm-dd"

NAVY = PatternFill("solid", fgColor="1F3864")
SUB = PatternFill("solid", fgColor="D9E1F2")
TOT = PatternFill("solid", fgColor="F2F2F2")
YEL = PatternFill("solid", fgColor="FFFF00")
thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
TOPLINE = Border(top=Side(style="thin", color="000000"), bottom=Side(style="double", color="000000"))

def f(color="000000", bold=False, size=10):
    return Font(name=FONT, color=color, bold=bold, size=size)
BLUE, BLACK, GREEN = "0000FF", "000000", "008000"

def title(ws, text, sub):
    ws["A1"] = text
    ws["A1"].font = f(bold=True, size=14)
    ws["A2"] = sub
    ws["A2"].font = f(GREEN, size=9)

def header(ws, row, cols, height=32):
    for i, h in enumerate(cols, 1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = Font(name=FONT, color="FFFFFF", bold=True, size=10)
        c.fill = NAVY
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BOX
    ws.row_dimensions[row].height = height

def put(ws, r, c, v, color=BLACK, fmt=None, bold=False, fill=None, align=None):
    cell = ws.cell(row=r, column=c, value=v)
    cell.font = f(color, bold)
    cell.border = BOX
    if fmt: cell.number_format = fmt
    if fill: cell.fill = fill
    if align: cell.alignment = Alignment(horizontal=align, vertical="center")
    return cell

def widths(ws, ws_w):
    for i, w in enumerate(ws_w, 1):
        ws.column_dimensions[CL(i)].width = w

wb = Workbook()

# ------------------------------------------------------------------ 안내
ws0 = wb.active
ws0.title = "안내"
title(ws0, "SAP 보험사 시산표 · 자산운용 테이블 예시 양식", "K-IFRS 17(보험계약) / K-IFRS 9(금융상품) 기준 · 예시 데이터는 가상 수치입니다")
ws0["A4"] = "기본 정보 (노란 칸을 수정하면 전 시트에 반영)"; ws0["A4"].font = f(bold=True, size=11)
params = [
    ("회사코드 (BUKRS)", "1000", None),
    ("회사명", "예시보험㈜", None),
    ("원장 (RLDNR)", "0L", None),
    ("회계연도 (GJAHR)", "2026", None),
    ("기간 (POPER)", "009", None),
    ("기준일", dt.date(2026, 9, 30), DATE),
    ("금액 단위", "백만원 (KRW mn)", None),
]
for i, (k, v, fm) in enumerate(params):
    r = 5 + i
    put(ws0, r, 1, k, bold=True, fill=SUB)
    c = put(ws0, r, 2, v, BLUE, fm, fill=YEL)
P = {k: f"안내!$B${5+i}" for i, (k, _, _) in enumerate(params)}
P_CO, P_LDG, P_YR, P_PER, P_DATE, P_UNIT = (P["회사코드 (BUKRS)"], P["원장 (RLDNR)"], P["회계연도 (GJAHR)"],
                                            P["기간 (POPER)"], P["기준일"], P["금액 단위"])

ws0["A13"] = "시트 구성"; ws0["A13"].font = f(bold=True, size=11)
sheets = [
    ("계정과목", "G/L 계정 마스터 (SKA1/SKAT/SKB1 대응). 보험사 K-IFRS 17·9 계정 체계 예시"),
    ("전표", "분개 라인 (BKPF/BSEG, S/4HANA는 ACDOCA). 시산표의 당기 차변·대변이 여기서 집계됨"),
    ("시산표", "합계잔액시산표 (FAGLFLEXT/ACDOCA, T-code FAGLB03·S_ALR_87012277 결과와 같은 형태)"),
    ("자산운용명세", "종목별 운용자산 포지션 (SAP TRM 증권·FS-CML 대출 등 서브원장 추출 형태)"),
    ("자산운용요약", "자산군별·IFRS 9 분류별 요약, 듀레이션/금리 가중평균, 명세 ↔ 시산표 GL 대사"),
]
header(ws0, 14, ["시트", "내용"], 20)
for i, (s, d) in enumerate(sheets):
    put(ws0, 15 + i, 1, s, bold=True)
    put(ws0, 15 + i, 2, d)

ws0["A21"] = "셀 색상 규칙"; ws0["A21"].font = f(bold=True, size=11)
legend = [
    ("파란 글씨", "직접 입력하는 값 (예시 데이터를 지우고 실제 값으로 교체)", BLUE, None),
    ("검정 글씨", "수식 (수정하지 마세요)", BLACK, None),
    ("초록 글씨", "다른 시트를 참조하는 수식", GREEN, None),
    ("노란 배경", "기본 정보 · 핵심 가정", BLACK, YEL),
]
for i, (k, d, col, fl) in enumerate(legend):
    put(ws0, 22 + i, 1, k, col, bold=True, fill=fl)
    put(ws0, 22 + i, 2, d)

ws0["A27"] = "SAP 필드 매핑 (시산표)"; ws0["A27"].font = f(bold=True, size=11)
header(ws0, 28, ["양식 열", "SAP 필드 (ECC FAGLFLEXT / S/4 ACDOCA)"], 20)
maps = [
    ("회사코드", "RBUKRS"), ("원장", "RLDNR (선행원장 0L)"), ("회계연도", "RYEAR / GJAHR"),
    ("기간", "POPER (001~012, 특별기간 013~016)"), ("G/L 계정", "RACCT"),
    ("기초잔액", "HSLVT (이월잔액) + 직전기간까지 HSL01~HSLnn 누계"),
    ("당기 차변 / 대변", "DRCRK = S(차변) / H(대변) 기준 HSL 합계"),
    ("금액", "HSL (회사코드 통화), 거래통화는 WSL / TSL"),
    ("부호 규칙", "SAP는 차변 +, 대변 − 로 저장. 이 양식도 동일 (기말잔액 음수 = 대변잔액)"),
]
for i, (k, d) in enumerate(maps):
    put(ws0, 29 + i, 1, k, bold=True)
    put(ws0, 29 + i, 2, d)
ws0["A39"] = ("※ 자산운용 서브원장은 회사별로 SAP TRM, FS-CML, 외부 운용시스템(인터페이스) 등 구성이 달라 "
              "테이블·필드명이 다릅니다. 자산운용명세 시트의 열 구성은 공통적으로 쓰는 항목 기준입니다.")
ws0["A39"].font = f("595959", size=9)
widths(ws0, [22, 95])

# ------------------------------------------------------------------ 계정과목
coa = [
    # code, name, 대분류, 중분류, BS/PL, 정상잔액
    ("10101000", "현금및현금성자산", "자산", "현금성자산", "BS", "차"),
    ("11101000", "FVPL금융자산-주식", "자산", "운용자산", "BS", "차"),
    ("11102000", "FVPL금융자산-수익증권", "자산", "운용자산", "BS", "차"),
    ("11201000", "FVOCI금융자산-국공채", "자산", "운용자산", "BS", "차"),
    ("11202000", "FVOCI금융자산-회사채", "자산", "운용자산", "BS", "차"),
    ("11203000", "FVOCI금융자산-외화채권", "자산", "운용자산", "BS", "차"),
    ("11301000", "AC금융자산-채권", "자산", "운용자산", "BS", "차"),
    ("11309000", "AC금융자산 손실충당금", "자산", "운용자산(차감)", "BS", "대"),
    ("11401000", "대출채권-기업대출", "자산", "운용자산", "BS", "차"),
    ("11402000", "대출채권-부동산담보대출", "자산", "운용자산", "BS", "차"),
    ("11409000", "대출채권 손실충당금", "자산", "운용자산(차감)", "BS", "대"),
    ("11501000", "투자부동산", "자산", "운용자산", "BS", "차"),
    ("11601000", "미수이자", "자산", "운용자산", "BS", "차"),
    ("12101000", "유형자산", "자산", "비운용자산", "BS", "차"),
    ("12201000", "재보험계약자산", "자산", "보험계약자산", "BS", "차"),
    ("12901000", "이연법인세자산", "자산", "비운용자산", "BS", "차"),
    ("21101000", "보험계약부채-최선추정부채(BEL)", "부채", "보험계약부채", "BS", "대"),
    ("21102000", "보험계약부채-위험조정(RA)", "부채", "보험계약부채", "BS", "대"),
    ("21103000", "보험계약부채-보험계약마진(CSM)", "부채", "보험계약부채", "BS", "대"),
    ("21201000", "재보험계약부채", "부채", "보험계약부채", "BS", "대"),
    ("22101000", "차입부채(후순위채)", "부채", "기타부채", "BS", "대"),
    ("22201000", "미지급비용", "부채", "기타부채", "BS", "대"),
    ("22301000", "당기법인세부채", "부채", "기타부채", "BS", "대"),
    ("31101000", "자본금", "자본", "자본금", "BS", "대"),
    ("31201000", "자본잉여금", "자본", "자본잉여금", "BS", "대"),
    ("31301000", "이익잉여금", "자본", "이익잉여금", "BS", "대"),
    ("31401000", "기타포괄손익누계액-FVOCI평가손익", "자본", "기타포괄손익", "BS", "대"),
    ("31402000", "기타포괄손익누계액-보험금융손익", "자본", "기타포괄손익", "BS", "대"),
    ("41101000", "보험수익", "수익", "보험손익", "PL", "대"),
    ("41201000", "재보험수익", "수익", "보험손익", "PL", "대"),
    ("42101000", "이자수익", "수익", "투자손익", "PL", "대"),
    ("42102000", "배당수익", "수익", "투자손익", "PL", "대"),
    ("42201000", "금융자산처분이익", "수익", "투자손익", "PL", "대"),
    ("42202000", "FVPL금융자산평가이익", "수익", "투자손익", "PL", "대"),
    ("42301000", "임대료수익", "수익", "투자손익", "PL", "대"),
    ("51101000", "보험서비스비용", "비용", "보험손익", "PL", "차"),
    ("51201000", "재보험비용", "비용", "보험손익", "PL", "차"),
    ("52101000", "보험금융비용", "비용", "투자손익", "PL", "차"),
    ("52201000", "금융자산처분손실", "비용", "투자손익", "PL", "차"),
    ("52202000", "FVPL금융자산평가손실", "비용", "투자손익", "PL", "차"),
    ("52301000", "손실충당금전입액(기대신용손실)", "비용", "투자손익", "PL", "차"),
    ("52401000", "이자비용", "비용", "투자손익", "PL", "차"),
    ("52501000", "투자부동산감가상각비", "비용", "투자손익", "PL", "차"),
    ("53101000", "법인세비용", "비용", "법인세", "PL", "차"),
]
NAME = {c[0]: c[1] for c in coa}
ws1 = wb.create_sheet("계정과목")
title(ws1, "G/L 계정과목 마스터", "SAP 대응: SKA1(계정과목표) · SKAT(계정명) · SKB1(회사코드별 설정)")
header(ws1, 4, ["G/L 계정\n(SAKNR)", "계정명\n(TXT50)", "대분류", "중분류", "재무제표\n(BS/PL)", "정상잔액\n(차/대)", "비고"])
for i, row in enumerate(coa):
    for j, v in enumerate(row):
        put(ws1, 5 + i, j + 1, v, BLUE, align="center" if j in (0, 2, 4, 5) else None)
    put(ws1, 5 + i, 7, None)
COA_LAST = 4 + len(coa) + 30  # 추가 입력 여유
ws1.freeze_panes = "B5"
ws1.auto_filter.ref = f"A4:G{4+len(coa)}"
widths(ws1, [13, 34, 9, 15, 10, 10, 30])
COA_A = f"계정과목!$A$5:$A${COA_LAST}"
def coa_lookup(cell_ref, col):
    return f'=IFERROR(INDEX(계정과목!${col}$5:${col}${COA_LAST},MATCH({cell_ref},{COA_A},0)),"계정 없음")'

# ------------------------------------------------------------------ 자산운용명세 데이터
# (pos, 포트폴리오, ISIN, 종목명, 자산군, IFRS9, 통화, 액면(통화,백만), 환율, 취득원가, 상각후원가, 공정가치, ECL, 미수이자, 금리, 등급, 만기, 수정듀레이션, 자산GL, 충당금GL)
H = [
    ("P0001", "일반계정", "KR7000000001", "국내상장주식 A (전기전자)", "주식", "FVPL", "KRW", None, None, 85000, 85000, 97400, 0, 0, None, "-", None, None, "11101000", ""),
    ("P0002", "일반계정", "KR7000000002", "국내상장주식 B (금융)", "주식", "FVPL", "KRW", None, None, 42000, 42000, 38600, 0, 0, None, "-", None, None, "11101000", ""),
    ("P0003", "일반계정", "KRZ500000003", "국내채권형 사모펀드", "수익증권", "FVPL", "KRW", None, None, 120000, 120000, 123500, 0, 0, None, "-", None, 3.1, "11102000", ""),
    ("P0004", "일반계정", "KYG000000004", "해외인프라 블라인드펀드", "수익증권", "FVPL", "USD", 60, 1380.0, 79800, 79800, 84200, 0, 0, None, "-", None, None, "11102000", ""),
    ("P0005", "일반계정", "KRZ500000005", "국내 오피스 부동산펀드", "수익증권", "FVPL", "KRW", None, None, 65000, 65000, 66100, 0, 0, None, "-", None, None, "11102000", ""),
    ("P0006", "일반계정", "KR1035000006", "국고채권 30년", "국공채", "FVOCI", "KRW", 500000, 1.0, 486200, 488900, 452300, 0, 8460, 0.0325, "AAA", dt.date(2056, 3, 10), 19.8, "11201000", ""),
    ("P0007", "일반계정", "KR1035000007", "국고채권 20년", "국공채", "FVOCI", "KRW", 300000, 1.0, 301500, 301100, 288400, 0, 450, 0.0275, "AAA", dt.date(2044, 9, 10), 14.6, "11201000", ""),
    ("P0008", "일반계정", "KR3500000008", "공사채 A (특수채)", "국공채", "FVOCI", "KRW", 200000, 1.0, 199400, 199400, 196800, 12, 2240, 0.0340, "AAA", dt.date(2034, 6, 15), 7.0, "11201000", ""),
    ("P0009", "일반계정", "KR6000000009", "회사채 A (AA+)", "회사채", "FVOCI", "KRW", 150000, 1.0, 150000, 150000, 148200, 45, 2560, 0.0385, "AA+", dt.date(2031, 4, 20), 4.2, "11202000", ""),
    ("P0010", "일반계정", "KR6000000010", "회사채 B (AA-)", "회사채", "FVOCI", "KRW", 100000, 1.0, 99800, 99880, 97900, 85, 4390, 0.0420, "AA-", dt.date(2029, 11, 30), 2.9, "11202000", ""),
    ("P0011", "일반계정", "US9120000011", "미국 국채 30년", "외화채권", "FVOCI", "USD", 200, 1380.0, 268000, 270500, 251600, 20, 1370, 0.0425, "AA+", dt.date(2054, 8, 15), 16.9, "11203000", ""),
    ("P0012", "일반계정", "US0000000012", "해외 회사채 A (USD)", "외화채권", "FVOCI", "USD", 80, 1380.0, 106500, 107200, 108900, 160, 2060, 0.0510, "A", dt.date(2035, 2, 15), 7.4, "11203000", ""),
    ("P0013", "일반계정", "KR1035000013", "국고채권 50년", "국공채", "AC", "KRW", 400000, 1.0, 372000, 376800, 331200, 0, 5480, 0.0250, "AAA", dt.date(2074, 3, 10), 27.5, "11301000", "11309000"),
    ("P0014", "일반계정", "KR2000000014", "지방채 A", "국공채", "AC", "KRW", 120000, 1.0, 119500, 119700, 116900, 18, 3410, 0.0310, "AAA", dt.date(2036, 12, 10), 8.6, "11301000", "11309000"),
    ("L0015", "일반계정", "LN-2026-0015", "PF대출 (선순위, 물류센터)", "대출채권", "AC", "KRW", 80000, 1.0, 80000, 80000, 79100, 960, 390, 0.0580, "BBB+", dt.date(2028, 12, 31), 2.1, "11401000", "11409000"),
    ("L0016", "일반계정", "LN-2026-0016", "인수금융 대출 (Term Loan A)", "대출채권", "AC", "KRW", 50000, 1.0, 50000, 50000, 49600, 750, 260, 0.0620, "BBB", dt.date(2030, 6, 30), 3.3, "11401000", "11409000"),
    ("L0017", "일반계정", "LN-POOL-0017", "부동산담보대출 (개인, Pool)", "대출채권", "AC", "KRW", 150000, 1.0, 150000, 150000, 148500, 300, 580, 0.0460, "-", dt.date(2045, 9, 30), 6.5, "11402000", "11409000"),
    ("R0018", "일반계정", "RE-0018", "투자부동산 (임대 오피스)", "투자부동산", "원가모형", "KRW", None, None, 210000, 182400, 265000, 0, 0, None, "-", None, None, "11501000", ""),
]
def book(h):
    return h[10] if h[5] in ("AC", "원가모형") else h[11]

# GL별 기말잔액 (명세 합계와 일치시키기 위한 예시값 계산)
end = {}
for h in H:
    end[h[18]] = end.get(h[18], 0) + book(h)
    if h[19]:
        end[h[19]] = end.get(h[19], 0) - h[12]
end["11601000"] = sum(h[13] for h in H)

# ------------------------------------------------------------------ 전표 (9월 예시 분개)
J = [  # 전기일, 유형, 차변계정, 대변계정, 금액, 적요
    (dt.date(2026, 9, 1), "SA", "10101000", "21101000", 95000, "보험료 수취 (보험계약부채 계상)"),
    (dt.date(2026, 9, 30), "SA", "21101000", "41101000", 62000, "보험수익 인식 - 예상 보험금·비용 해제"),
    (dt.date(2026, 9, 30), "SA", "21103000", "41101000", 6500, "보험수익 인식 - CSM 상각"),
    (dt.date(2026, 9, 30), "SA", "21102000", "41101000", 1800, "보험수익 인식 - RA 해제"),
    (dt.date(2026, 9, 15), "SA", "51101000", "10101000", 48000, "보험금 지급"),
    (dt.date(2026, 9, 30), "SA", "51101000", "22201000", 9500, "사업비 발생 (미지급)"),
    (dt.date(2026, 9, 30), "SA", "52101000", "21101000", 7600, "보험금융비용 (이자부리)"),
    (dt.date(2026, 9, 30), "SA", "51201000", "12201000", 3200, "재보험비용 (출재보험료 배분)"),
    (dt.date(2026, 9, 30), "SA", "12201000", "41201000", 2100, "재보험수익 (회수 예상액)"),
    (dt.date(2026, 9, 30), "SA", "11601000", "42101000", 8900, "채권·대출 이자 발생 (경과이자)"),
    (dt.date(2026, 9, 10), "SA", "10101000", "11601000", 6200, "이표 이자 수취"),
    (dt.date(2026, 9, 5), "SA", "11201000", "10101000", 199400, "공사채 A 매입 (P0008)"),
    (dt.date(2026, 9, 30), "SA", "11101000", "42202000", 4300, "FVPL 주식 평가이익"),
    (dt.date(2026, 9, 30), "SA", "52202000", "11102000", 1200, "FVPL 수익증권 평가손실"),
    (dt.date(2026, 9, 30), "SA", "31401000", "11201000", 5800, "FVOCI 국공채 평가손실 (금리 상승)"),
    (dt.date(2026, 9, 18), "SA", "10101000", "11202000", 50000, "회사채 처분 - 장부금액 제거"),
    (dt.date(2026, 9, 18), "SA", "10101000", "42201000", 1200, "회사채 처분이익"),
    (dt.date(2026, 9, 30), "SA", "52301000", "11409000", 140, "대출채권 기대신용손실 추가 적립"),
    (dt.date(2026, 9, 25), "SA", "10101000", "11401000", 10000, "기업대출 원금 상환"),
    (dt.date(2026, 9, 30), "SA", "52501000", "11501000", 650, "투자부동산 감가상각"),
    (dt.date(2026, 9, 5), "SA", "10101000", "42301000", 1450, "임대료 수취"),
    (dt.date(2026, 9, 20), "SA", "10101000", "42102000", 900, "배당금 수취"),
    (dt.date(2026, 9, 30), "SA", "52401000", "22201000", 320, "후순위채 이자비용 발생"),
    (dt.date(2026, 9, 30), "SA", "11301000", "42101000", 410, "AC 채권 할인 상각 (유효이자)"),
    (dt.date(2026, 9, 30), "SA", "53101000", "22301000", 2100, "법인세비용 (월할)"),
    (dt.date(2026, 9, 28), "SA", "22201000", "10101000", 8000, "미지급비용 지급"),
]
dr = {}; cr = {}
for _, _, d, c, a, _ in J:
    dr[d] = dr.get(d, 0) + a
    cr[c] = cr.get(c, 0) + a

ws2 = wb.create_sheet("전표")
title(ws2, "분개 라인 (전표)", "SAP 대응: BKPF(전표헤더)·BSEG(라인) / S/4HANA ACDOCA · 금액 단위: 백만원")
header(ws2, 4, ["회사코드\n(BUKRS)", "전표번호\n(BELNR)", "회계연도\n(GJAHR)", "기간\n(POPER)", "전기일\n(BUDAT)",
                "전표유형\n(BLART)", "라인\n(BUZEI)", "G/L 계정\n(HKONT)", "계정명", "차대\n(SHKZG)",
                "금액\n(DMBTR)", "부호금액\n(차+/대−)", "적요\n(SGTXT)"])
r = 5
for n, (d, t, da, ca, a, txt) in enumerate(J, 1):
    belnr = f"{100000000 + n}"
    for line, (acc, sh) in enumerate([(da, "S"), (ca, "H")], 1):
        put(ws2, r, 1, f"={P_CO}", GREEN, align="center")
        put(ws2, r, 2, belnr, BLUE, align="center")
        put(ws2, r, 3, f"={P_YR}", GREEN, align="center")
        put(ws2, r, 4, f'=TEXT(MONTH(E{r}),"000")', align="center")
        put(ws2, r, 5, d, BLUE, DATE, align="center")
        put(ws2, r, 6, t, BLUE, align="center")
        put(ws2, r, 7, f"{line:03d}", BLUE, align="center")
        put(ws2, r, 8, acc, BLUE, align="center")
        put(ws2, r, 9, coa_lookup(f"H{r}", "B"), GREEN)
        put(ws2, r, 10, sh, BLUE, align="center")
        put(ws2, r, 11, a, BLUE, NUM)
        put(ws2, r, 12, f'=IF(J{r}="S",K{r},-K{r})', fmt=NUM)
        put(ws2, r, 13, txt, BLUE)
        r += 1
J_LAST_DATA = r - 1
J_LAST = 1000  # 추가 입력 여유
put(ws2, r, 9, "합계", bold=True, fill=TOT)
put(ws2, r, 10, "차변", bold=True, fill=TOT, align="center")
put(ws2, r, 11, f'=SUMIF(J5:J{J_LAST_DATA},"S",K5:K{J_LAST_DATA})', fmt=NUM, bold=True, fill=TOT)
put(ws2, r, 12, f"=SUM(L5:L{J_LAST_DATA})", fmt=NUM, bold=True, fill=TOT)
put(ws2, r, 13, f'=IF(ABS(L{r})<0.5,"OK - 차대 일치","ERROR - 차대 불일치")', bold=True, fill=TOT)
put(ws2, r + 1, 10, "대변", bold=True, fill=TOT, align="center")
put(ws2, r + 1, 11, f'=SUMIF(J5:J{J_LAST_DATA},"H",K5:K{J_LAST_DATA})', fmt=NUM, bold=True, fill=TOT)
ws2.cell(row=r + 3, column=1, value="※ 행을 추가할 때는 합계 행 위에 삽입하세요. 시산표는 이 시트의 5~1000행을 집계합니다.").font = f("595959", size=9)
ws2.freeze_panes = "A5"
ws2.auto_filter.ref = f"A4:M{J_LAST_DATA}"
widths(ws2, [9, 12, 9, 8, 11, 9, 7, 11, 30, 7, 12, 12, 36])
dv = DataValidation(type="list", formula1='"S,H"', allow_blank=True); ws2.add_data_validation(dv); dv.add(f"J5:J{J_LAST_DATA}")

# ------------------------------------------------------------------ 시산표
begin_other = {
    "10101000": 168000, "12101000": 64000, "12201000": 41000, "12901000": 28500,
    "21101000": -2030000, "21102000": -118000, "21103000": -372000, "21201000": -6800,
    "22101000": -50000, "22201000": -23400, "22301000": -11200,
    "31101000": -100000, "31201000": -85000,
    "31401000": 49500, "31402000": -62000,
    "41101000": -560000, "41201000": -17500, "42101000": -71000, "42102000": -6800,
    "42201000": -9400, "42202000": -14200, "42301000": -11300,
    "51101000": 455000, "51201000": 26100, "52101000": 60500, "52201000": 3100,
    "52202000": 8700, "52301000": 950, "52401000": 2550, "52501000": 5200, "53101000": 17600,
}
begin = dict(begin_other)
for acc, e in end.items():  # 운용자산: 기말 = 명세 합계가 되도록 기초 역산
    begin[acc] = e - dr.get(acc, 0) + cr.get(acc, 0)
begin["31301000"] = 0
begin["31301000"] = -sum(begin.values())  # 이익잉여금으로 기초 대차 맞춤
assert begin["31301000"] < 0, begin["31301000"]
for acc, v in begin.items():
    if acc in end: assert v >= 0 or acc in ("11309000", "11409000"), (acc, v)

ws3 = wb.create_sheet("시산표")
title(ws3, "합계잔액시산표 (Trial Balance)", "SAP 대응: FAGLFLEXT / ACDOCA · T-code FAGLB03, S_ALR_87012277 · 부호: 차변 +, 대변 −")
ws3["A3"] = f'="회사: "&{P["회사명"]}&" ("&{P_CO}&")   회계연도: "&{P_YR}&"   기간: "&{P_PER}&"   기준일: "&TEXT({P_DATE},"yyyy-mm-dd")&"   단위: "&{P_UNIT}'
ws3["A3"].font = f(GREEN, size=9)
header(ws3, 5, ["회사코드\n(RBUKRS)", "원장\n(RLDNR)", "회계연도\n(RYEAR)", "기간\n(POPER)", "G/L 계정\n(RACCT)", "계정명",
                "대분류", "BS/PL", "기초잔액\n(전기말 누계)", "당기 차변", "당기 대변", "기말잔액\n(차+/대−)",
                "차변잔액", "대변잔액"])
JR = f"전표!$H$5:$H${J_LAST}"; JS = f"전표!$J$5:$J${J_LAST}"; JK = f"전표!$K$5:$K${J_LAST}"
TB0 = 6
for i, (code, *_rest) in enumerate(coa):
    r = TB0 + i
    put(ws3, r, 1, f"={P_CO}", GREEN, align="center")
    put(ws3, r, 2, f"={P_LDG}", GREEN, align="center")
    put(ws3, r, 3, f"={P_YR}", GREEN, align="center")
    put(ws3, r, 4, f"={P_PER}", GREEN, align="center")
    put(ws3, r, 5, code, BLUE, align="center")
    put(ws3, r, 6, coa_lookup(f"E{r}", "B"), GREEN)
    put(ws3, r, 7, coa_lookup(f"E{r}", "C"), GREEN, align="center")
    put(ws3, r, 8, coa_lookup(f"E{r}", "E"), GREEN, align="center")
    put(ws3, r, 9, begin[code], BLUE, NUM)
    put(ws3, r, 10, f'=SUMIFS({JK},{JR},E{r},{JS},"S")', GREEN, NUM)
    put(ws3, r, 11, f'=SUMIFS({JK},{JR},E{r},{JS},"H")', GREEN, NUM)
    put(ws3, r, 12, f"=I{r}+J{r}-K{r}", fmt=NUM)
    put(ws3, r, 13, f"=MAX(L{r},0)", fmt=NUM)
    put(ws3, r, 14, f"=MAX(-L{r},0)", fmt=NUM)
TB1 = TB0 + len(coa) - 1
t = TB1 + 1
put(ws3, t, 6, "합계", bold=True, fill=TOT)
for c in range(1, 15):
    ws3.cell(row=t, column=c).fill = TOT
for c in range(9, 15):
    col = CL(c)
    put(ws3, t, c, f"=SUM({col}{TB0}:{col}{TB1})", fmt=NUM, bold=True, fill=TOT).border = TOPLINE
chk = t + 1
put(ws3, chk, 6, "검증", bold=True)
put(ws3, chk, 9, f'=IF(ABS(I{t})<0.5,"OK","기초 불일치")', bold=True, align="center")
put(ws3, chk, 10, f'=IF(ABS(J{t}-K{t})<0.5,"OK","차대 불일치")', bold=True, align="center")
ws3.merge_cells(start_row=chk, start_column=10, end_row=chk, end_column=11)
put(ws3, chk, 12, f'=IF(ABS(L{t})<0.5,"OK","기말 불일치")', bold=True, align="center")
put(ws3, chk, 13, f'=IF(ABS(M{t}-N{t})<0.5,"OK","잔액 불일치")', bold=True, align="center")
ws3.merge_cells(start_row=chk, start_column=13, end_row=chk, end_column=14)

s = chk + 2
ws3.cell(row=s, column=6, value="대분류별 요약 (재무제표 표시 부호: 양수)").font = f(bold=True, size=11)
header(ws3, s + 1, [None] * 5 + ["대분류", "", "", "기초", "당기 차변", "당기 대변", "기말"], 20)
for c in range(1, 6):
    ws3.cell(row=s + 1, column=c).fill = PatternFill(None); ws3.cell(row=s + 1, column=c).border = Border()
cats = [("자산", 1), ("부채", -1), ("자본", -1), ("수익", -1), ("비용", 1)]
for i, (cat, sign) in enumerate(cats):
    r = s + 2 + i
    put(ws3, r, 6, cat, bold=True)
    for c in (9, 10, 11, 12):
        col = CL(c)
        sgn = "" if (sign == 1 or c in (10, 11)) else "-"
        put(ws3, r, c, f'={sgn}SUMIFS({col}${TB0}:{col}${TB1},$G${TB0}:$G${TB1},$F{r})', fmt=NUM)
r_ni = s + 2 + len(cats)
put(ws3, r_ni, 6, "당기순이익 (수익−비용, 연누계)", bold=True, fill=SUB)
put(ws3, r_ni, 9, f"=I{r_ni-2}-I{r_ni-1}", fmt=NUM, bold=True, fill=SUB)
put(ws3, r_ni, 12, f"=L{r_ni-2}-L{r_ni-1}", fmt=NUM, bold=True, fill=SUB)
put(ws3, r_ni + 1, 6, "대차검증 (자산 = 부채+자본+순이익)", bold=True)
put(ws3, r_ni + 1, 12, f'=IF(ABS(L{s+2}-L{s+3}-L{s+4}-L{r_ni})<0.5,"OK","불일치")', bold=True, align="center")
ws3.cell(row=r_ni + 3, column=6, value="※ 기초잔액은 전기말까지의 누계(손익계정은 연초~전월 누계)입니다. 당기 차변·대변은 '전표' 시트에서 자동 집계됩니다.").font = f("595959", size=9)
ws3.freeze_panes = "G6"
ws3.auto_filter.ref = f"A5:N{TB1}"
widths(ws3, [9, 7, 9, 7, 11, 34, 8, 7, 14, 13, 13, 14, 13, 13])

# ------------------------------------------------------------------ 자산운용명세
ws4 = wb.create_sheet("자산운용명세")
title(ws4, "운용자산 종목별 명세 (Investment Holdings)", "서브원장 추출 형태: SAP TRM(증권) · FS-CML(대출) · 부동산 관리 등 · 금액 단위: 백만원")
ws4["A3"] = f'="회사: "&{P["회사명"]}&" ("&{P_CO}&")   기준일: "&TEXT({P_DATE},"yyyy-mm-dd")&"   단위: "&{P_UNIT}'
ws4["A3"].font = f(GREEN, size=9)
cols4 = ["포지션번호", "회사코드", "포트폴리오\n(계정구분)", "종목코드\n(ISIN 등)", "종목명", "자산군", "IFRS 9\n분류",
         "통화", "액면금액\n(통화, 백만)", "환율\n(KRW/통화)", "취득원가", "상각후원가", "공정가치",
         "장부금액\n(총액)", "평가손익\n(FV−상각후원가)", "기대신용손실\n(ECL)", "순장부금액", "미수이자",
         "표면금리", "신용등급", "만기일", "잔존만기\n(년)", "수정\n듀레이션", "자산 G/L", "충당금 G/L"]
header(ws4, 5, cols4, 42)
H0 = 6
for i, h in enumerate(H):
    r = H0 + i
    (pos, pf, isin, nm, cls, ifrs, cur, face, fx, cost, ac, fv, ecl, accr, cpn, rating, mat, dur, gl, glp) = h
    put(ws4, r, 1, pos, BLUE, align="center")
    put(ws4, r, 2, f"={P_CO}", GREEN, align="center")
    put(ws4, r, 3, pf, BLUE, align="center")
    put(ws4, r, 4, isin, BLUE)
    put(ws4, r, 5, nm, BLUE)
    put(ws4, r, 6, cls, BLUE, align="center")
    put(ws4, r, 7, ifrs, BLUE, align="center")
    put(ws4, r, 8, cur, BLUE, align="center")
    put(ws4, r, 9, face, BLUE, NUM)
    put(ws4, r, 10, fx, BLUE, NUM2)
    put(ws4, r, 11, cost, BLUE, NUM)
    put(ws4, r, 12, ac, BLUE, NUM)
    put(ws4, r, 13, fv, BLUE, NUM)
    put(ws4, r, 14, f'=IF(OR(G{r}="AC",G{r}="원가모형"),L{r},M{r})', fmt=NUM)
    put(ws4, r, 15, f"=M{r}-L{r}", fmt=NUM)
    put(ws4, r, 16, ecl, BLUE, NUM)
    put(ws4, r, 17, f'=N{r}-IF(G{r}="AC",P{r},0)', fmt=NUM)
    put(ws4, r, 18, accr, BLUE, NUM)
    put(ws4, r, 19, cpn, BLUE, PCT2)
    put(ws4, r, 20, rating, BLUE, align="center")
    put(ws4, r, 21, mat, BLUE, DATE, align="center")
    put(ws4, r, 22, f'=IF(U{r}="","",ROUND((U{r}-{P_DATE})/365.25,2))', GREEN, NUM2)
    put(ws4, r, 23, dur, BLUE, NUM2)
    put(ws4, r, 24, gl, BLUE, align="center")
    put(ws4, r, 25, glp, BLUE, align="center")
H1 = H0 + len(H) - 1
HL = 500  # 집계 범위 (추가 입력 여유)
t4 = H1 + 1
put(ws4, t4, 5, "합계", bold=True, fill=TOT)
for c in range(1, 26):
    ws4.cell(row=t4, column=c).fill = TOT
for c in (11, 12, 13, 14, 15, 16, 17, 18):
    col = CL(c)
    put(ws4, t4, c, f"=SUM({col}{H0}:{col}{H1})", fmt=NUM, bold=True, fill=TOT).border = TOPLINE
notes4 = [
    "※ 장부금액(총액): FVPL·FVOCI는 공정가치, AC·원가모형(투자부동산)은 상각후원가(감가상각 후 금액).",
    "※ 순장부금액: AC 자산은 기대신용손실(ECL)을 차감. FVOCI의 ECL은 기타포괄손익에 반영되어 장부금액에서 차감하지 않음.",
    "※ 주식·수익증권은 상각후원가 칸에 취득원가를 입력. 원가모형 투자부동산의 평가손익은 미인식 평가차익(주석 공시용).",
    "※ 행을 추가할 때는 합계 행 위에 삽입하세요. 요약 시트는 6~500행을 집계합니다.",
]
for i, n in enumerate(notes4):
    ws4.cell(row=t4 + 2 + i, column=1, value=n).font = f("595959", size=9)
ws4.freeze_panes = "F6"
ws4.auto_filter.ref = f"A5:Y{H1}"
widths(ws4, [10, 8, 10, 15, 28, 10, 9, 6, 11, 10, 11, 11, 11, 12, 13, 11, 12, 10, 8, 8, 11, 9, 8, 11, 11])
CLASSES = ["주식", "수익증권", "국공채", "회사채", "외화채권", "대출채권", "투자부동산"]
IFRS = ["FVPL", "FVOCI", "AC", "원가모형"]
dv1 = DataValidation(type="list", formula1='"' + ",".join(CLASSES) + '"', allow_blank=True)
dv2 = DataValidation(type="list", formula1='"' + ",".join(IFRS) + '"', allow_blank=True)
dv3 = DataValidation(type="list", formula1='"일반계정,특별계정"', allow_blank=True)
for d, rng in ((dv1, f"F{H0}:F{H1}"), (dv2, f"G{H0}:G{H1}"), (dv3, f"C{H0}:C{H1}")):
    ws4.add_data_validation(d); d.add(rng)

# ------------------------------------------------------------------ 자산운용요약
ws5 = wb.create_sheet("자산운용요약")
title(ws5, "운용자산 요약 및 GL 대사", "자산운용명세 시트를 자동 집계 · 금액 단위: 백만원")
ws5["A3"] = f'="회사: "&{P["회사명"]}&" ("&{P_CO}&")   기준일: "&TEXT({P_DATE},"yyyy-mm-dd")'
ws5["A3"].font = f(GREEN, size=9)
R = lambda col: f"자산운용명세!${col}${H0}:${col}${HL}"

ws5["A5"] = "1. 자산군별 현황"; ws5["A5"].font = f(bold=True, size=11)
header(ws5, 6, ["자산군", "종목수", "장부금액(총액)", "ECL(AC 차감분)", "순장부금액", "비중", "평가손익",
                "가중평균\n표면금리", "가중평균\n수정듀레이션"])
a0 = 7
for i, cls in enumerate(CLASSES):
    r = a0 + i
    put(ws5, r, 1, cls, BLUE, bold=True)
    put(ws5, r, 2, f'=COUNTIF({R("F")},A{r})', GREEN, NUM)
    put(ws5, r, 3, f'=SUMIFS({R("N")},{R("F")},A{r})', GREEN, NUM)
    put(ws5, r, 4, f'=SUMIFS({R("P")},{R("F")},A{r},{R("G")},"AC")', GREEN, NUM)
    put(ws5, r, 5, f"=C{r}-D{r}", fmt=NUM)
    put(ws5, r, 7, f'=SUMIFS({R("O")},{R("F")},A{r})', GREEN, NUM)
    put(ws5, r, 8, f'=IFERROR(SUMPRODUCT(({R("F")}=A{r})*{R("N")}*{R("S")})/SUMPRODUCT(({R("F")}=A{r})*({R("S")}<>"")*{R("N")}),0)', GREEN, PCT2)
    put(ws5, r, 9, f'=IFERROR(SUMPRODUCT(({R("F")}=A{r})*{R("N")}*{R("W")})/SUMPRODUCT(({R("F")}=A{r})*({R("W")}<>"")*{R("N")}),0)', GREEN, NUM2)
a1 = a0 + len(CLASSES) - 1
at = a1 + 1
for r in range(a0, at):
    put(ws5, r, 6, f"=IF($E${at}=0,0,E{r}/$E${at})", fmt=PCT)
put(ws5, at, 1, "합계", bold=True, fill=TOT)
for c in (2, 3, 4, 5, 6, 7):
    col = CL(c)
    put(ws5, at, c, f"=SUM({col}{a0}:{col}{a1})", fmt=PCT if c == 6 else NUM, bold=True, fill=TOT).border = TOPLINE
put(ws5, at, 8, f'=IFERROR(SUMPRODUCT({R("N")}*{R("S")})/SUMPRODUCT(({R("S")}<>"")*{R("N")}),0)', GREEN, PCT2, bold=True, fill=TOT).border = TOPLINE
put(ws5, at, 9, f'=IFERROR(SUMPRODUCT({R("N")}*{R("W")})/SUMPRODUCT(({R("W")}<>"")*{R("N")}),0)', GREEN, NUM2, bold=True, fill=TOT).border = TOPLINE

b = at + 3
ws5.cell(row=b, column=1, value="2. IFRS 9 분류별 현황").font = f(bold=True, size=11)
header(ws5, b + 1, ["IFRS 9 분류", "종목수", "장부금액(총액)", "공정가치", "비중(장부금액)", "평가손익", "손익 반영처"])
where = {"FVPL": "당기손익", "FVOCI": "기타포괄손익(OCI)", "AC": "미반영 (주석 공시)", "원가모형": "미반영 (주석 공시)"}
b0 = b + 2
for i, k in enumerate(IFRS):
    r = b0 + i
    put(ws5, r, 1, k, BLUE, bold=True)
    put(ws5, r, 2, f'=COUNTIF({R("G")},A{r})', GREEN, NUM)
    put(ws5, r, 3, f'=SUMIFS({R("N")},{R("G")},A{r})', GREEN, NUM)
    put(ws5, r, 4, f'=SUMIFS({R("M")},{R("G")},A{r})', GREEN, NUM)
    put(ws5, r, 6, f'=SUMIFS({R("O")},{R("G")},A{r})', GREEN, NUM)
    put(ws5, r, 7, where[k], BLUE)
b1 = b0 + len(IFRS) - 1
bt = b1 + 1
for r in range(b0, bt):
    put(ws5, r, 5, f"=IF($C${bt}=0,0,C{r}/$C${bt})", fmt=PCT)
put(ws5, bt, 1, "합계", bold=True, fill=TOT)
for c in (2, 3, 4, 5, 6):
    col = CL(c)
    put(ws5, bt, c, f"=SUM({col}{b0}:{col}{b1})", fmt=PCT if c == 5 else NUM, bold=True, fill=TOT).border = TOPLINE
put(ws5, bt, 7, None, fill=TOT)

g = bt + 3
ws5.cell(row=g, column=1, value="3. 명세 ↔ 시산표 GL 대사 (Sub-ledger vs General Ledger)").font = f(bold=True, size=11)
header(ws5, g + 1, ["G/L 계정", "계정명", "명세 집계 기준", "명세 합계 (A)", "시산표 기말잔액 (B)", "차이 (A−B)", "판정"])
recon = [(gl, "자산 G/L별 장부금액(총액)", "asset") for gl in
         ["11101000", "11102000", "11201000", "11202000", "11203000", "11301000", "11401000", "11402000", "11501000"]]
recon += [("11309000", "충당금 G/L별 ECL (음수)", "contra"), ("11409000", "충당금 G/L별 ECL (음수)", "contra"),
          ("11601000", "미수이자 전체", "accr")]
TB_E = f"시산표!$E${TB0}:$E${TB1}"; TB_L = f"시산표!$L${TB0}:$L${TB1}"
g0 = g + 2
for i, (gl, basis, kind) in enumerate(recon):
    r = g0 + i
    put(ws5, r, 1, gl, BLUE, align="center")
    put(ws5, r, 2, coa_lookup(f"A{r}", "B"), GREEN)
    put(ws5, r, 3, basis, BLUE)
    if kind == "asset":
        fa = f'=SUMIFS({R("N")},{R("X")},A{r})'
    elif kind == "contra":
        fa = f'=-SUMIFS({R("P")},{R("Y")},A{r})'
    else:
        fa = f'=SUMPRODUCT((LEN({R("X")})>0)*{R("R")})'
    put(ws5, r, 4, fa, GREEN, NUM)
    put(ws5, r, 5, f"=SUMIFS({TB_L},{TB_E},A{r})", GREEN, NUM)
    put(ws5, r, 6, f"=D{r}-E{r}", fmt=NUM)
    put(ws5, r, 7, f'=IF(ABS(F{r})<0.5,"OK","차이 확인")', align="center")
g1 = g0 + len(recon) - 1
gt = g1 + 1
put(ws5, gt, 1, "합계", bold=True, fill=TOT); put(ws5, gt, 2, None, fill=TOT); put(ws5, gt, 3, None, fill=TOT)
for c in (4, 5, 6):
    col = CL(c)
    put(ws5, gt, c, f"=SUM({col}{g0}:{col}{g1})", fmt=NUM, bold=True, fill=TOT).border = TOPLINE
put(ws5, gt, 7, f'=IF(COUNTIF(G{g0}:G{g1},"OK")={len(recon)},"전체 일치","차이 있음")', bold=True, fill=TOT, align="center")
ws5.cell(row=gt + 2, column=1, value="※ 차이가 나면 미전기 거래, 평가 전표 누락, 외화환산 시점 차이, G/L 매핑 오류 순으로 확인합니다.").font = f("595959", size=9)
widths(ws5, [14, 34, 26, 15, 17, 13, 18, 12, 12])
ws5.freeze_panes = "A5"

for ws in wb.worksheets:
    ws.sheet_view.showGridLines = False
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print("saved", OUT)
print("RE plug", begin["31301000"])
print({k: begin[k] for k in end})
